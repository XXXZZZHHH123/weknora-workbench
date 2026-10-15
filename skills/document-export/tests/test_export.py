import importlib.util
from pathlib import Path
import shutil

import pytest
from docx import Document
from pypdf import PdfReader

spec = importlib.util.spec_from_file_location("document_export", Path(__file__).parents[1] / "scripts/export_document.py")
renderer = importlib.util.module_from_spec(spec)
spec.loader.exec_module(renderer)


def content():
    return {"title": "数据核对报告", "subtitle": "独立测试内容", "sections": [
        {"heading": "核对结果", "paragraphs": ["本段保留单位、缺失范围及原始来源，不推断未提供的信息。"],
         "table": {"headers": ["指标", "原值", "来源"], "rows": [["样本甲", "12.50 kg", "测试.xlsx 数据页 C7"], ["样本乙", "缺失", "测试.xlsx 数据页 D7"]]}}
    ], "sources": ["测试.xlsx 数据页 A7:D7；文件指纹以实际来源为准"]}


def test_word_preserves_values_and_source_coordinates(tmp_path):
    path = tmp_path / "report.docx"
    renderer.build_document(content(), path)
    doc = Document(path)
    assert doc.paragraphs[0].text == content()["title"]
    assert doc.tables[0].rows[1].cells[1].text == "12.50 kg"
    assert doc.tables[0].rows[1].cells[2].text.endswith("C7")
    assert doc.tables[0].rows[2].cells[1].text == "缺失"


@pytest.mark.parametrize("bad", [
    {"title": "标题", "sections": []},
    {"title": "标题", "sections": [{"paragraphs": [3]}]},
    {"title": "标题", "sections": [{"table": {"headers": ["甲", "乙"], "rows": [["少一列"]]}}]},
    {"title": "标题\x01", "sections": [{}]},
])
def test_invalid_content_does_not_create_deliverable(tmp_path, bad):
    with pytest.raises(ValueError):
        renderer.build_document(bad, tmp_path / "report.docx")
    assert not (tmp_path / "report.docx").exists()


@pytest.mark.parametrize("name", ["../escape", "/etc/report", "a/b", "a\\b", ".", ""])
def test_output_path_cannot_escape_destination(tmp_path, name):
    with pytest.raises(ValueError, match="文件名"):
        renderer.export(content(), tmp_path, name, ["docx"])
    assert not list(tmp_path.iterdir())


def test_missing_converter_never_reports_success(tmp_path, monkeypatch):
    monkeypatch.setattr(shutil, "which", lambda _: None)
    with pytest.raises(ValueError, match="LibreOffice"):
        renderer.export(content(), tmp_path, "report", ["docx", "pdf"])
    assert not list(tmp_path.iterdir())


@pytest.mark.skipif(not (shutil.which("libreoffice") or shutil.which("soffice")), reason="requires LibreOffice Writer")
def test_actual_conversion_reopens_both_files_and_preserves_chinese(tmp_path):
    result = renderer.export(content(), tmp_path, "report", ["docx", "pdf"])
    assert result["status"] == "created" and result["pages"] >= 1
    assert len(result["files"]) == 2
    for entry in result["files"]:
        path = Path(entry["path"])
        assert path.stat().st_size == entry["bytes"] > 0
        assert len(entry["sha256"]) == 64
        if path.suffix == ".pdf":
            text = "".join(page.extract_text() for page in PdfReader(path).pages)
            assert "12.50" in text and "C7" in text and "来源" in text
        else:
            assert Document(path).tables
    second = renderer.export(content(), tmp_path, "report", ["docx"])
    assert second["files"][0]["path"] != result["files"][0]["path"]
    assert len(list(tmp_path.iterdir())) == 3, "retries preserve earlier files; PDF validation remains private"
