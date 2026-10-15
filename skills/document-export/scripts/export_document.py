"""Render model-authored document content; no question interpretation or research."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import uuid

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader


def text(value, field):
    if not isinstance(value, str):
        raise ValueError(f"{field} 必须为文字")
    # XML 1.0 cannot encode these controls. Reject rather than silently altering evidence.
    if any(ord(c) < 32 and c not in "\n\r\t" for c in value):
        raise ValueError(f"{field} 含不支持的控制字符")
    return value


def build_document(spec, target):
    if not isinstance(spec, dict) or not text(spec.get("title"), "title").strip():
        raise ValueError("需要非空文档标题")
    sections = spec.get("sections", [])
    sources = spec.get("sources", [])
    if not isinstance(sections, list) or not isinstance(sources, list):
        raise ValueError("sections 和 sources 必须为列表")
    if not sections:
        raise ValueError("文档没有正文章节")
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2"]:
        style = document.styles[style_name]
        style.font.name = "Noto Sans CJK SC"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal = document.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    document.styles["Title"].font.size = Pt(22)
    document.styles["Heading 1"].font.size = Pt(14)
    document.add_paragraph(spec["title"], "Title")
    if spec.get("subtitle"):
        document.add_paragraph(text(spec["subtitle"], "subtitle"), "Subtitle")
    for entry in sections:
        if not isinstance(entry, dict):
            raise ValueError("章节必须为对象")
        if entry.get("heading"):
            document.add_heading(text(entry["heading"], "heading"), 1)
        paragraphs = entry.get("paragraphs", [])
        if not isinstance(paragraphs, list):
            raise ValueError("paragraphs 必须为列表")
        for paragraph in paragraphs:
            document.add_paragraph(text(paragraph, "paragraph"))
        if entry.get("table") is not None:
            data = entry["table"]
            if not isinstance(data, dict):
                raise ValueError("table 必须为对象")
            headers, rows = data.get("headers"), data.get("rows")
            if not isinstance(headers, list) or not 1 <= len(headers) <= 8 or not isinstance(rows, list):
                raise ValueError("表格需要 1–8 列 headers 以及 rows 列表；宽表请拆分")
            table = document.add_table(rows=1, cols=len(headers))
            table.style = "Light Shading Accent 1"
            for cell, value in zip(table.rows[0].cells, headers):
                cell.text = text(value, "table.headers")
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows:
                if not isinstance(row, list) or len(row) != len(headers):
                    raise ValueError("表格行列数与表头不一致")
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = text(value, "table.rows")
            document.add_paragraph()
    if sources:
        document.add_heading("来源", 1)
        for source in sources:
            document.add_paragraph(text(source, "sources"))
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.save(target)
    reopened = Document(target)
    if not reopened.paragraphs or reopened.paragraphs[0].text != spec["title"]:
        raise ValueError("Word 内容重新读取失败")


def export(spec, output_dir, name, formats):
    if not name or Path(name).name != name or name in {".", ".."} or any(c in name for c in "/\\\x00"):
        raise ValueError("文件名不能包含路径")
    if not formats or not set(formats) <= {"docx", "pdf"}:
        raise ValueError("仅支持 docx 和 pdf")
    office = shutil.which("libreoffice") or shutil.which("soffice")
    if not office:
        raise ValueError("缺少 LibreOffice Writer，文件尚未生成")
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    # Unique immutable names preserve previous deliverables across retries/turns.
    stem = f"{name}-{uuid.uuid4().hex[:10]}"
    with tempfile.TemporaryDirectory(prefix="document-export-") as temporary:
        work = Path(temporary)
        docx_path = work / f"{stem}.docx"
        build_document(spec, docx_path)
        result = subprocess.run(
            [office, f"-env:UserInstallation={(work / 'profile').as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", str(work), str(docx_path)],
            capture_output=True, text=True, timeout=120,
        )
        pdf_path = work / f"{stem}.pdf"
        if result.returncode != 0 or not pdf_path.is_file():
            raise ValueError("Word 转 PDF 失败；未交付不完整文件。" + result.stderr[-500:])
        pdf = PdfReader(pdf_path, strict=True)
        if not pdf.pages or not any((page.extract_text() or "").strip() for page in pdf.pages):
            raise ValueError("PDF 没有可读取的正文")
        files = []
        for fmt in dict.fromkeys(formats):
            original = work / f"{stem}.{fmt}"
            content = original.read_bytes()
            if not content:
                raise ValueError("生成文件为空")
            target = output_dir / original.name
            # Exclusive create also refuses symlinks and accidental overwrites.
            with target.open("xb") as stream:
                stream.write(content)
            files.append({"path": str(target), "bytes": len(content), "sha256": hashlib.sha256(content).hexdigest()})
        return {"status": "created", "pages": len(pdf.pages), "files": files}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--name", required=True)
    parser.add_argument("--formats", nargs="+", choices=["docx", "pdf"], default=["docx", "pdf"])
    parser.add_argument("--output-dir", type=Path, default=Path("/workspace/output"))
    args = parser.parse_args()
    if args.input.stat().st_size > 2 * 1024 * 1024:
        raise ValueError("文档输入超过 2 MiB，请拆分文档")
    print(json.dumps(export(json.loads(args.input.read_text()), args.output_dir, args.name, args.formats), ensure_ascii=False))


if __name__ == "__main__":
    main()
