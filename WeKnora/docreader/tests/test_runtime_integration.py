"""Local gRPC parsing without storage, model providers, or external sources."""

import io
import os
from pathlib import Path
import secrets
import socket
import subprocess
import sys
import tempfile
import unittest

import grpc
from grpc_health.v1 import health_pb2, health_pb2_grpc
from docx import Document

from docreader.proto import docreader_pb2, docreader_pb2_grpc


WEKNORA = Path(__file__).resolve().parents[2]
REPO = WEKNORA.parent


def text_pdf(text):
    content = b"BT /F1 18 Tf 50 720 Td (" + text.encode("ascii") + b") Tj ET\n"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"endstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>",
    ]
    out = bytearray(b"%PDF-1.7\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out.extend(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = len(out)
    out.extend(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.extend(b"%010d 00000 n \n" % offset)
    out.extend(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return bytes(out)


class LocalRuntimeIntegrationTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        with socket.socket() as sock:
            sock.bind(("127.0.0.1", 0))
            port = sock.getsockname()[1]
        cls.address = f"127.0.0.1:{port}"
        cls.token = secrets.token_urlsafe(32)
        cls.metadata = (("authorization", "Bearer " + cls.token),)
        cls.env = {
            "PATH": os.environ.get("PATH", "/usr/bin:/bin"), "LANG": "C.UTF-8",
            "PYTHONPATH": str(WEKNORA), "GRPC_AUTH_TOKEN": cls.token,
            "DOCREADER_GRPC_HOST": "127.0.0.1", "DOCREADER_GRPC_PORT": str(port),
            "DOCREADER_ADDR": cls.address, "DOCREADER_ODL_HYBRID": "off",
            "HTTP_PROXY": "http://127.0.0.1:1", "HTTPS_PROXY": "http://127.0.0.1:1",
            "NO_PROXY": "127.0.0.1,localhost", "DOCREADER_PDF_RENDER_PARALLELISM": "1",
            "LOG_LEVEL": "ERROR",
        }
        cls.log = open(Path(cls.tmp.name) / "server.log", "w+")
        cls.server = subprocess.Popen([sys.executable, "-m", "docreader.main"], cwd=WEKNORA,
                                      env=cls.env, stdout=cls.log, stderr=subprocess.STDOUT)
        cls.channel = grpc.insecure_channel(cls.address)
        try:
            grpc.channel_ready_future(cls.channel).result(timeout=30)
        except Exception:
            cls.server.terminate()
            cls.server.wait(timeout=5)
            cls.log.seek(0)
            detail = cls.log.read()[-3000:]
            cls.log.close()
            cls.tmp.cleanup()
            raise RuntimeError("DocReader did not start: " + detail)
        cls.stub = docreader_pb2_grpc.DocReaderStub(cls.channel)

    @classmethod
    def tearDownClass(cls):
        cls.channel.close()
        cls.server.terminate()
        cls.server.wait(timeout=5)
        cls.log.close()
        cls.tmp.cleanup()

    def test_health_and_authenticated_registry(self):
        health = health_pb2_grpc.HealthStub(self.channel).Check(health_pb2.HealthCheckRequest(), timeout=3)
        self.assertEqual(health.status, health_pb2.HealthCheckResponse.SERVING)
        result = self.stub.ListEngines(docreader_pb2.ListEnginesRequest(), metadata=self.metadata, timeout=3)
        self.assertTrue(any(engine.name == "builtin" and engine.available for engine in result.engines))
        probe = subprocess.run([sys.executable, str(REPO / "scripts/check-local-docreader.py")],
                               env=self.env, capture_output=True, timeout=10)
        self.assertEqual(probe.returncode, 0, probe.stderr.decode())

    def test_server_listens_on_configured_ipv4_loopback_only(self):
        with socket.socket(socket.AF_INET6) as sock:
            sock.settimeout(1)
            with self.assertRaises(OSError):
                sock.connect(("::1", int(self.address.rsplit(":", 1)[1])))

    def test_registry_rejects_missing_or_wrong_authentication(self):
        for metadata in ((), (("authorization", "Bearer invalid-fixture-token"),)):
            with self.subTest(metadata=bool(metadata)):
                with self.assertRaises(grpc.RpcError) as error:
                    self.stub.ListEngines(docreader_pb2.ListEnginesRequest(), metadata=metadata, timeout=3)
                self.assertEqual(error.exception.code(), grpc.StatusCode.UNAUTHENTICATED)
        probe = subprocess.run([sys.executable, str(REPO / "scripts/check-local-docreader.py")],
                               env={**self.env, "GRPC_AUTH_TOKEN": "invalid-fixture-token"},
                               capture_output=True, timeout=10)
        self.assertNotEqual(probe.returncode, 0)

    def test_parses_plain_markdown_docx_and_native_text_pdf(self):
        document = Document()
        document.add_heading("Local parser validation", 1)
        document.add_paragraph("文档解析离线验收：保留中文内容。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "字段"
        table.cell(0, 1).text = "值"
        table.cell(1, 0).text = "测试编号"
        table.cell(1, 1).text = "DOCX-42"
        output = io.BytesIO()
        document.save(output)
        fixtures = [
            ("plain.md", "md", "# 纯文本解析\n\n离线文本内容保留。\n".encode(), "离线文本内容保留"),
            ("table.docx", "docx", output.getvalue(), "DOCX-42"),
            ("native.pdf", "pdf", text_pdf("Native PDF parser validation"), "Native PDF parser validation"),
        ]
        for name, file_type, content, expected in fixtures:
            for streaming in (False, True):
                with self.subTest(file_type=file_type, streaming=streaming):
                    request = docreader_pb2.ReadRequest(file_name=name, file_type=file_type, file_content=content)
                    if streaming:
                        responses = list(self.stub.ReadStream(request, metadata=self.metadata, timeout=20))
                        result = responses[0].meta
                        self.assertEqual(len(responses), 1)
                    else:
                        result = self.stub.Read(request, metadata=self.metadata, timeout=20)
                        self.assertEqual(len(result.image_refs), 0)
                    self.assertEqual(result.error, "")
                    self.assertIn(expected, result.markdown_content)

    def test_invalid_document_is_reported_and_server_stays_usable(self):
        response = self.stub.Read(docreader_pb2.ReadRequest(file_name="invalid.pdf", file_type="pdf",
                                                           file_content=b"not a PDF"), metadata=self.metadata, timeout=10)
        self.assertTrue(response.error)
        response = self.stub.Read(docreader_pb2.ReadRequest(file_name="after.md", file_type="md",
                                                           file_content=b"Service remains ready after invalid input."),
                                  metadata=self.metadata, timeout=10)
        self.assertEqual(response.error, "")
        self.assertIn("remains ready", response.markdown_content)


if __name__ == "__main__":
    unittest.main()
