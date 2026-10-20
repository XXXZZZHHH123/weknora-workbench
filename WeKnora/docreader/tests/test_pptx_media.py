import base64
import io
import subprocess
import struct
import time
import re
import shutil
import unittest
import zipfile
from unittest.mock import patch

from docreader.parser import pptx_media as media

P = media._NS["p"]
A = media._NS["a"]
R = media._NS["r"]
REL = media._NS["rel"]


def picture(identifier, name, rid):
    return (f'<p:pic><p:nvPicPr><p:cNvPr id="{identifier}" name="{name}"/>'
            f'</p:nvPicPr><p:blipFill><a:blip r:embed="{rid}"/></p:blipFill></p:pic>')


def fixture():
    parts = {
        # Include an unused asset before the real pictures: it must never be mapped.
        "ppt/media/unrelated.emf": b"wrong-report",
        "ppt/media/chart-b.emf": b"chart-b",
        "ppt/media/logo.png": b"logo",
        "ppt/media/chart-a.emf": b"chart-a",
        "ppt/presentation.xml": (f'<p:presentation xmlns:p="{P}" xmlns:r="{R}">'
                                 '<p:sldIdLst><p:sldId r:id="first"/>'
                                 '<p:sldId r:id="second"/></p:sldIdLst></p:presentation>').encode(),
        "ppt/_rels/presentation.xml.rels": (f'<Relationships xmlns="{REL}">'
                                          '<Relationship Id="first" Target="slides/slide9.xml"/>'
                                          '<Relationship Id="second" Target="slides/slide2.xml"/>'
                                          '</Relationships>').encode(),
        "ppt/slides/slide9.xml": (f'<p:sld xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}">'
                                 + picture(1, "Company logo", "logo")
                                 + '<p:graphicFrame><p:oleObj r:id="excel">'
                                 + picture(2, "01-05 chart", "chart")
                                 + '</p:oleObj></p:graphicFrame></p:sld>').encode(),
        "ppt/slides/_rels/slide9.xml.rels": (f'<Relationships xmlns="{REL}">'
                                           '<Relationship Id="logo" Target="../media/logo.png"/>'
                                           '<Relationship Id="chart" Target="../media/chart-a.emf"/>'
                                           '<Relationship Id="excel" TargetMode="External" '
                                           'Target="file:///C:/private.xlsx"/>'
                                           '</Relationships>').encode(),
        "ppt/slides/slide2.xml": (f'<p:sld xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}">'
                                 + '<p:grpSp>' + picture(3, "Other chart", "chart")
                                 + '</p:grpSp>' + picture(4, "External image", "external")
                                 + '</p:sld>').encode(),
        "ppt/slides/_rels/slide2.xml.rels": (f'<Relationships xmlns="{REL}">'
                                           '<Relationship Id="chart" Target="../media/chart-b.emf"/>'
                                           '<Relationship Id="external" TargetMode="External" '
                                           'Target="https://example.test/private.png"/>'
                                           '</Relationships>').encode(),
    }
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return out.getvalue()


MARKDOWN = ("<!-- Slide number: 1 -->\n# First title\nFirst text\n"
            "![Company logo](data:image/png;base64,bG9nbw==)\n"
            "### Notes:\nFirst notes\n\n"
            "<!-- Slide number: 2 -->\n# Second title\nSecond text\n"
            "![Other chart](wrong.jpg)\n")


class TestPptxMedia(unittest.TestCase):
    def test_trim_applies_only_to_converted_canvas_and_preserves_alpha(self):
        from PIL import Image, ImageDraw
        canvas = Image.new("RGB", (100, 120), "white")
        ImageDraw.Draw(canvas).rectangle((30, 40, 59, 69), fill="black")
        raw = io.BytesIO()
        canvas.save(raw, format="PNG")
        original = Image.open(io.BytesIO(media._rasterize_with_pillow(raw.getvalue())))
        cropped = Image.open(io.BytesIO(media._rasterize_with_pillow(raw.getvalue(), trim_canvas=True)))
        self.assertEqual(original.size, (100, 120))
        self.assertEqual(cropped.size, (46, 46))
        logo = Image.new("RGBA", (8, 8), (255, 0, 0, 128))
        raw = io.BytesIO()
        logo.save(raw, format="PNG")
        preserved = Image.open(io.BytesIO(media._rasterize_with_pillow(raw.getvalue())))
        self.assertEqual(preserved.getpixel((0, 0)), (255, 0, 0, 128))

    def test_emf_exact_title_joins_only_adjacent_original_text_runs(self):
        header = bytearray(88)
        struct.pack_into("<II", header, 0, 1, 88)
        header[40:44] = b" EMF"

        def text_record(text, x, y, width, flags=0):
            encoded = text.encode("utf-16-le")
            chars = len(encoded) // 2
            advance_offset = (76 + len(encoded) + 3) // 4 * 4
            size = advance_offset + chars * 4
            record = bytearray(size)
            struct.pack_into("<II", record, 0, 84, size)
            struct.pack_into("<iiIII", record, 36, x, y, chars, 76, flags)
            struct.pack_into("<I", record, 72, advance_offset)
            record[76:76 + len(encoded)] = encoded
            struct.pack_into("<" + "i" * chars, record, advance_offset, *([width] * chars))
            return bytes(record)

        emf = bytes(header) + b"".join([
            text_record("郑棉", 346, 40, 25), text_record("01", 396, 40, 13),
            text_record("-", 422, 40, 8), text_record("05", 430, 40, 13),
            text_record("价差", 456, 40, 25), text_record("2401", 800, 423, 10),
            text_record("glyph indices", 1, 1, 10, flags=0x10),
        ])
        self.assertEqual(media.extract_emf_labels(emf), ["郑棉01-05价差", "2401"])
        self.assertEqual(media.extract_emf_labels(b"untrusted bytes"), [])
        self.assertEqual(media.extract_emf_labels(emf[:-10]), ["郑棉01-05价差", "2401"])
        # Excel emits runs at (0,0), positioned with MWT_LEFTMULTIPLY.
        def transform(mode, x=0, y=0):
            return struct.pack("<IIffffffI", 36, 36, 1, 0, 0, 1, x, y, mode)
        positioned = bytes(header) + b"".join([
            transform(2, 346, 40), text_record("郑棉", 0, 0, 25),
            transform(1), transform(2, 396, 40), text_record("01", 0, 0, 13),
        ])
        self.assertEqual(media.extract_emf_labels(positioned), ["郑棉01"])

    def test_batch_converts_only_supplied_internal_caches_and_validates_outputs(self):
        from PIL import Image
        from pathlib import Path
        calls = []

        def convert(command, timeout):
            calls.append(command)
            for value in command[command.index("--outdir") + 2:]:
                Image.new("RGB", (8, 8)).save(Path(value).with_suffix(".png"))
            return True

        with zipfile.ZipFile(io.BytesIO(fixture())) as archive, \
                patch.object(media, "find_soffice", return_value="/test/soffice"), \
                patch.object(media, "_run_conversion", side_effect=convert):
            images = media._rasterize_vector_batches(
                archive, ["ppt/media/chart-a.emf", "ppt/media/chart-b.emf"], time.monotonic() + 10)
        self.assertEqual(set(images), {"ppt/media/chart-a.emf", "ppt/media/chart-b.emf"})
        self.assertEqual(len(calls), 1)
        self.assertNotIn("unrelated", " ".join(calls[0]))
        self.assertNotIn("xlsx", " ".join(calls[0]))

    def test_relationships_keep_each_chart_on_correct_display_page(self):
        calls = []
        def raster(name, data, **kwargs):
            calls.append((name, data))
            return b"png:" + data
        with patch.object(media, "rasterize_media_bytes", side_effect=raster), \
                patch.object(media, "_rasterize_vector_batches", return_value={}):
            text, images = media.attach_pptx_media_to_markdown(MARKDOWN, fixture())
        first, second = text.split("<!-- Slide number: 2 -->")
        self.assertIn("# 第1页", first)
        self.assertIn("## First title", first)
        self.assertIn("First notes", first)
        self.assertIn("ppt/media/chart-a.emf", first)
        self.assertNotIn("chart-b.emf", first)
        self.assertIn("# 第2页", second)
        self.assertIn("ppt/media/chart-b.emf", second)
        self.assertNotIn("chart-a.emf", second)
        self.assertEqual(first.count("![第1页 · Company logo]"), 1)
        self.assertNotIn("data:image", text)
        self.assertNotIn("wrong.jpg", text)
        self.assertNotIn("C:/private", text)
        self.assertNotIn("https://", text)
        self.assertIn("未能解析", second)  # Explicit external image unavailability.
        self.assertEqual([data for _, data in calls], [b"logo", b"chart-a", b"chart-b"])
        self.assertEqual({base64.b64decode(v) for v in images.values()},
                         {b"png:logo", b"png:chart-a", b"png:chart-b"})

    def test_inconsistent_page_markers_cannot_guess_image_assignment(self):
        with self.assertRaisesRegex(ValueError, "refusing"):
            media.attach_pptx_media_to_markdown("one unidentified slide", fixture())
        with self.assertRaisesRegex(ValueError, "refusing"):
            media.attach_pptx_media_to_markdown(MARKDOWN.replace("number: 2", "number: 9"), fixture())

    def test_failed_cache_is_explicit_and_text_survives(self):
        with patch.object(media, "rasterize_media_bytes", return_value=None), \
                patch.object(media, "_rasterize_vector_batches", return_value={}):
            text, images = media.attach_pptx_media_to_markdown(MARKDOWN, fixture())
        self.assertEqual({base64.b64decode(v) for v in images.values()}, {b"chart-a", b"chart-b"})
        self.assertTrue(all(ref.endswith(".emf") for ref in images))
        self.assertIn("First text", text)
        self.assertIn("不能据此判断图例或数值", text)
        self.assertIn("此原始矢量图尚未生成预览", text)
        self.assertNotIn("原图曲线几何证据", text)

    def test_exhausted_vector_budget_preserves_existing_raster_logo(self):
        with patch.object(media, "_MEDIA_BUDGET_SECONDS", -1), \
                patch.object(media, "_rasterize_vector_batches", return_value={}), \
                patch.object(media, "rasterize_media_bytes", return_value=b"logo-png") as raster:
            text, images = media.attach_pptx_media_to_markdown(MARKDOWN, fixture())
        self.assertEqual(len(images), 3)
        self.assertIn("![第1页 · Company logo]", text)
        self.assertIn("此原始矢量图尚未生成预览", text)
        self.assertEqual({base64.b64decode(v) for ref,v in images.items() if ref.endswith('.emf')}, {b"chart-a",b"chart-b"})
        self.assertEqual(raster.call_args.args[:2], ("ppt/media/logo.png", b"logo"))
        self.assertEqual(raster.call_count, 1)

    def test_ole_without_markdown_placeholder_still_triggers_attachment(self):
        self.assertTrue(media.markdown_needs_pptx_media_attach(
            "<!-- Slide number: 12 -->\nWorksheet\n"))

    def test_emf_prefers_standalone_libreoffice_not_imagemagick(self):
        with patch.object(media, "_rasterize_with_libreoffice", return_value=b"png") as office, \
                patch.object(media, "_rasterize_with_imagemagick") as imagemagick:
            self.assertEqual(media.rasterize_media_bytes("image84.emf", b"emf"), b"png")
        office.assert_called_once()
        imagemagick.assert_not_called()

    def test_size_and_extension_limits_before_launching_converter(self):
        with patch.object(media, "_rasterize_with_libreoffice") as office:
            self.assertIsNone(media.rasterize_media_bytes("large.emf", b"0" * (media._MAX_MEDIA_BYTES + 1)))
            self.assertIsNone(media.rasterize_media_bytes("ole.bin", b"private workbook"))
        office.assert_not_called()

    def test_svg_external_reference_cannot_be_fetched(self):
        with patch.object(media, "_rasterize_with_imagemagick") as converter:
            self.assertIsNone(media.rasterize_media_bytes(
                "external.svg", b'<svg><image href="https://example.test/private.png"/></svg>'))
            self.assertIsNone(media.rasterize_media_bytes(
                "external.svg", b'<svg><style>@import url(https://example.test/style);</style></svg>'))
        converter.assert_not_called()

    def test_conversion_timeout_kills_process_group(self):
        with patch.object(media.subprocess, "Popen") as popen, patch.object(media.os, "killpg") as kill:
            proc = popen.return_value
            proc.pid = 1234
            proc.wait.side_effect = [subprocess.TimeoutExpired("soffice", 1), 0]
            self.assertFalse(media._run_conversion(["soffice", "input.emf"], 1))
        kill.assert_called_once_with(1234, media.signal.SIGKILL)
        self.assertTrue(popen.call_args.kwargs["start_new_session"])

    def test_generated_deck_preserves_text_notes_and_shared_logo_once_per_page(self):
        from PIL import Image
        from pptx import Presentation
        from pptx.util import Inches
        from docreader.parser.markitdown_parser import MarkitdownParser

        png = io.BytesIO()
        Image.new("RGBA", (8, 8), (255, 0, 0, 128)).save(png, format="PNG")
        deck = Presentation()
        for number in (1, 2):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = f"正文 {number}"
            slide.shapes.add_picture(io.BytesIO(png.getvalue()), Inches(1), Inches(2))
            slide.notes_slide.notes_text_frame.text = f"备注 {number}"
        package = io.BytesIO()
        deck.save(package)
        doc = MarkitdownParser(file_type="pptx").parse_into_text(package.getvalue())
        self.assertEqual(len(doc.images), 1)  # Same asset, each slide retains its own reference.
        for number in (1, 2):
            self.assertIn(f"# 第{number}页", doc.content)
            self.assertIn(f"正文 {number}", doc.content)
            self.assertIn(f"备注 {number}", doc.content)
            self.assertEqual(doc.content.count(f"![第{number}页 ·"), 1)
        self.assertNotIn("data:image", doc.content)
from lxml import etree
from markitdown import MarkItDown
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from docreader.parser import pptx_media
from docreader.parser.markitdown_parser import MarkitdownParser, StdMarkitdownParser

RED = (220, 30, 40)
BLUE = (30, 80, 220)
GREEN = (20, 170, 70)
IMAGE = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")


def png(color):
    output = io.BytesIO()
    Image.new("RGB", (16, 16), color).save(output, format="PNG")
    return output.getvalue()


def deck(colors, names=None, grouped=False, reverse_slides=False):
    presentation = Presentation()
    for index, color in enumerate(colors):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shapes = slide.shapes.add_group_shape().shapes if grouped else slide.shapes
        picture = shapes.add_picture(io.BytesIO(png(color)), Inches(1), Inches(1))
        picture.name = names[index] if names else f"picture{index}"
        picture._element._nvXxPr.cNvPr.set("descr", f"position{index}")
    if reverse_slides:
        ids = presentation.slides._sldIdLst
        for child in list(ids):
            ids.remove(child)
            ids.insert(0, child)
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def rewrite_archive(data, *, reverse_media=False, unused=False, corrupt_first=False):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    media = [name for name in files if name.startswith("ppt/media/")]
    if corrupt_first:
        files[media[0]] = b"not an image"
    if reverse_media:
        media.reverse()
    if unused:
        name = "ppt/media/unused.png"
        files[name] = png(GREEN)
        media.insert(0, name)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, blob in files.items():
            if name not in media:
                archive.writestr(name, blob)
        for name in media:
            archive.writestr(name, files[name])
    return output.getvalue()


def fallback(data):
    markdown = (
        MarkItDown()
        .convert(io.BytesIO(data), file_extension=".pptx", keep_data_uris=False)
        .text_content
    )
    return pptx_media.attach_pptx_media_to_markdown(markdown, data)


def image_colors(markdown, images):
    colors = []
    for target in IMAGE.findall(markdown):
        if target not in images:
            colors.append(None)
            continue
        raw = base64.b64decode(images[target])
        with Image.open(io.BytesIO(raw)) as image:
            colors.append(image.convert("RGB").getpixel((8, 8)))
    return colors


class TestPptxMediaAssociation(unittest.TestCase):
    def test_document_order_control(self):
        markdown, images = fallback(deck([RED, BLUE]))
        self.assertEqual(image_colors(markdown, images), [RED, BLUE])

    def test_zip_entry_order_does_not_change_which_picture_is_attached(self):
        data = rewrite_archive(deck([RED, BLUE]), reverse_media=True)
        markdown, images = fallback(data)
        self.assertEqual(image_colors(markdown, images), [RED, BLUE])

    def test_unused_media_does_not_replace_a_visible_picture(self):
        data = rewrite_archive(deck([RED, BLUE]), unused=True)
        markdown, images = fallback(data)
        self.assertEqual(image_colors(markdown, images), [RED, BLUE])

    def test_shared_image_is_attached_at_every_placement(self):
        markdown, images = fallback(deck([RED, RED, BLUE]))
        self.assertEqual(image_colors(markdown, images), [RED, RED, BLUE])

    def test_failed_conversion_does_not_take_the_next_picture(self):
        data = rewrite_archive(deck([RED, BLUE]), corrupt_first=True)
        # A missing converter is a supported deployment condition. The bad
        # PNG still exercises Pillow's actual decode failure.
        with patch.object(pptx_media.shutil, "which", return_value=None):
            markdown, images = fallback(data)
        self.assertEqual(image_colors(markdown, images), [None, BLUE])

    def test_duplicate_shape_names_keep_each_placement(self):
        data = deck([RED, BLUE], names=["Picture 1", "Picture 1"])
        markdown, images = fallback(data)
        self.assertEqual(image_colors(markdown, images), [RED, BLUE])

    def test_names_that_share_a_placeholder_keep_each_placement(self):
        data = rewrite_archive(
            deck([RED, BLUE], names=["Picture-1", "Picture 1"]), reverse_media=True
        )
        markdown, images = fallback(data)
        self.assertEqual(image_colors(markdown, images), [RED, BLUE])

    def test_presentation_order_takes_precedence_over_slide_filenames(self):
        markdown, images = fallback(deck([RED, BLUE], reverse_slides=True))
        self.assertEqual(image_colors(markdown, images), [BLUE, RED])

    def test_same_slide_duplicate_names_follow_position_order(self):
        # MarkItDown orders shapes by (top, left), not z-order.
        for grouped in (False, True):
            with self.subTest(grouped=grouped):
                presentation = Presentation()
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                shapes = (
                    slide.shapes.add_group_shape().shapes if grouped else slide.shapes
                )
                for color, top in ((RED, 4), (BLUE, 1)):
                    picture = shapes.add_picture(
                        io.BytesIO(png(color)), Inches(1), Inches(top)
                    )
                    picture.name = "Picture 1"
                output = io.BytesIO()
                presentation.save(output)
                markdown, images = fallback(output.getvalue())
                self.assertEqual(image_colors(markdown, images), [BLUE, RED])

    def test_grouped_pictures_keep_their_source(self):
        data = rewrite_archive(deck([RED, BLUE], grouped=True), reverse_media=True)
        markdown, images = fallback(data)
        self.assertEqual(image_colors(markdown, images), [RED, BLUE])

    def test_external_picture_does_not_consume_the_next_embedded_image(self):
        presentation = Presentation(io.BytesIO(deck([RED, BLUE])))
        slide = presentation.slides[0]
        picture = slide.shapes[0]
        blip = picture._element.blipFill.blip
        slide.part.drop_rel(blip.rEmbed)
        del blip.attrib[qn("r:embed")]
        relationship = slide.part.relate_to(
            "https://example.invalid/external.png",
            RELATIONSHIP_TYPE.IMAGE,
            is_external=True,
        )
        blip.set(qn("r:link"), relationship)
        output = io.BytesIO()
        presentation.save(output)
        markdown, images = fallback(output.getvalue())
        self.assertEqual(image_colors(markdown, images), [None, BLUE])

    def test_unrelated_links_and_unknown_placeholders_stay_unchanged(self):
        data = deck([RED])
        original = (
            MarkItDown()
            .convert(io.BytesIO(data), file_extension=".pptx", keep_data_uris=False)
            .text_content
        )
        prefix = (
            "![web](https://example.invalid/image.png)\n"
            "![inline](data:image/png;base64,eA==)\n"
            "![unknown](not-a-picture.jpg)\n"
        )
        markdown, images = pptx_media.attach_pptx_media_to_markdown(
            prefix + original, data
        )
        self.assertTrue(markdown.startswith(prefix))
        self.assertEqual(image_colors(markdown, images)[-1], RED)


def deck_with_broken_picture(mutate):
    """One slide: text, a good RED picture, then a BLUE picture `mutate` breaks."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(0.1), Inches(4), Inches(0.5))
    textbox.text_frame.text = "SLIDE TEXT OK"
    good = slide.shapes.add_picture(io.BytesIO(png(RED)), Inches(1), Inches(1))
    good.name = "good"
    bad = slide.shapes.add_picture(io.BytesIO(png(BLUE)), Inches(1), Inches(3))
    bad.name = "bad"
    media = mutate(slide, bad)
    output = io.BytesIO()
    presentation.save(output)
    data = output.getvalue()
    if media is None:
        return data
    # Drop the picture's media part from the package: python-pptx then drops
    # the relationship at load time while the slide XML still points at it.
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    del files[media]
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, blob in files.items():
            archive.writestr(name, blob)
    return output.getvalue()


def remove_blip(slide, picture):
    blip = picture._element.blipFill.blip
    blip.getparent().remove(blip)


def remove_blip_fill(slide, picture):
    blip_fill = picture._element.blipFill
    blip_fill.getparent().remove(blip_fill)


def dangling_embed(slide, picture):
    picture._element.blipFill.blip.set(qn("r:embed"), "rId999")


def external_embed(slide, picture):
    relationship = slide.part.relate_to(
        "https://example.invalid/external.png",
        RELATIONSHIP_TYPE.IMAGE,
        is_external=True,
    )
    picture._element.blipFill.blip.set(qn("r:embed"), relationship)


def missing_media_part(slide, picture):
    part = slide.part.related_part(picture._element.blipFill.blip.rEmbed)
    return str(part.partname).lstrip("/")


class TestPptxMediaBrokenPictureRelationships(unittest.TestCase):
    def test_broken_picture_keeps_the_document_and_other_pictures(self):
        for mutate in (
            remove_blip,
            remove_blip_fill,
            dangling_embed,
            external_embed,
            missing_media_part,
        ):
            with self.subTest(mutate.__name__):
                data = deck_with_broken_picture(mutate)
                document = StdMarkitdownParser(
                    file_name="deck.pptx", file_type="pptx"
                ).parse_into_text(data)
                self.assertIn("SLIDE TEXT OK", document.content)
                self.assertEqual(
                    image_colors(document.content, document.images), [RED, None]
                )
                self.assertIn("(bad.jpg)", document.content)

                parsed = MarkitdownParser(
                    file_name="deck.pptx", file_type="pptx"
                ).parse(data)
                self.assertIn("SLIDE TEXT OK", parsed.content)


class TestPptxMediaRealFallback(unittest.TestCase):
    @unittest.skipUnless(shutil.which("convert"), "ImageMagick is not installed")
    def test_svg_fallback_preserves_picture_order(self):
        # SVG is valid Office media, but Pillow cannot identify it. That makes
        # pinned MarkItDown 0.1.3 naturally enter the actual fallback path.
        data = deck([RED, BLUE])
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        original = "ppt/media/image2.png"
        replacement = "ppt/media/image2.svg"
        files[replacement] = (
            b'<svg xmlns="http://www.w3.org/2000/svg" width="16" height="16">'
            b'<rect width="16" height="16" fill="#1e50dc"/></svg>'
        )
        del files[original]
        for name, blob in list(files.items()):
            if name.endswith(".rels"):
                files[name] = blob.replace(b"image2.png", b"image2.svg")
        types = etree.fromstring(files["[Content_Types].xml"])
        etree.SubElement(
            types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
            Extension="svg",
            ContentType="image/svg+xml",
        )
        files["[Content_Types].xml"] = etree.tostring(types)
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(replacement, files.pop(replacement))
            for name, blob in files.items():
                archive.writestr(name, blob)

        document = MarkitdownParser(file_type="pptx").parse(output.getvalue())
        self.assertEqual(image_colors(document.content, document.images), [RED, BLUE])


if __name__ == "__main__":
    unittest.main()
