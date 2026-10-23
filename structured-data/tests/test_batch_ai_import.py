"""No model calls: serial budgets, absence evidence, cancellation and atomic guards."""
from copy import deepcopy
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
import pytest

from structured_data.batch_ai_import import import_all_with_ai, inspect_sheet_data, mark_no_data
from structured_data.imports import ImportManager
from structured_data.review_worker import run_review_task
from test_review_api import FakeAuthClient, Reply, client, make_workbook, poll


class Manager:
    def __init__(self, statuses=('needs_review',), published=()):
        self.doc = {'import_id': 'a' * 24, 'dataset_id': 'old', 'summary': {},
                    'sheets': [{'name': f'sheet-{i}', 'status': status,
                        'published_in_dataset': i in published,
                        'fingerprint': 'fingerprint',
                        'mapping': {'sheet': f'sheet-{i}', 'fingerprint': 'fingerprint', 'blocks': [{}]}}
                        for i, status in enumerate(statuses)]}

    def get(self, _):
        return deepcopy(self.doc)

    def publication_view(self, doc):
        return doc

    def sheet(self, _, name):
        return next(row for row in self.doc['sheets'] if row['name'] == name)

    def _result(self, doc):
        return {key: doc[key] for key in ('import_id', 'dataset_id', 'summary')}


def fake_task(manager, calls, *, preview_count=1, before_commit=None):
    def task(action, **kwargs):
        name = kwargs['mapping']['sheet']
        calls.append((action, name))
        if action == 'evidence':
            return {'status': 'uncertain'}
        if action == 'mark_skip':
            kwargs['publish_guard']()
            return {'sheet': name, 'message': 'No actual numeric values'}
        if action == 'preview':
            return {'series_count': 1, 'observation_count': 1, 'valid_observation_count': preview_count}
        if action == 'review':
            return {'status': 'confirmed', 'blocks': [], 'diagnostics': []}
        if action == 'confirm':
            if before_commit:
                before_commit(name)
            kwargs['publish_guard']()
            manager.doc['dataset_id'] = 'new-' + name
            manager.sheet('', name)['published_in_dataset'] = True
            return manager._result(manager.doc)
        raise AssertionError(action)
    return task


def fake_assist(counter, *, unresolved=False):
    def assist(manager, import_id, sheet, **kwargs):
        kwargs['budget'].consume()
        counter.append(sheet)
        return {'mapping': kwargs['mapping'], 'notice': 'source evidence missing' if unresolved else '',
                'review_state': {'status': 'needs_review' if unresolved else 'confirmed'},
                'model_requests': 1}
    return assist


def execute(manager, task, **kwargs):
    return import_all_with_ai(manager, 'a' * 24, task=task,
        check_write=kwargs.pop('check_write', lambda: None),
        cancelled=kwargs.pop('cancelled', lambda: False), **kwargs)


def test_only_actual_published_sheets_skipped_not_forged_import_state():
    manager = Manager(('needs_review', 'auto_ready'), published=(0,))
    manager.doc['sheets'][1]['import_state'] = 'imported'
    calls, models = [], []
    result = execute(manager, fake_task(manager, calls), assist=fake_assist(models))
    assert result['sheets_existing'] == ['sheet-0']
    assert result['sheets_imported'] == ['sheet-1']
    assert not models and all(name != 'sheet-0' for _, name in calls)


def test_batch_budget_is_global_and_serial_continues_non_ai_sheets():
    manager = Manager(('needs_review', 'needs_review', 'needs_review', 'auto_ready'))
    calls, models = [], []
    result = execute(manager, fake_task(manager, calls), assist=fake_assist(models),
                     total_max_requests=2, per_sheet_max_requests=1)
    assert result['requests_used'] == 2 and models == ['sheet-0', 'sheet-1']
    assert result['sheets_imported'] == ['sheet-0', 'sheet-1', 'sheet-3']
    assert result['sheets_pending'] == ['sheet-2']
    assert '总 AI 请求预算' in result['notice']
    assert [name for action, name in calls if action == 'confirm'] == result['sheets_imported']


def test_no_values_even_with_blank_records_skips_without_model_and_zero_is_valid():
    manager = Manager(('auto_ready',))
    calls = []
    result = execute(manager, fake_task(manager, calls, preview_count=0),
                     assist=lambda *args, **kwargs: pytest.fail('no-data sheet must not consume AI'))
    assert result['sheets_skipped'][0]['code'] == 'no_indicator_data'
    assert result['dataset_id'] == 'old' and not result['sheets_imported']
    assert not any(action == 'confirm' for action, _ in calls)


def test_unreliable_ai_mapping_never_published():
    manager = Manager()
    calls, models = [], []
    result = execute(manager, fake_task(manager, calls), assist=fake_assist(models, unresolved=True))
    assert result['sheets_pending'] == ['sheet-0'] and result['requests_used'] == 1
    assert result['dataset_id'] == 'old'
    assert not any(action == 'confirm' for action, _ in calls)


def test_model_claim_cannot_bypass_full_parser_validation():
    manager = Manager()
    calls, models = [], []
    original = fake_task(manager, calls)
    previews = 0

    def task(action, **kwargs):
        nonlocal previews
        if action == 'preview':
            previews += 1
            if previews == 2:
                raise ValueError('AI mapping is not valid')
        return original(action, **kwargs)

    result = execute(manager, task, assist=fake_assist(models))
    assert result['sheets_failed'][0]['message'] == 'AI mapping is not valid'
    assert result['dataset_id'] == 'old'


def test_cancel_keeps_completed_sheets_and_returns_all_partial_results():
    manager = Manager(('auto_ready', 'auto_ready', 'auto_ready'))
    calls = []
    cancelled = False

    def record(row):
        nonlocal cancelled
        if row['status'] == 'imported':
            cancelled = True

    result = execute(manager, fake_task(manager, calls), cancelled=lambda: cancelled, sheet_result=record)
    assert result['cancelled'] is True and result['stopped'] == 'cancelled'
    assert result['sheets_imported'] == ['sheet-0']
    assert result['sheets_pending'] == ['sheet-1', 'sheet-2']
    assert result['dataset_id'] == 'new-sheet-0'
    assert [row['status'] for row in result['details']] == ['imported', 'not_started', 'not_started']


def test_permission_rechecked_at_commit_and_stops_following_sheets():
    manager = Manager(('auto_ready', 'auto_ready', 'auto_ready'))
    calls = []
    revoked = False

    def check():
        if revoked:
            raise ValueError('access revoked')

    def before(name):
        nonlocal revoked
        if name == 'sheet-1':
            revoked = True

    result = execute(manager, fake_task(manager, calls, before_commit=before), check_write=check)
    assert result['stopped'] == 'permission_denied' and not result['cancelled']
    assert result['sheets_imported'] == ['sheet-0']
    assert result['dataset_id'] == 'new-sheet-0'
    assert ('preview', 'sheet-2') not in calls


def test_failed_replacement_keeps_existing_and_bulk_ignores_explicit_draft_for_it():
    manager = Manager(('confirmed',), published=(0,))
    result = execute(manager, lambda *_args, **_kwargs: pytest.fail('already imported'),
                     mappings=[{'sheet': 'sheet-0', 'blocks': ['untrusted']}])
    assert result['sheets_existing'] == ['sheet-0'] and result['dataset_id'] == 'old'


def content_workbook():
    book = Workbook()
    text = book.active
    text.title = 'ordinary-text'
    text.append(['标题', '说明'])
    text.append(['这是说明内容', '没有数值'])
    identifier = book.create_sheet('index-with-ids')
    identifier.append(['index', 'link'])
    identifier.append([1, 'Some destination'])
    formula = book.create_sheet('uncached-formula')
    formula['A1'] = '=1+1'
    dated = book.create_sheet('目录')
    dated.append(['date', 'actual measurement'])
    dated.append([datetime(2025, 1, 1), 0])
    dated.append([datetime(2025, 1, 2), 2])
    empty_values = book.create_sheet('日报')
    empty_values.append(['date', 'unavailable measurement'])
    empty_values.append([datetime(2025, 1, 1), None])
    empty_values.append([datetime(2025, 1, 2), None])
    raw = BytesIO()
    book.save(raw)
    return raw.getvalue()


def test_evidence_only_proves_absence_not_numbered_indexes_or_formulas(tmp_path):
    path = tmp_path / 'source.xlsx'
    path.write_bytes(content_workbook())
    manager = ImportManager(tmp_path / 'state')
    uploaded = manager.ingest(path, use_reviewed_mappings=False)
    identifier = uploaded['import_id']
    assert inspect_sheet_data(manager, identifier, 'ordinary-text')['status'] == 'no_indicator_data'
    assert inspect_sheet_data(manager, identifier, 'index-with-ids')['status'] == 'uncertain'
    assert inspect_sheet_data(manager, identifier, 'uncached-formula')['status'] == 'uncertain'
    assert inspect_sheet_data(manager, identifier, 'ordinary-text', max_cells=1)['status'] == 'uncertain'
    # Names do not exclude usable data. A real zero counts as available data.
    doc = manager.get(identifier)
    assert next(row for row in doc['sheets'] if row['name'] == '目录')['import_state'] == 'imported'
    unavailable = next(row for row in doc['sheets'] if row['name'] == '日报')
    assert unavailable['import_state'] == 'skipped'
    assert unavailable['automatic_skip']['code'] == 'no_indicator_data'
    ready = manager.sheet(identifier, '目录')['mapping']
    assert manager.preview(identifier, ready)['valid_observation_count'] == 2
    blank = manager.sheet(identifier, '日报')['mapping']
    assert manager.preview(identifier, blank)['valid_observation_count'] == 0
    # An explicitly checked single-sheet action can override automatic skip.
    manager.confirm(identifier, blank, save_template=False)
    assert next(row for row in manager.get(identifier)['sheets'] if row['name'] == '日报')['import_state'] == 'imported'
    assert 'automatic_skip' not in manager.sheet(identifier, '日报')


def test_bulk_api_requires_risk_and_strict_budgets_and_preserves_lock(client):
    uploaded = poll(client, client.post('/api/import', files={'file': ('review.xlsx', make_workbook())}))
    url = '/api/imports/' + uploaded['import_id'] + '/assist-confirm-all'
    assert client.post(url, json={}).status_code == 422
    assert client.post(url, json={'acknowledged_risk': False}).status_code == 400
    assert client.post(url, json={'acknowledged_risk': True, 'total_max_requests': True}).status_code == 422
    assert client.post(url, json={'acknowledged_risk': True, 'total_max_requests': 513}).status_code == 422
    assert client.post(url, json={'acknowledged_risk': True,
        'mappings': [{'sheet': 'parallel', 'note': 'x' * (260 * 1024)}]}).status_code == 413
    assert poll(client, client.post('/api/imports/' + uploaded['import_id'] + '/confirm-all', json={}))


def test_bulk_api_real_parser_imports_assisted_sheet_and_skips_empty(client, monkeypatch):
    uploaded = poll(client, client.post('/api/import', files={'file': ('review.xlsx', make_workbook())}))
    calls = []
    monkeypatch.setattr('structured_data.ai_review.assist_mapping', fake_assist(calls))
    result = poll(client, client.post('/api/imports/' + uploaded['import_id'] + '/assist-confirm-all',
        json={'acknowledged_risk': True, 'total_max_requests': 8}))
    # The real helper would see a fully valid parallel mapping and use no model;
    # this fake spends exactly one request to exercise the actual HTTP contract.
    assert result['sheets_imported'] == ['parallel']
    assert result['sheets_existing'] == ['simple']
    assert result['sheets_skipped'][0]['sheet'] == 'empty'
    assert result['requests_used'] == 1 and not result['sheets_failed']
    assert len(result['details']) == 3
    summary = client.get('/api/imports/' + uploaded['import_id'], params={'summary_only': True}).json()
    saved_empty = next(row for row in summary['sheets'] if row['name'] == 'empty')
    assert saved_empty['import_state'] == 'skipped'
    assert saved_empty['automatic_skip']['verified_source_sha256'] == summary['sha256']
    assert saved_empty['automatic_skip']['verified_fingerprint'] == saved_empty['fingerprint']
    # A second batch trusts only hash-bound skip evidence and does not call AI
    # or reinterpret the already published data.
    second = poll(client, client.post('/api/imports/' + uploaded['import_id'] + '/assist-confirm-all',
        json={'acknowledged_risk': True}))
    assert second['requests_used'] == 0 and second['sheets_imported'] == []
    assert second['sheets_skipped'][0]['sheet'] == 'empty'


def test_bulk_api_early_cancel_and_job_read_are_same_user_scoped(client, monkeypatch):
    uploaded = poll(client, client.post('/api/import', files={'file': ('review.xlsx', make_workbook())}))
    identifier = 'c' * 32
    original_get = FakeAuthClient.get

    async def another_user(self, url, headers):
        if headers.get('Authorization') == 'Bearer another':
            return Reply({'data': {'user': {'id': 'another', 'tenant_id': 10000},
                'memberships': [{'tenant_id': 10000, 'role': 'owner'}]}})
        return await original_get(self, url, headers)

    monkeypatch.setattr(FakeAuthClient, 'get', another_user)
    assert client.post('/api/jobs/' + identifier + '/cancel').status_code == 200
    response = client.post('/api/imports/' + uploaded['import_id'] + '/assist-confirm-all',
                           json={'acknowledged_risk': True, 'request_id': identifier})
    assert response.json() == {'id': identifier, 'status': 'cancelled'}
    assert client.get('/api/jobs/' + identifier).json()['status'] == 'cancelled'
    assert client.get('/api/jobs/' + identifier, headers={'Authorization': 'Bearer another'}).status_code == 404
    assert client.get('/api/imports/' + uploaded['import_id']).json()['dataset_id'] == uploaded['dataset_id']


def test_bulk_api_live_cancel_keeps_completed_new_sheet_and_original_snapshot(client, monkeypatch):
    import threading
    import time

    book = Workbook()
    for index in range(2):
        ws = book.active if index == 0 else book.create_sheet()
        ws.title = f'pending-{index}'
        ws.append(['date', 'price'])
        ws.append([None, 'source group'])
        ws.append([datetime(2025, 1, 1), index])
        ws.append([datetime(2025, 1, 2), index + 1])
    raw = BytesIO()
    book.save(raw)
    uploaded = poll(client, client.post('/api/import', files={'file': ('pending.xlsx', raw.getvalue())}))
    url = '/api/imports/' + uploaded['import_id']
    entered, release = threading.Event(), threading.Event()

    def assist(manager, import_id, sheet, **kwargs):
        kwargs['budget'].consume()
        if sheet == 'pending-1':
            entered.set()
            assert release.wait(30), 'Cancellation fixture did not release the model call'
        return {'mapping': kwargs['mapping'], 'notice': '',
                'review_state': {'status': 'confirmed'}, 'model_requests': 1}

    monkeypatch.setattr('structured_data.ai_review.assist_mapping', assist)
    response = client.post(url + '/assist-confirm-all', json={'acknowledged_risk': True})
    assert response.status_code == 200
    job_url = '/api/jobs/' + response.json()['id']
    try:
        # Several real parser subprocesses precede this event. Synchronize on
        # the event without treating shared-runner scheduling as a latency SLA.
        assert entered.wait(30), client.get(job_url).json()
        working = client.get(job_url).json()
        assert working['result']['sheets_imported'] == ['pending-0']
        assert working['progress']['total'] == 2
        assert client.post(job_url + '/cancel').status_code == 200
    finally:
        release.set()
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        job = client.get(job_url).json()
        if job['status'] != 'processing':
            break
        time.sleep(.01)
    assert job['status'] == 'cancelled'
    assert job['result']['sheets_imported'] == ['pending-0']
    assert job['result']['sheets_pending'] == ['pending-1']
    assert job['result']['requests_used'] == 2
    current = client.get(url).json()
    assert current['dataset_id'] != uploaded['dataset_id']
    coverage = {sheet['name']: sheet['published_in_dataset'] for sheet in current['sheets']}
    assert coverage == {'pending-0': True, 'pending-1': False}


def test_single_sheet_ai_does_not_auto_publish_all_blank_observations(client, monkeypatch):
    uploaded = poll(client, client.post('/api/import', files={'file': ('content.xlsx', content_workbook())}))
    url = '/api/imports/' + uploaded['import_id']
    calls = []
    monkeypatch.setattr('structured_data.ai_review.assist_mapping', fake_assist(calls))
    mapping = client.get(url + '/sheet', params={'name': '日报'}).json()['mapping']
    result = poll(client, client.post(url + '/assist-confirm', json={'sheet': '日报', 'mapping': mapping}))
    assert not result['imported'] and result['skipped']['code'] == 'no_indicator_data'
    assert result['dataset_id'] == uploaded['dataset_id']
    assert not next(row for row in client.get(url).json()['sheets'] if row['name'] == '日报')['published_in_dataset']


def test_mark_skip_requires_real_absence_and_final_permission_guard(tmp_path):
    path = tmp_path / 'content.xlsx'
    path.write_bytes(content_workbook())
    manager = ImportManager(tmp_path / 'state')
    uploaded = manager.ingest(path, use_reviewed_mappings=False)
    identifier = uploaded['import_id']
    mapping = manager.sheet(identifier, 'ordinary-text')['mapping']
    before = (manager.path(identifier) / 'review.json').read_bytes()

    def revoked():
        raise ValueError('permission revoked')

    with pytest.raises(ValueError, match='permission revoked'):
        mark_no_data(manager, identifier, mapping, before_write=revoked)
    assert (manager.path(identifier) / 'review.json').read_bytes() == before
    with pytest.raises(ValueError, match='可能有数据'):
        mark_no_data(manager, identifier, manager.sheet(identifier, 'index-with-ids')['mapping'], before_write=lambda: None)
    with pytest.raises(ValueError, match='指纹'):
        mark_no_data(manager, identifier, {**mapping, 'fingerprint': 'forged'}, before_write=lambda: None)
    marked = mark_no_data(manager, identifier, mapping, before_write=lambda: None)
    assert marked['verified_source_sha256'] == manager.get(identifier)['sha256']
    manager.source_path(manager.get(identifier)).write_bytes(b'changed original')
    with pytest.raises(ValueError, match='改变'):
        mark_no_data(manager, identifier, mapping, before_write=lambda: None)


def mixed_values_workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = 'mixed-blocks'
    sheet.append(['date', 'measurement', None, 'date', 'unavailable'])
    sheet.append([None, 'group A', None, None, 'group B'])
    sheet.append([datetime(2025, 1, 1), 0, None, datetime(2025, 1, 1), None])
    sheet.append([datetime(2025, 1, 2), 2, None, datetime(2025, 1, 2), None])
    raw = BytesIO()
    book.save(raw)
    return raw.getvalue()


def test_strict_review_keeps_empty_block_pending_and_manual_override_available(tmp_path):
    source = tmp_path / 'mixed.xlsx'
    source.write_bytes(mixed_values_workbook())
    manager = ImportManager(tmp_path / 'state')
    uploaded = manager.ingest(source, use_reviewed_mappings=False)
    identifier = uploaded['import_id']
    mapping = manager.sheet(identifier, 'mixed-blocks')['mapping']
    checked = manager.preview(identifier, mapping)
    assert [block['valid_observation_count'] for block in checked['blocks']] == [2, 0]
    limited = run_review_task(manager.root, 'preview', import_id=identifier, mapping=mapping)
    assert [block['valid_observation_count'] for block in limited['blocks']] == [2, 0]
    reviewed = manager.review_mapping(identifier, mapping)
    assert reviewed['status'] == 'needs_review'
    assert [block['status'] for block in reviewed['blocks']] == ['confirmed', 'needs_review']
    assert reviewed['blocks'][1]['diagnostics'][0]['field'] == 'columns'
    assert reviewed['blocks'][1]['diagnostics'][0]['block_index'] == 1
    # Manual import is explicit approval and may retain missing-data series.
    manager.confirm(identifier, mapping, save_template=False)
    published = manager.publication_view(manager.get(identifier))
    assert published['sheets'][0]['published_in_dataset'] is True


@pytest.mark.parametrize('batch', [False, True])
def test_ai_claim_of_confirmation_cannot_publish_a_mixed_empty_block(client, monkeypatch, batch):
    uploaded = poll(client, client.post('/api/import', files={'file': ('mixed.xlsx', mixed_values_workbook())}))
    url = '/api/imports/' + uploaded['import_id']
    mapping = client.get(url + '/sheet', params={'name': 'mixed-blocks'}).json()['mapping']
    calls = []
    monkeypatch.setattr('structured_data.ai_review.assist_mapping', fake_assist(calls))
    if batch:
        result = poll(client, client.post(url + '/assist-confirm-all', json={'acknowledged_risk': True}), timeout=30)
        assert result['sheets_pending'] == ['mixed-blocks'] and not result['sheets_imported']
        state = result['details'][0]['review_state']
    else:
        result = poll(client, client.post(url + '/assist-confirm', json={'sheet': 'mixed-blocks', 'mapping': mapping}))
        assert result['imported'] is False
        state = result['review_state']
    assert result['dataset_id'] == uploaded['dataset_id']
    assert [block['status'] for block in state['blocks']] == ['confirmed', 'needs_review']
    assert client.get(url).json()['sheets'][0]['published_in_dataset'] is False
