from copy import deepcopy
from datetime import datetime
import os
from pathlib import Path

import pytest
from openpyxl import Workbook

from structured_data.scanner import scan_workbook
from structured_data.generic_parser import parse_mapped_workbook
from structured_data.imports import ImportManager


def fixture(tmp_path, extra=None):
    path=tmp_path/'sales.xlsx'; book=Workbook(); sheet=book.active; sheet.title='销售日报'
    sheet.append([None,'日报 2023-07-10'])
    sheet.append([None,'日期','销售额'])
    sheet.append([None,datetime(2023,7,20),17145])
    sheet.append([None,datetime(2023,7,21),17165])
    sheet.append([None,'日变化',20])
    sheet.append([None,'趋势'])
    if extra:sheet.append(extra)
    book.save(path)
    return path


def test_summary_preview_confirm_query(tmp_path):
    path=fixture(tmp_path); manager=ImportManager(tmp_path/'data'); imported=manager.ingest(path)
    sheet=manager.sheet(imported['import_id'],'销售日报'); mapping=sheet['mapping']
    assert sheet['status']=='needs_review' and len(mapping['blocks'])==2
    assert mapping['blocks'][0]['data_end_row']==4
    preview=manager.preview(imported['import_id'],mapping)
    assert preview['observation_count']==3
    result=manager.confirm(imported['import_id'],mapping)
    series=manager.store.list_series(result['dataset_id'])['series']
    dated=next(s for s in series if s['date_mode']=='date')
    change=next(s for s in series if s['date_mode']=='none')
    price=manager.store.query_series(result['dataset_id'],dated['series_id'],'2023-07-21')['rows'][0]
    delta=manager.store.query_series(result['dataset_id'],change['series_id'])['rows'][0]
    assert price['value']==17165 and price['source_cell']=='C4'
    assert delta['value']==20 and delta['date'] is None and delta['category_cell']=='B5'
    assert change['dimensions']['统计口径']=='日变化'


@pytest.mark.parametrize('extra',[
    [None,datetime(2023,7,22),17170], [None,'其他',10], [None,None,'=1+2'],
])
def test_unknown_or_resumed_tail_is_not_silently_removed(tmp_path,extra):
    path=fixture(tmp_path,extra); sheet=scan_workbook(path)['sheets'][0]
    assert not sheet['candidate_revision']
    with pytest.raises(ValueError,match='日期无效'):
        parse_mapped_workbook(path,[sheet['mapping']])


def test_unknown_label_is_not_treated_as_summary(tmp_path):
    path=fixture(tmp_path)
    from openpyxl import load_workbook
    book=load_workbook(path);book.active['B5']='2023-99-99';book.save(path)
    sheet=scan_workbook(path)['sheets'][0]
    assert not sheet['candidate_revision']
    with pytest.raises(ValueError,match='日期无效'):parse_mapped_workbook(path,[sheet['mapping']])


@pytest.mark.skipif(os.getenv('COTTON_ACCEPTANCE')!='1',reason='Original workbook opt-in')
def test_real_daily_report_preserves_price_change_and_formula(tmp_path):
    path=Path('/Users/mac/Documents/中哲物产/棉花测试集/棉花/棉花数据库(2).xlsx')
    scan=scan_workbook(path);sheet=next(s for s in scan['sheets'] if s['name']=='日报')
    assert sheet['candidate_revision']=='row-summary-sections-v1'
    result=parse_mapped_workbook(path,[sheet['mapping']],scan=scan)
    assert len(result['series'])==34 and len(result['observations'])==119
    cells={o['source_cell']:o for o in result['observations']}
    assert cells['C11']['value']==17165 and cells['C11']['date']=='2023-07-21'
    assert cells['C12']['value']==20 and cells['C12']['date'] is None
    assert cells['C12']['category_cell']=='B12' and cells['C12']['category']=='日变化'
    assert cells['C12']['formula']=='=C11-C10'
    assert 'C13' not in cells
    old=deepcopy(sheet['mapping']);old['blocks']=old['blocks'][:1];old['blocks'][0]['data_end_row']=None
    with pytest.raises(ValueError,match='日报!B12 日期无效'):
        parse_mapped_workbook(path,[old],scan=scan)
