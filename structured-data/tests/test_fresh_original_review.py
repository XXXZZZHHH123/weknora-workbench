"""Fresh review must not consult packaged or previously imported case mappings."""
from copy import deepcopy
from datetime import datetime
import hashlib
import json

from openpyxl import Workbook
import pytest

from structured_data.imports import ImportManager
from structured_data.scanner import scan_workbook
from structured_data import review_worker


def original(path, changed=False):
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Measurements"
    sheet.append(["date", "quantity"])
    sheet.append([datetime(2024, 1, 1), 4 if changed else 3])
    sheet.append([datetime(2024, 1, 2), 5])
    extra = workbook.create_sheet("Layered")
    extra.append(["date", "quantity"]); extra.append([None, "group A"])
    extra.append([datetime(2024, 1, 1), 7]); extra.append([datetime(2024, 1, 2), 9])
    workbook.save(path)
    return path


def prohibit_profiles(monkeypatch):
    def forbidden(*_args, **_kwargs):
        raise AssertionError("A fresh original must not consult packaged reviewed mappings")
    monkeypatch.setattr("structured_data.profiles.reviewed.reviewed_mapping", forbidden)
    monkeypatch.setattr("structured_data.regions.reviewed_regions", forbidden)
    monkeypatch.setattr("structured_data.profiles.reviewed.parse_workbook", forbidden)


def test_fresh_scanner_and_initial_publish_do_not_consult_packaged_profiles(tmp_path, monkeypatch):
    path = original(tmp_path / "source.xlsx")
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    prohibit_profiles(monkeypatch)
    scan = scan_workbook(path, use_reviewed_mappings=False)
    assert scan["use_reviewed_mappings"] is False and scan["sheet_count"] == 2
    assert scan["sheets"][0]["mapping"]["blocks"][0]["columns"][0]["metric"] == "quantity"
    manager = ImportManager(tmp_path / "data")
    result = manager.ingest(path, use_reviewed_mappings=False)
    doc = manager.get(result["import_id"])
    assert doc["use_reviewed_mappings"] is False
    assert not any(sheet["status"] == "trusted_existing" for sheet in doc["sheets"])
    assert manager.store.manifest(result["dataset_id"])["use_reviewed_mappings"] is False
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_fresh_import_and_dataset_never_reuse_default_ids_or_local_default_templates(tmp_path, monkeypatch):
    path = original(tmp_path / "source.xlsx")
    manager = ImportManager(tmp_path / "data")
    default = manager.ingest(path)
    manager.confirm(default["import_id"], manager.sheet(default["import_id"], "Layered")["mapping"])
    prohibit_profiles(monkeypatch)
    fresh = manager.ingest(path, use_reviewed_mappings=False)
    assert fresh["import_id"] != default["import_id"]
    assert fresh["dataset_id"] != default["dataset_id"]
    assert manager.sheet(fresh["import_id"], "Layered")["status"] == "needs_review"
    assert manager.store.manifest(fresh["dataset_id"])["import_id"] == fresh["import_id"]


def test_fresh_confirm_restart_and_current_scope_template_reuse_stay_fresh(tmp_path, monkeypatch):
    path = original(tmp_path / "source.xlsx")
    manager = ImportManager(tmp_path / "data")
    prohibit_profiles(monkeypatch)
    fresh = manager.ingest(path, use_reviewed_mappings=False)
    mapping = deepcopy(manager.sheet(fresh["import_id"], "Layered")["mapping"])
    assert manager.preview(fresh["import_id"], mapping)["observation_count"] == 2
    confirmed = manager.confirm(fresh["import_id"], mapping, save_template=True)
    restarted = ImportManager(tmp_path / "data")
    assert restarted.sheet(fresh["import_id"], "Layered")["status"] == "confirmed"
    assert restarted.ingest(path, use_reviewed_mappings=False)["dataset_id"] == confirmed["dataset_id"]
    another = original(tmp_path / "update.xlsx", changed=True)
    reused = restarted.ingest(another, use_reviewed_mappings=False)
    assert restarted.sheet(reused["import_id"], "Layered")["status"] == "template_ready"
    assert restarted.get(reused["import_id"])["use_reviewed_mappings"] is False
    assert all(restarted.store.manifest(identifier)["use_reviewed_mappings"] is False
               for identifier in (fresh["dataset_id"], confirmed["dataset_id"], reused["dataset_id"]))


def test_native_source_lookup_distinguishes_default_and_fresh_review_records(tmp_path):
    path = original(tmp_path / "source.xlsx")
    manager = ImportManager(tmp_path / "data")
    source = {"source_knowledge_id": "new-source", "file_size": path.stat().st_size,
              "file_hash": hashlib.md5(path.read_bytes()).hexdigest(), "knowledge_base_id": "kb-new"}
    old = manager.ingest(path)
    manager.bind_native_source(old["import_id"], source)
    assert manager.find_native_source(source, use_reviewed_mappings=False) is None
    fresh = manager.ingest(path, use_reviewed_mappings=False)
    manager.bind_native_source(fresh["import_id"], source)
    assert manager.find_native_source(source, use_reviewed_mappings=False)["import_id"] == fresh["import_id"]
    assert manager.find_native_source(source)["import_id"] == old["import_id"]


def test_fresh_catalog_reports_import_identity_and_empty_sheet_without_a_repair_task(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    path = original(tmp_path / 'source.xlsx')
    workbook = load_workbook(path)
    workbook.create_sheet('Intentionally empty')
    workbook.save(path)
    prohibit_profiles(monkeypatch)
    manager = ImportManager(tmp_path / 'data')
    result = manager.ingest(path, use_reviewed_mappings=False)
    catalog = manager.store.list_datasets()
    assert len(catalog) == 1
    item = catalog[0]
    assert item['dataset_id'] == result['dataset_id'] and item['import_id'] == result['import_id']
    assert item['use_reviewed_mappings'] is False
    assert item['empty_sheets'] == ['Intentionally empty']
    assert 'Intentionally empty' not in item['pending_sheets']
    assert 'Layered' in item['pending_sheets']


def test_catalog_keeps_pending_status_and_reasons_bound_to_each_dataset(tmp_path, monkeypatch):
    path = original(tmp_path / 'source.xlsx')
    prohibit_profiles(monkeypatch)
    manager = ImportManager(tmp_path / 'data')
    first = manager.ingest(path, use_reviewed_mappings=False)
    original(path, changed=True)
    second = manager.ingest(path, use_reviewed_mappings=False)
    assert first['dataset_id'] != second['dataset_id']

    expected = {
        first['dataset_id']: {
            'Measurements': ('needs_review', ['first snapshot review reason']),
            'Layered': ('unsupported', ['first snapshot unsupported reason']),
        },
        second['dataset_id']: {
            'Measurements': ('unsupported', ['second snapshot unsupported reason']),
            'Layered': ('needs_review', ['second snapshot review reason']),
        },
    }
    for dataset_id, states in expected.items():
        manifest_path = manager.store._path(dataset_id) / 'manifest.json'
        manifest = json.loads(manifest_path.read_text())
        for sheet in manifest['coverage']:
            status, reasons = states[sheet['name']]
            sheet['status'], sheet['reasons'] = status, reasons
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2))

    listed = {item['dataset_id']: item for item in manager.store.list_datasets()}
    for dataset_id, states in expected.items():
        item = listed[dataset_id]
        expected_states = [
            {'name': name, 'status': status, 'reasons': reasons}
            for name, (status, reasons) in states.items()
        ]
        assert item['pending_sheet_states'] == expected_states
        assert item['pending_sheets'] == [state['name'] for state in expected_states]
        assert all(isinstance(name, str) for name in item['pending_sheets'])
        assert 'pending_sheet_states' in item['coverage_notice']
        assert 'unsupported不能表述为完成映射后即可使用' in item['coverage_notice']

    assert listed[first['dataset_id']]['pending_sheet_states'][0]['status'] == 'needs_review'
    assert listed[second['dataset_id']]['pending_sheet_states'][0]['status'] == 'unsupported'


def test_catalog_tolerates_legacy_pending_coverage_without_reasons(tmp_path, monkeypatch):
    path = original(tmp_path / 'source.xlsx')
    prohibit_profiles(monkeypatch)
    manager = ImportManager(tmp_path / 'data')
    result = manager.ingest(path, use_reviewed_mappings=False)
    manifest_path = manager.store._path(result['dataset_id']) / 'manifest.json'
    manifest = json.loads(manifest_path.read_text())
    pending = next(sheet for sheet in manifest['coverage'] if sheet['name'] == 'Layered')
    pending.pop('reasons', None)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2))

    item = next(item for item in manager.store.list_datasets() if item['dataset_id'] == result['dataset_id'])
    assert {'name': 'Layered', 'status': 'needs_review', 'reasons': []} in item['pending_sheet_states']
    assert 'Layered' in item['pending_sheets']


def test_signed_tool_catalog_keeps_fresh_identity_for_current_and_historical_snapshots(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    from structured_data.server import create_app
    from test_scoped_service import signed
    path = original(tmp_path / 'source.xlsx')
    root = tmp_path / 'data'
    manager = ImportManager(root)
    prohibit_profiles(monkeypatch)
    first = manager.ingest(path, use_reviewed_mappings=False)
    key = b'fresh-tool-contract-test-key-1234567890'
    key_path = tmp_path / 'read-key'
    key_path.write_bytes(key)
    monkeypatch.setenv('STRUCTURED_SERVICE_TOKEN_FILE', str(key_path))
    monkeypatch.setenv('STRUCTURED_TENANT_ID', '10000')
    with TestClient(create_app(root)) as client:
        body = {'operation': 'list_datasets', 'arguments': {}}
        raw, headers = signed(key, body=body)
        response = client.post('/api/tools/read', content=raw, headers=headers)
        assert response.status_code == 200, response.text
        catalog = response.json()['datasets']
        assert {item['dataset_id'] for item in catalog} == {first['dataset_id']}
        first_item = catalog[0]
        assert first_item['pending_sheets'] == [state['name'] for state in first_item['pending_sheet_states']]
        coverage = manager.store.manifest(first['dataset_id'])['coverage']
        expected_states = [
            {'name': sheet['name'], 'status': sheet['status'], 'reasons': sheet['reasons']}
            for sheet in coverage if sheet['status'] in {'needs_review', 'unsupported'}
        ]
        assert first_item['pending_sheet_states'] == expected_states
        assert first_item['coverage_notice'] == manager.store.list_datasets()[0]['coverage_notice']

        final = manager.confirm(first['import_id'], manager.sheet(first['import_id'], 'Layered')['mapping'])
        assert first['dataset_id'] != final['dataset_id']
        raw, headers = signed(key, body=body)
        response = client.post('/api/tools/read', content=raw, headers=headers)
        assert response.status_code == 200, response.text
        catalog = response.json()['datasets']
        assert {item['dataset_id'] for item in catalog} == {first['dataset_id'], final['dataset_id']}
        assert all(item['use_reviewed_mappings'] is False for item in catalog)
        assert all(item['import_id'] == first['import_id'] for item in catalog)
        assert all(item['sha256'] == hashlib.sha256(path.read_bytes()).hexdigest() for item in catalog)
        latest = next(item for item in catalog if item['dataset_id'] == final['dataset_id'])
        assert latest['parent_dataset_id'] == first['dataset_id']
        historical = next(item for item in catalog if item['dataset_id'] == first['dataset_id'])
        assert 'pending_sheet_states' not in historical


def test_spawn_worker_carries_fresh_mode_to_saved_review_and_confirmation(tmp_path):
    path = original(tmp_path / "source.xlsx")
    result = review_worker.run_review_task(tmp_path / "data", "import", path=path, use_reviewed_mappings=False)
    manager = ImportManager(tmp_path / "data")
    assert manager.get(result["import_id"])["use_reviewed_mappings"] is False
    mapping = manager.sheet(result["import_id"], "Layered")["mapping"]
    saved = review_worker.run_review_task(tmp_path / "data", "confirm", import_id=result["import_id"], mapping=mapping)
    assert manager.get(saved["import_id"])["use_reviewed_mappings"] is False
    assert manager.store.manifest(saved["dataset_id"])["use_reviewed_mappings"] is False
