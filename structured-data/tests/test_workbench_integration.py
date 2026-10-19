from datetime import datetime
from pathlib import Path
import json
import httpx
from fastapi.testclient import TestClient
from openpyxl import Workbook
from structured_data.server import create_app
from structured_data.scanner import scan_workbook
from structured_data.generic_parser import parse_mapped_workbook
from structured_data.model_settings import read_model
import pytest
from test_native_charts import chart_store


def test_merged_headers_keep_category_measure_and_evidence(tmp_path):
    w=Workbook();s=w.active;s.title='通用零售'
    s.append(['日期','销售额(亿元)',None,'销量(万件)',None,'均价(元/件)',None,'合计'])
    s.append([None,'鞋','衣','鞋','衣','鞋','衣','亿元'])
    for cell in ('B1:C1','D1:E1','F1:G1'):s.merge_cells(cell)
    s.append([datetime(2023,1,1),10,20,30,40,50,60,30])
    s.append([datetime(2023,2,1),11,21,31,41,51,61,32])
    p=tmp_path/'retail.xlsx';w.save(p)
    scan=scan_workbook(p);m=scan['sheets'][0]['mapping'];cols=m['blocks'][0]['columns']
    assert [(c['metric'],c['unit']) for c in cols[:6]]==[(a+' / '+b,u) for a,u in [('销售额(亿元)','亿元'),('销量(万件)','万件'),('均价(元/件)','元/件')] for b in ('鞋','衣')]
    assert cols[1]['header_cells']==['B1','C2']
    assert cols[-1]['metric']=='合计' and cols[-1]['unit']=='亿元'
    result=parse_mapped_workbook(p,[m],scan=scan)
    assert len(result['series'])==7
    assert next(o for o in result['observations'] if o['source_cell']=='C3')['value']==20


@pytest.mark.parametrize('labels',[['1901','1902'],['2021','2022'],['2017年','2018年']])
def test_literal_chart_labels_do_not_override_verified_source_dates(chart_store,monkeypatch,labels):
    from structured_data import native_charts
    store,dataset,_=chart_store
    actual_extract=native_charts.extract
    def renamed(path):
        charts=actual_extract(path);charts[0]['curves']=charts[0]['curves'][:2]
        for curve,label in zip(charts[0]['curves'],labels):curve['label']=label
        return charts
    monkeypatch.setattr(native_charts,'extract',renamed)
    (store._path(dataset) / f'{native_charts.VERSION}.json').unlink()
    curves=store.chart_index(dataset)['charts'][0]['curves']
    assert [c['label'] for c in curves]==labels
    assert all(c['status']=='linked' for c in curves)
    assert curves[0]['first_date']=='2023-09-07' and curves[1]['first_date']=='2024-01-01'
    assert store.query_series(dataset,curves[0]['series_id'],operation='sum')['aggregate']==30
    assert store.query_series(dataset,curves[1]['series_id'],operation='sum')['aggregate']==99


def test_shared_login_checks_identity_and_model_settings_redact_key(tmp_path,monkeypatch):
    class AuthClient:
        def __init__(self,**kw):pass
        async def __aenter__(self):return self
        async def __aexit__(self,*args):pass
        async def get(self,url,headers):
            token=headers['Authorization']
            if token=='Bearer expired':return httpx.Response(401,json={})
            tenant=10001 if token=='Bearer other' else 10000
            return httpx.Response(200,json={'data':{'user':{'id':'user-a','tenant_id':tenant},
                'memberships':[{'tenant_id':tenant,'role':'owner'}]}})
    monkeypatch.setattr('httpx.AsyncClient',AuthClient)
    monkeypatch.setenv('STRUCTURED_TENANT_ID','10000')
    monkeypatch.setenv('STRUCTURED_MODEL_CONFIG',str(tmp_path/'absent'))
    with TestClient(create_app(tmp_path)) as c:
        assert c.get('/api/datasets').status_code==401
        assert c.get('/api/datasets',headers={'Authorization':'Bearer expired'}).status_code==401
        assert c.get('/api/datasets',headers={'Authorization':'Bearer other'}).json()=={'datasets':[]}
        good={'Authorization':'Bearer valid','Origin':'http://testserver','X-Tenant-ID':'10000'}
        assert c.get('/api/datasets',headers=good).status_code==200
        assert c.get('/api/datasets',headers=good|{'X-Tenant-ID':'10001'}).status_code==403
        payload={'model':'custom-chat','base_url':'https://example.com/v1','api_key':'test-private-key'}
        assert c.post('/api/model-settings',json=payload,headers=good|{'Origin':'https://evil.invalid'}).status_code==403
        saved=c.post('/api/model-settings',json=payload,headers=good)
        assert saved.status_code==200 and 'test-private-key' not in saved.text
        assert 'api_key' not in c.get('/api/model-settings',headers=good).json()
        assert read_model(tmp_path)['api_key']=='test-private-key'
        assert (tmp_path/'model-settings.json').stat().st_mode & 0o777==0o600
        payload.update(api_key='',model='changed-model')
        assert c.post('/api/model-settings',json=payload,headers=good).status_code==200
        assert read_model(tmp_path)['model']=='changed-model' and read_model(tmp_path)['api_key']=='test-private-key'


def test_review_error_locations_survive_http_and_background_job(tmp_path, monkeypatch):
    from structured_data.imports import ImportManager
    import time
    class AuthClient:
        def __init__(self, **kwargs): pass
        async def __aenter__(self): return self
        async def __aexit__(self, *args): pass
        async def get(self, url, headers):
            return httpx.Response(200, json={'data': {'user': {'id': 'review-user', 'tenant_id': 10000},
                'memberships': [{'tenant_id': 10000, 'role': 'owner'}]}})
    monkeypatch.setattr('httpx.AsyncClient', AuthClient)
    monkeypatch.setenv('STRUCTURED_TENANT_ID', '10000')
    source=tmp_path/'review-fixture.xlsx'
    wb=Workbook();ws=wb.active;ws.title='录入表'
    ws.append(['日期','数量']);ws.append([datetime(2023,1,1),1]);ws.append([datetime(2023,1,2),2]);wb.save(source)
    manager=ImportManager(tmp_path)
    imported=manager.ingest(source)
    mapping=manager.sheet(imported['import_id'],'录入表')['mapping']
    mapping['blocks'][0]['header_rows']=[]
    headers={'Authorization':'Bearer fixture','Origin':'http://testserver','X-Tenant-ID':'10000'}
    with TestClient(create_app(tmp_path)) as client:
        response=client.post(f'/api/imports/{imported["import_id"]}/sample-preview',json={'mapping':mapping},headers=headers)
        assert response.status_code==400
        assert response.json()['review_target']['field']=='header_rows'
        assert response.json()['review_target']['sheet']=='录入表'
        job=client.post(f'/api/imports/{imported["import_id"]}/preview?background=true',json={'mapping':mapping},headers=headers)
        assert job.status_code==200
        deadline=time.monotonic()+10
        while time.monotonic()<deadline:
            result=client.get('/api/jobs/'+job.json()['id'],headers=headers).json()
            if result['status']=='failed': break
            time.sleep(0.05)
        assert result['status']=='failed'
        assert result['review_target']==response.json()['review_target']
    assert manager.get(imported['import_id'])['dataset_id']==imported['dataset_id']
