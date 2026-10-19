from datetime import datetime
from hashlib import sha256
import json

from openpyxl import Workbook
import pytest

from structured_data.store import DataStore
from structured_data.source_context import formula_references, protect_duplicate_bindings, verify_selection
from structured_data.field_resolver import FieldResolver


@pytest.fixture
def duplicate_source(tmp_path):
    book=Workbook();sheet=book.active;sheet.title='原料'
    sheet.append(['日期','3128级棉花','20天原料','双29棉花','20天原料','CC Index 3128B','20天原料','双29棉花','20天原料'])
    sheet.append([None,*['元/吨']*8]);sheet.merge_cells('B1:B1')
    sheet.append([datetime(2023,1,1),100,'=B3',130,'=D3',200,'=F3',230,'=H3'])
    path=tmp_path/'input.xlsx';book.save(path)
    series=[];observations=[]
    for col,metric,value,formula in [('B','3128级棉花',100,None),('C','20天原料',100,'=B3'),('D','双29棉花',130,'=B3+30'),('E','20天原料',130,'=D3'),('F','CC Index 3128B',200,None),('G','20天原料',200,'=F3'),('H','双29棉花',230,'=F3+30'),('I','20天原料',230,'=H3')]:
        series.append(dict(series_id=col,sheet='原料',metric=metric,unit='元/吨',dimensions={'原表列':col},header_cells=[col+'1',col+'2'],date_mode='date'))
        observations.append(dict(series_id=col,date='2023-01-01',value=value,formula=formula,source_cell=col+'3',date_cell='A3',raw_value=value,value_status='valid'))
    store=DataStore(tmp_path/'data')
    published=store.publish(dict(source_file=str(path),sha256=sha256(path.read_bytes()).hexdigest(),rule_version='fixture',series=series,observations=observations,sheets=[],excluded_regions=[],issues=[]))
    return store,published['dataset_id']


def test_context_keeps_actual_headers_recursive_formula_and_original(duplicate_source):
    store,dataset=duplicate_source;before=store.manifest(dataset);catalog=store.catalog(dataset)
    context=next(s['source_context'] for s in catalog if s['series_id']=='E')
    parent=context['formula_source']['references'][0]['fields'][0]
    assert parent['metric']=='双29棉花'
    assert parent['source_context']['formula_source']['references'][0]['fields'][0]['metric']=='3128级棉花'
    assert context['headers'][0]=={'cell':'E1','value':'20天原料'}
    assert store.catalog(dataset)==catalog
    assert store.manifest(dataset)==before
    assert len(store.list_datasets())==1


def test_formula_reference_tokens_do_not_follow_external_or_text():
    refs=formula_references('=SUM($B$2:C9)+\'另一张表\'!D3+"E4"+\'[other.xlsx]Sheet\'!F1','本表')
    assert [(r['sheet'],r['columns']) for r in refs]==[('本表',['B','C']),('另一张表',['D'])]
    assert formula_references('=INDIRECT("B3")','本表')==[]


@pytest.mark.parametrize('question',['20天原料','不选第一套，用另一套','这个字段的滚动值','20天原料，原表C列'])
def test_singleton_requires_complete_model_binding_for_aliases_and_negation(duplicate_source,question):
    store,dataset=duplicate_source;catalog=store.catalog(dataset)
    proposal={'decision':'resolved','fields':[{'role':'v1','candidate_ids':['C']}]}
    with pytest.raises(ValueError,match='source_binding'):
        protect_duplicate_bindings(proposal,catalog,question)
    binding={'considered_ids':['C','E','G','I'],'selected_ids':['G'],'reason':'模型根据完整问题及真实来源解释，选择第二来源的对应字段'}
    selected={'decision':'resolved','fields':[{'role':'v1','candidate_ids':['G'],'source_binding':binding}]}
    assert protect_duplicate_bindings(selected,catalog,question)['fields'][0]['candidate_ids']==['G']
    assert verify_selection(store,dataset,['G'],question,lambda p:{'decision':'resolved','bindings':[binding]}) is None
    mismatch=verify_selection(store,dataset,['C'],question,lambda p:{'decision':'resolved','bindings':[binding]})
    assert mismatch['status']=='plan_mismatch' and 'rows' not in mismatch


def test_resolver_repairs_source_plan_in_model_instead_of_fixed_clarification(duplicate_source):
    store,dataset=duplicate_source;calls=[]
    def model(payload):
        calls.append(payload)
        if len(calls)==1:
            return {'intent':'lookup','decision':'resolved','fields':[{'role':'v1','label':'滚动输入','candidate_ids':['C']}],'query':{'operation':'rows'}}
        assert 'source_binding' in payload['repair_instruction']
        assert {s['series_id'] for group in payload['duplicate_sources'] for s in group}=={'C','E','G','I','D','H'}
        return {'intent':'lookup','decision':'clarify','question':'模型发现不同输入依赖，请说明使用的那一组来源。',
            'fields':[{'role':'v1','label':'滚动输入','candidate_ids':['C','E','G','I']}],'query':{'operation':'rows'}}
    resolver=FieldResolver(store,model)
    plan=resolver.resolve(dataset,'这个滚动输入是什么？不要选第一组',date='2023-01-01')
    assert len(calls)==2 and plan['status']=='needs_clarification'
    assert {c['series_id'] for c in plan['fields'][0]['candidates']}=={'C','E','G','I'}
    assert plan['question'].startswith('模型发现')
    with pytest.raises(ValueError,match='尚未确认'):resolver.execute(plan['resolution_id'])
    confirmed=resolver.confirm(plan['resolution_id'],{'v1':'G'})
    assert resolver.execute(confirmed['resolution_id'])['sources']['v1']['rows'][0]['value']==200


def test_direct_source_review_does_not_reuse_question_only_cache(duplicate_source):
    store,dataset=duplicate_source;calls=[]
    def model(p):
        calls.append(p)
        chosen='C' if len(calls)==1 else 'G'
        return {'decision':'resolved','bindings':[{'considered_ids':['C','E','G','I'],'selected_ids':[chosen],'reason':'本次模型根据本轮实际上下文选择来源'}]}
    assert verify_selection(store,dataset,['C'],'还是这个字段',model) is None
    assert verify_selection(store,dataset,['C'],'还是这个字段',model) is None
    assert len(calls)==1  # Exact successful binding is reusable, not question-only.
    assert verify_selection(store,dataset,['C'],'改用另一来源，还是这个字段',model)['status']=='plan_mismatch'
    assert len(calls)==2
    assert verify_selection(store,dataset,['C'],None,model)['status']=='error'
    def unavailable(p):raise ValueError('model timeout')
    failed=verify_selection(store,dataset,['C'],'滚动字段',unavailable)
    assert failed['status']=='error' and failed['retryable'] and 'question' not in failed


@pytest.mark.parametrize('binding',[
    {'considered_ids':['C'],'selected_ids':['C'],'reason':'遗漏实际来源'},
    {'considered_ids':['C','E','G','I'],'selected_ids':['imaginary'],'reason':'虚构ID'},
    {'considered_ids':['C','E','G','I'],'selected_ids':['C'],'reason':''},
])
def test_bad_source_plans_fail_closed_without_querying(duplicate_source,binding):
    store,dataset=duplicate_source
    result=verify_selection(store,dataset,['C'],'滚动字段',lambda p:{'decision':'resolved','bindings':[binding]})
    assert result['status']=='error' and 'rows' not in result


def test_model_clarification_can_leave_selection_empty_but_must_cover_all_real_sources(duplicate_source):
    store,dataset=duplicate_source;calls=[]
    review={'decision':'clarify','question':'模型发现多个实际依赖，请明确希望使用的来源',
            'bindings':[{'considered_ids':['C','E','G','I'],'selected_ids':[],
                         'reason':'完整问题未指定不同来源，尚未选择任意一个字段'}]}
    def model(payload):calls.append(payload);return review
    result=verify_selection(store,dataset,['C'],'请解释这个滚动字段',model)
    assert result['status']=='needs_clarification' and result['question']==review['question']
    assert result['bindings'][0]['selected_ids']==[] and len(calls)==1
    assert 'rows' not in result and 'aggregate' not in result
    invalid=dict(review,decision='resolved')
    assert verify_selection(store,dataset,['C'],'实际需要明确的选择',lambda p:invalid)['status']=='error'
    missing_question=dict(review);missing_question.pop('question')
    assert verify_selection(store,dataset,['C'],'另一次业务问题',lambda p:missing_question)['status']=='error'


def test_signed_query_and_sql_use_real_model_source_plan(duplicate_source, monkeypatch):
    from fastapi.testclient import TestClient
    from structured_data.server import create_app
    from structured_data.field_resolver import FieldResolver
    from test_scoped_service import signed
    store,dataset=duplicate_source;calls=[]
    def model(payload):
        calls.append(payload)
        assert payload['_source_binding'] and payload['question']=='完整问题和实际上下文'
        assert payload['source_groups'][0][0]['source_context']['formula_source']
        return {'decision':'clarify','question':'模型需要确认实际输入依赖',
                'bindings':[{'considered_ids':['C','E','G','I'],'selected_ids':['C','E','G','I'],'reason':'真实来源依赖不同，完整问题尚未区分'}]}
    monkeypatch.setattr(FieldResolver,'__init__',lambda self,data: self.__dict__.update(store=data,model_fn=model))
    key=b'actual-source-selection-guard-1234567890'
    path=store.root/'read-key';path.write_bytes(key)
    monkeypatch.setenv('STRUCTURED_SERVICE_TOKEN_FILE',str(path));monkeypatch.setenv('STRUCTURED_TENANT_ID','10000')
    with TestClient(create_app(store.root)) as client:
        for operation,args in [('query_series',{'series_id':'C'}),('analyze_table',{'series_ids':['C'],'sql':'SELECT sum(value) FROM data'}),('compare_series',{'left_series_id':'C','right_series_id':'G','date':'2023-01-01'}),('convert_price',{'price_series_id':'C','fx_series_id':'B','date':'2023-01-01'})]:
            raw,headers=signed(key,body={'operation':operation,'arguments':{'dataset_id':dataset,'source_question':'完整问题和实际上下文',**args}})
            response=client.post('/api/tools/read',content=raw,headers=headers)
            assert response.status_code==200,response.text
            assert response.json()['status']=='needs_clarification'
            assert response.json()['question']=='模型需要确认实际输入依赖'
            assert 'rows' not in response.json() and 'aggregate' not in response.json()
    assert len(calls)==4


def test_merged_header_anchor_and_scope_cache(duplicate_source,tmp_path):
    from structured_data.source_context import _headers
    import zipfile,time
    book=Workbook();w=book.active;w.title='分组'
    w.append([None,'来源组',None]);w.merge_cells('B1:C1');w.append(['日期','原料','20天原料'])
    path=tmp_path/'merged.xlsx';book.save(path)
    with zipfile.ZipFile(path) as archive:
        result=_headers(archive,{'分组'},{'分组':{1,2}},time.monotonic()+10)
    assert result['分组']['merged'][0]['anchor']=='B1'
    assert result['分组']['merged'][0]['last_column']==3
    assert result['分组']['cells']['B1']=='来源组'
    store,dataset=duplicate_source
    assert len(store.catalog(dataset,sheet='原料'))==len(store.catalog(dataset))


def test_initial_catalog_recalls_all_literal_duplicate_headers(duplicate_source,monkeypatch):
    from fastapi.testclient import TestClient
    from structured_data.server import create_app
    from test_scoped_service import signed
    store,dataset=duplicate_source;key=b'literal-field-discovery-key-1234567890'
    path=store.root/'read-key';path.write_bytes(key)
    monkeypatch.setenv('STRUCTURED_SERVICE_TOKEN_FILE',str(path));monkeypatch.setenv('STRUCTURED_TENANT_ID','10000')
    with TestClient(create_app(store.root)) as client:
        raw,headers=signed(key,body={'operation':'list_datasets','arguments':{'question':'20天原料'}})
        result=client.post('/api/tools/read',content=raw,headers=headers).json()
        assert {f['series_id'] for f in result['literal_fields']}=={'C','E','G','I'}
        assert result['literal_field_count']==4
        assert all(f['source_context']['formula_source'] for f in result['literal_fields'])
        assert not result['literal_fields_truncated']
        assert 'rows' not in result


def test_successful_binding_reuse_keeps_signed_scope_checks_and_queries_live(duplicate_source, monkeypatch):
    from fastapi.testclient import TestClient
    from structured_data.server import create_app
    from test_scoped_service import signed
    store, dataset = duplicate_source
    model_calls, query_calls = [], []
    def model(payload):
        model_calls.append(payload)
        return {'decision': 'resolved', 'bindings': [
            {'considered_ids': ['C', 'E', 'G', 'I'], 'selected_ids': ['C'],
             'reason': 'The complete question selects this actual formula source.'}]}
    monkeypatch.setattr(FieldResolver, '__init__', lambda self, data: self.__dict__.update(store=data, model_fn=model))
    original_query = DataStore.query_series
    def query(self, *args, **kwargs):
        query_calls.append(True)
        return original_query(self, *args, **kwargs)
    monkeypatch.setattr(DataStore, 'query_series', query)
    key = b'success-binding-scope-key-1234567890'
    path = store.root / 'read-key'
    path.write_bytes(key)
    monkeypatch.setenv('STRUCTURED_SERVICE_TOKEN_FILE', str(path))
    monkeypatch.setenv('STRUCTURED_TENANT_ID', '10000')
    body = {'operation': 'query_series', 'arguments': {
        'dataset_id': dataset, 'series_id': 'C', 'source_question': 'Complete source question'}}
    with TestClient(create_app(store.root)) as client:
        raw, headers = signed(key, body=body)
        for _ in range(2):
            result = client.post('/api/tools/read', content=raw, headers=headers)
            assert result.status_code == 200
            assert result.json()['rows'][0]['value'] == 100
        assert len(model_calls) == 1
        assert len(query_calls) == 2  # Only the semantic receipt is reused.
        assert client.post('/api/tools/read', content=raw + b' ', headers=headers).status_code == 401
        other_raw, other_headers = signed(key, body=body, tenant=10001)
        assert client.post('/api/tools/read', content=other_raw, headers=other_headers).status_code == 400
        assert len(model_calls) == 1 and len(query_calls) == 2
