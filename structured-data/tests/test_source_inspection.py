import hashlib
import json
import threading
import zipfile
import xml.etree.ElementTree as ET

from fastapi.testclient import TestClient
from openpyxl import Workbook
import pytest

from structured_data import source_inspection
from structured_data.server import create_app
from structured_data.chart_source_cache import SourceEvidenceReadBudgetExceeded
from structured_data.source_inspection import (
    SourceInspectionReadError,
    inspect_query_source_context,
    inspect_source,
)
from structured_data.store import DataStore, dump
from test_scoped_service import signed


@pytest.fixture
def original(tmp_path):
    store = DataStore(tmp_path)
    book = Workbook()
    daily = book.active
    daily.title = '月度'
    daily.append(['日期', '产量', '进口'])
    daily.append(['2022-09-01', 650, 127.008062964705883])
    daily['D2'] = '=B2+C2'  # openpyxl does not invent an Excel formula cache.
    for row in range(3, 14):
        daily.cell(row, 2, 10)
        daily.cell(row, 3, 1)
    balance = book.create_sheet('年度及小表')
    for cell, value in {'AF77': '年度', 'AG77': '产量', 'AH77': '进口',
                        'AF83': '22/23', 'AG83': 650, 'AH83': 125,
                        'AF85': '年度', 'AG85': '产量', 'AH85': '进口',
                        'AF90': '22/23', 'AG90': 650, 'AH90': 127,
                        'BX25': '棉花产量/进口', 'BX31': '22/23', 'BY31': 650, 'BZ31': 125}.items():
        balance[cell] = value
    layout = book.create_sheet('横向标签')
    layout.append(['Metric', '2022/23', '2023/24'])
    layout.append(['Supply', 7, 8])
    quoted = book.create_sheet("研究员's 表")
    quoted['B2'], quoted['C2'] = 12, 3
    source = tmp_path / 'fixture.xlsx'
    book.save(source)
    sha = hashlib.sha256(source.read_bytes()).hexdigest()
    dataset = hashlib.sha256((sha + ':test').encode()).hexdigest()[:24]
    path = store.datasets / dataset
    path.mkdir()
    (path / 'source.xlsx').write_bytes(source.read_bytes())
    dump(path / 'manifest.json', {'dataset_id': dataset, 'source_file': source.name, 'sha256': sha,
        'sheets': [{'name': '年度及小表', 'parsed_range': 'A2:Y170'}],
        'excluded_regions': [{'sheet': '年度及小表', 'range': 'AF77:BZ90', 'reason': 'not parsed'}]})
    return store, dataset, path


def test_searches_all_columns_and_layouts_without_choosing_business_definition(original):
    store, dataset, _ = original
    result = inspect_source(store, dataset, searches=['22/23', '2022/23'], max_cells=400)
    coords = {(cell['sheet'], cell['cell']) for cell in result['matches']}
    assert coords == {('年度及小表', 'AF83'), ('年度及小表', 'AF90'),
                      ('年度及小表', 'BX31'), ('横向标签', 'B1')}
    assert result['coverage']['年度及小表']['parsed_ranges'] == ['A2:Y170']
    assert result['hash_verified'] and result['scan']['complete'] and not result['truncated']
    assert result['matches'][0]['context_cells']
    assert 'unit' not in result and 'preferred_result' not in result


def test_exact_ranges_and_decimal_calculation_keep_conflicting_source_values(original):
    store, dataset, _ = original
    result = inspect_source(store, dataset, sheet='年度及小表', ranges=['AF77:AH90'],
                            sums=[['AG83', 'AH83'], ['AG90', 'AH90'], ['BY31', 'BZ31']])
    cells = {cell['cell']: cell['value'] for cell in result['ranges'][0]['cells']}
    assert len(cells) == 42 and cells['AH83'] == 125 and cells['AH90'] == 127
    assert [calc['value'] for calc in result['calculations']] == ['775', '777', '775']
    assert [cell['cell'] for cell in result['calculations'][1]['cells']] == ['AG90', 'AH90']
    assert result['calculations'][0]['sha256'] == result['sha256']
    monthly = inspect_source(store, dataset, sheet='月度', sums=[['B2', 'C2']])
    assert monthly['calculations'][0]['value'] == '777.0080629647059'


def test_excel_references_normalize_format_without_changing_source_or_scope(original):
    store, dataset, _ = original
    result = inspect_source(store, dataset, sheet='年度及小表',
        ranges=["'年度及小表'!$af$83:$ah$83", 'AF90:AH90'],
        sums=[['年度及小表!$ag$83', 'ah83'], ["'年度及小表'!ag90", "$AH$90"]])
    assert [item['range'] for item in result['ranges']] == ['AF83:AH83', 'AF90:AH90']
    assert [item['value'] for item in result['calculations']] == ['775', '777']
    assert result['sheets'] == ['年度及小表']
    assert [cell['cell'] for cell in result['calculations'][0]['cells']] == ['AG83', 'AH83']
    inferred = inspect_source(store, dataset, ranges=["'年度及小表'!af83:ah83"],
        sums=[['年度及小表!ag83', '年度及小表!ah83']])
    assert inferred['sheets'] == ['年度及小表'] and inferred['calculations'][0]['value'] == '775'
    apostrophe = inspect_source(store, dataset, sums=[["'研究员''s 表'!$b$2", "'研究员''s 表'!c2"]])
    assert apostrophe['sheets'] == ["研究员's 表"] and apostrophe['calculations'][0]['value'] == '15'


def test_bounded_sum_ranges_expand_every_source_cell_and_use_decimal(original):
    store, dataset, _ = original
    result = inspect_source(store, dataset, sheet='月度', sums=[['$b$2:$b$13', "'月度'!c2:c13"]])
    calculation = result['calculations'][0]
    assert calculation['value'] == '898.0080629647059'
    assert len(calculation['cells']) == 24
    assert [cell['cell'] for cell in calculation['cells']] == [f'B{i}' for i in range(2, 14)] + [f'C{i}' for i in range(2, 14)]
    assert calculation['cells'][0]['value'] == 650 and calculation['cells'][12]['value'] == 127.0080629647059
    assert all(cell['sheet'] == '月度' and cell['value_status'] == 'original' for cell in calculation['cells'])
    assert calculation['sha256'] == result['sha256']
    one_range = inspect_source(store, dataset, sums=[["'年度及小表'!AG83:AH83"]])
    assert one_range['calculations'][0]['value'] == '775'
    separate_groups = inspect_source(store, dataset, sheet='月度', sums=[['B2:B13'], ['B2:B13']], max_cells=24)
    assert [item['value'] for item in separate_groups['calculations']] == ['760', '760']
    assert all(len(item['cells']) == 12 for item in separate_groups['calculations'])


@pytest.mark.parametrize('arguments, message', [
    ({'sheet': '月度', 'sums': [['B2:B13', 'B13:C13']]}, '重复坐标'),
    ({'sheet': '月度', 'sums': [['B2:B13', '$b$2']]}, '重复坐标'),
    ({'sums': [['月度!B2:B13', '年度及小表!AG83:AH83']]}, '不同工作表'),
    ({'sheet': '月度', 'sums': [['A1:XFD1048576']]}, '超过 max_cells'),
    ({'sheet': '月度', 'sums': [['B2:B13'], ['B2:B13']], 'max_cells': 23}, '超过 max_cells'),
    ({'sheet': '月度', 'sums': [['B2']]}, '至少需要 2'),
    ({'sheet': '月度', 'sums': [['B2:D2']]}, '有效数值'),
    ({'sheet': '月度', 'sums': [['C3:D3']]}, '有效数值'),
    ({'sheet': '月度', 'sums': [['B:B', 'C:C']]}, 'A1 引用'),
    ({'sheet': '月度', 'sums': [['SUM(B2:C2)']]}, 'A1 引用'),
])
def test_sum_ranges_reject_overlap_cross_sheet_budget_and_missing_values(original, arguments, message):
    store, dataset, _ = original
    with pytest.raises(ValueError, match=message):
        inspect_source(store, dataset, **arguments)


def test_sum_budget_rejected_before_source_read_or_large_expansion(original, monkeypatch):
    store, dataset, _ = original
    def unexpected_read(*args):
        pytest.fail('Oversized references must be rejected before inspecting any dataset source')
    monkeypatch.setattr(store, '_path', unexpected_read)
    with pytest.raises(ValueError, match='超过 max_cells'):
        inspect_source(store, dataset, sheet='月度', sums=[['A1:XFD1048576']])


@pytest.mark.parametrize('arguments', [
    {'sheet': '月度', 'ranges': ['年度及小表!A1']},
    {'ranges': ['月度!A1', '年度及小表!A1']},
    {'ranges': ['月度!A1', 'A2']},
    {'sums': [['月度!B2', '年度及小表!AH83']]},
    {'ranges': ["'[other.xlsx]月度'!A1"]},
    {'sheet': '月度', 'ranges': ['../book!A1']},
    {'sheet': '月度', 'ranges': ['=SUM(B2:C2)']},
    {'ranges': ["'不存在的工作表'!A1"]},
    {'sheet': '月度', 'sums': [['$B$2', 'b2']]},
])
def test_excel_reference_compatibility_does_not_enable_cross_sheet_or_external_reads(original, arguments):
    store, dataset, _ = original
    with pytest.raises(ValueError):
        inspect_source(store, dataset, **arguments)


def test_formula_raw_value_and_missing_cache_never_invented(original):
    store, dataset, _ = original
    cell = inspect_source(store, dataset, sheet='月度', ranges=['D2'])['ranges'][0]['cells'][0]
    assert cell['formula'] == '=B2+C2' and cell['cached_value'] is None
    assert cell['value'] is None and cell['value_status'] == 'missing_formula_cache'
    for refs in [['B2', 'D2'], ['B2', 'Z2'], ['B2', 'A2']]:
        with pytest.raises(ValueError, match='有效数值'):
            inspect_source(store, dataset, sheet='月度', sums=[refs])


def test_saved_formula_cache_is_used_without_recalculating(original):
    store, dataset, path = original
    source = path / 'source.xlsx'
    with zipfile.ZipFile(source) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    # Set deliberately stale cache 999. Source inspection must preserve it,
    # rather than calculate the mathematically expected B2+C2 itself.
    sheet = ET.fromstring(members['xl/worksheets/sheet1.xml'])
    namespace = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    formula = sheet.find('.//s:c[@r="D2"]/s:f', namespace)
    value = sheet.find('.//s:c[@r="D2"]/s:v', namespace)
    assert formula.text == 'B2+C2'
    value.text = '999'
    members['xl/worksheets/sheet1.xml'] = ET.tostring(sheet)
    with zipfile.ZipFile(source, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    manifest = store.manifest(dataset)
    manifest['sha256'] = hashlib.sha256(source.read_bytes()).hexdigest()
    dump(path / 'manifest.json', manifest)
    result = inspect_source(store, dataset, sheet='月度', ranges=['D2'], sums=[['B2', 'D2']])
    cell = result['ranges'][0]['cells'][0]
    assert cell['formula'] == '=B2+C2' and cell['value'] == cell['cached_value'] == 999
    assert result['calculations'][0]['value'] == '1649'


def test_output_budget_refuses_cutting_source_values(original, monkeypatch):
    store, dataset, _ = original
    monkeypatch.setattr('structured_data.source_inspection.MAX_RESULT_BYTES', 100)
    with pytest.raises(ValueError, match='输出预算'):
        inspect_source(store, dataset, sheet='月度', ranges=['B2:C2'])


def test_query_source_context_classifies_budget_and_safe_read_failure(original, monkeypatch, caplog):
    store, dataset, path = original
    record = [{'source_cell': 'B2', 'value': 650}]
    scan_budget = source_inspection.MAX_SCAN_SECONDS
    monkeypatch.setattr('structured_data.source_inspection.MAX_SCAN_SECONDS', -1)
    budget = inspect_query_source_context(store, dataset, '月度', record)
    assert budget['status'] == 'unavailable' and budget['reason'] == 'source_read_budget_exceeded'
    assert budget['phase'] == 'range_read'
    assert budget['scope']['failed_source_scope'] == {'sheet': '月度', 'range': 'B2'}
    assert 'cells' not in budget and '不表示' in budget['notice']
    assert SourceEvidenceReadBudgetExceeded(phase='range_read').phase == 'range_read'

    monkeypatch.setattr('structured_data.source_inspection.MAX_SCAN_SECONDS', scan_budget)
    def fail_open(*args, **kwargs):
        raise OSError(f'{path / "private-source.xlsx"}: private workbook detail')
    monkeypatch.setattr('structured_data.source_inspection.load_workbook', fail_open)
    failed = inspect_query_source_context(store, dataset, '月度', record)
    assert failed['status'] == 'unavailable'
    assert failed['reason'] == 'original_source_unavailable_or_verification_failed'
    assert failed['phase'] == 'workbook_open' and failed['exception_type'] == 'OSError'
    serialized = json.dumps(failed)
    assert 'private-source.xlsx' not in serialized and 'private workbook detail' not in serialized
    assert 'private workbook detail' in caplog.text


def test_budget_failure_is_typed_and_supplementary_api_reports_range(original, monkeypatch):
    store, dataset, _ = original
    monkeypatch.setattr('structured_data.source_inspection.MAX_SCAN_SECONDS', -1)
    with pytest.raises(SourceEvidenceReadBudgetExceeded) as caught:
        inspect_source(store, dataset, sheet='月度', ranges=['B2:C2'])
    assert caught.value.phase == 'range_read'
    assert caught.value.scope == {'sheet': '月度', 'range': 'B2:C2'}

    key = b'budget-inspection-key-with-enough-length'
    key_path = store.root / 'key'
    key_path.write_bytes(key)
    monkeypatch.setenv('STRUCTURED_SERVICE_TOKEN_FILE', str(key_path))
    monkeypatch.setenv('STRUCTURED_TENANT_ID', '10000')
    with TestClient(create_app(store.root)) as client:
        raw, headers = signed(key, body={'operation': 'inspect_source', 'arguments': {
            'dataset_id': dataset, 'sheet': '月度', 'ranges': ['B2:C2']}})
        response = client.post('/api/tools/read', content=raw, headers=headers)
    result = response.json()
    assert response.status_code == 200
    assert result['status'] == 'error' and result['error'] == 'source_read_budget_exceeded'
    assert result['phase'] == 'range_read' and result['scope'] == {'sheet': '月度', 'range': 'B2:C2'}
    assert result['retryable'] and '不表示原表没有证据' in result['next_step']
    assert 'ranges' not in result and 'cells' not in result


def test_source_read_failure_returns_limited_type_and_phase(original, monkeypatch, caplog):
    store, dataset, path = original
    key = b'format-inspection-key-with-enough-length'
    key_path = store.root / 'key'
    key_path.write_bytes(key)
    monkeypatch.setenv('STRUCTURED_SERVICE_TOKEN_FILE', str(key_path))
    monkeypatch.setenv('STRUCTURED_TENANT_ID', '10000')

    def fail_open(*args, **kwargs):
        raise OSError(f'{path / "private-source.xlsx"}: private format detail')

    monkeypatch.setattr('structured_data.source_inspection.load_workbook', fail_open)
    with pytest.raises(SourceInspectionReadError) as caught:
        inspect_source(store, dataset, sheet='月度', ranges=['B2'])
    assert caught.value.phase == 'workbook_open' and caught.value.exception_type == 'OSError'

    with TestClient(create_app(store.root)) as client:
        raw, headers = signed(key, body={'operation': 'inspect_source', 'arguments': {
            'dataset_id': dataset, 'sheet': '月度', 'ranges': ['B2']}})
        response = client.post('/api/tools/read', content=raw, headers=headers)
    result = response.json()
    assert response.status_code == 200 and result['status'] == 'error'
    assert result['error'] == 'source_read_failed'
    assert result['phase'] == 'workbook_open' and result['exception_type'] == 'OSError'
    assert not result['retryable']
    serialized = json.dumps(result)
    assert 'private-source.xlsx' not in serialized and 'private format detail' not in serialized
    assert 'private format detail' in caplog.text


def test_hash_change_rejected_and_other_store_cannot_access(original, tmp_path):
    store, dataset, path = original
    with pytest.raises(ValueError, match='不存在'):
        inspect_source(DataStore(tmp_path / 'other'), dataset, searches=['22/23'])
    (path / 'source.xlsx').write_bytes((path / 'source.xlsx').read_bytes() + b'changed')
    with pytest.raises(ValueError, match='SHA256'):
        inspect_source(store, dataset, searches=['22/23'])


@pytest.mark.parametrize('arguments', [
    {'ranges': ['A1']}, {'sheet': '月度', 'ranges': ['A1:B999']},
    {'sheet': '月度', 'ranges': ['A0']}, {'sheet': '月度', 'ranges': ['B2:A1']},
    {'sheet': '月度', 'ranges': ['XFE1']}, {'searches': ['']},
    {'searches': ['x'] * 9}, {'searches': ['x' * 201]},
    {'searches': ['x'], 'max_cells': 401}, {'searches': ['x'], 'max_matches': True},
    {'searches': ['x'], 'padding_cols': 5}, {'sheet': 'missing', 'searches': ['x']},
    {'sheet': '月度', 'sums': [['B2', 'B2']]}, {'sheet': '月度', 'sums': [['B2:C2', 'C2']]},
])
def test_invalid_or_oversized_inputs_rejected(original, arguments):
    store, dataset, _ = original
    with pytest.raises(ValueError):
        inspect_source(store, dataset, **arguments)


def test_scan_and_output_limits_are_explicit_not_a_false_absence(original, monkeypatch):
    store, dataset, _ = original
    result = inspect_source(store, dataset, searches=['22/23'], max_matches=1)
    assert len(result['matches']) == 1 and result['truncated']
    assert 'max_matches' in result['truncation_reasons'] and not result['scan']['complete']
    monkeypatch.setattr('structured_data.source_inspection.MAX_SCAN_CELLS', 100)
    result = inspect_source(store, dataset, sheet='年度及小表', searches=['22/23'])
    assert not result['matches'] and result['truncated'] and result['scan']['cells_scanned'] == 100
    assert 'scan_cells_limit' in result['truncation_reasons']


def test_whole_workbook_budget_names_unvisited_sheets_and_directs_targeted_retry(original, monkeypatch):
    store, dataset, _ = original
    monkeypatch.setattr('structured_data.source_inspection.MAX_SCAN_CELLS', 6)
    result = inspect_source(store, dataset, searches=['2022/23'])
    assert result['matches'] == [] and result['truncated'] and not result['scan']['complete']
    assert result['scan']['sheets_scanned'] == ['月度']
    assert result['scan']['incomplete_sheets'] == ['月度']
    assert result['scan']['unscanned_sheets'] == ['年度及小表', '横向标签', "研究员's 表"]
    assert '不能证明' in result['next_step'] and '真实 sheet' in result['next_step']
    assert '不要再次省略 sheet' in result['next_step']
    # Retry is scoped by the real source sheet, not a guessed first sheet or
    # a larger budget. The existing six-cell limit already covers this layout.
    targeted = inspect_source(store, dataset, sheet='横向标签', searches=['2022/23'], padding_rows=0, padding_cols=0)
    assert not targeted['truncated'] and targeted['scan']['complete']
    assert targeted['scan']['unscanned_sheets'] == targeted['scan']['incomplete_sheets'] == []
    assert targeted['matches'][0]['sheet'] == '横向标签' and targeted['matches'][0]['cell'] == 'B1'
    assert 'next_step' not in targeted


def test_native_schema_and_signature_keep_workspace_boundaries(original, monkeypatch):
    store, dataset, _ = original
    key = b'original-inspection-key-with-enough-length'
    key_path = store.root / 'key'
    key_path.write_bytes(key)
    monkeypatch.setenv('STRUCTURED_SERVICE_TOKEN_FILE', str(key_path))
    monkeypatch.setenv('STRUCTURED_TENANT_ID', '10000')
    with TestClient(create_app(store.root)) as client:
        def call(operation, arguments, tenant=10000):
            raw, headers = signed(key, body={'operation': operation, 'arguments': arguments}, tenant=tenant)
            return client.post('/api/tools/read', content=raw, headers=headers)
        schema = call('describe_tools', {'operations': ['inspect_source']}).json()['tools'][0]['inputSchema']
        assert schema['required'] == ['dataset_id'] and 'sums' in schema['properties']
        args = {'dataset_id': dataset, 'sheet': '年度及小表', 'sums': [['AG83', 'AH83']]}
        response = call('inspect_source', args)
        assert response.status_code == 200 and response.json()['calculations'][0]['value'] == '775'
        other_space = call('inspect_source', args, tenant=10001)
        assert other_space.status_code == 200 and other_space.json()['status'] == 'error'
        assert '数据集不存在' in other_space.json()['error']
        assert 'source_file' not in other_space.json() and 'sha256' not in other_space.json()
        assert call('inspect_source', args | {'tenant_id': 10001}).status_code == 400
        assert call('inspect_source', {'ranges': ['A1']}).status_code == 400
        oversized = call('inspect_source', {'dataset_id': dataset, 'sheet': '月度', 'ranges': ['A1:B999']})
        assert oversized.status_code == 200
        result = oversized.json()
        assert result['kind'] == 'structured_source_inspection' and result['status'] == 'error'
        assert result['retryable'] and '超过 max_cells' in result['error']
        assert '缩小' in result['next_step'] and 'ranges' not in result
        source = store._path(dataset) / 'source.xlsx'
        source.write_bytes(source.read_bytes() + b'changed')
        damaged = call('inspect_source', args).json()
        assert damaged['status'] == 'error' and not damaged['retryable']
        assert 'SHA256' in damaged['error'] and '重新导入' in damaged['next_step']
        assert damaged['phase'] == 'source_verification'
        assert damaged['exception_type'] == 'SourceVerificationError'


def test_read_admission_wait_is_typed_scoped_and_happens_after_argument_validation(original, monkeypatch):
    store, dataset, _ = original

    class DeniedSemaphore:
        called = False

        def acquire(self, timeout=None):
            self.called = True
            return False

        def release(self):
            pytest.fail('a semaphore slot that was not acquired must not be released')

    denied = DeniedSemaphore()
    monkeypatch.setattr(source_inspection, '_SOURCE_READ_SEMAPHORE', denied)
    monkeypatch.setattr(source_inspection, 'MAX_READ_ADMISSION_WAIT_SECONDS', 0.01)
    with pytest.raises(ValueError):
        inspect_source(store, dataset, sheet='月度', ranges=['A0'])
    assert not denied.called

    def source_lookup_must_not_run(*args, **kwargs):
        pytest.fail('admission timeout must happen before source lookup or workbook reads')

    monkeypatch.setattr(store, '_path', source_lookup_must_not_run)
    with pytest.raises(SourceEvidenceReadBudgetExceeded) as caught:
        inspect_source(store, dataset, sheet='月度', ranges=['B2:C2'], sums=[['B2', 'C2']])
    assert denied.called
    assert caught.value.phase == 'read_admission'
    assert caught.value.scope == {'sheet': '月度', 'ranges': ['B2:C2'],
                                  'sum_ranges': [['B2', 'C2']], 'search_count': 0}


def test_concurrent_source_reads_do_not_overlap_workbook_loads(original, monkeypatch):
    store, dataset, _ = original
    inner = threading.BoundedSemaphore(1)
    state_lock = threading.Lock()
    holders = 0
    contended = threading.Event()

    class ObservedSemaphore:
        def acquire(self, timeout=None):
            nonlocal holders
            with state_lock:
                if holders:
                    contended.set()
            acquired = inner.acquire(timeout=timeout)
            if acquired:
                with state_lock:
                    holders += 1
            return acquired

        def release(self):
            nonlocal holders
            with state_lock:
                holders -= 1
            inner.release()

    monkeypatch.setattr(source_inspection, '_SOURCE_READ_SEMAPHORE', ObservedSemaphore())
    real_load = source_inspection.load_workbook
    first_load_entered = threading.Event()
    release_first_load = threading.Event()
    load_lock = threading.Lock()
    active_loads = 0
    max_active_loads = 0
    load_count = 0

    def controlled_load(*args, **kwargs):
        nonlocal active_loads, max_active_loads, load_count
        with load_lock:
            load_count += 1
            this_load = load_count
            active_loads += 1
            max_active_loads = max(max_active_loads, active_loads)
        try:
            if this_load == 1:
                first_load_entered.set()
                assert release_first_load.wait(timeout=5)
            return real_load(*args, **kwargs)
        finally:
            with load_lock:
                active_loads -= 1

    monkeypatch.setattr(source_inspection, 'load_workbook', controlled_load)
    results, failures = [], []

    def read_source():
        try:
            results.append(inspect_source(store, dataset, sheet='月度', ranges=['B2:C2']))
        except Exception as exc:  # surfaced in the test thread below
            failures.append(exc)

    first = threading.Thread(target=read_source)
    second = threading.Thread(target=read_source)
    first.start()
    assert first_load_entered.wait(timeout=5)
    second.start()
    try:
        assert contended.wait(timeout=5)
        with load_lock:
            assert load_count == 1 and active_loads == max_active_loads == 1
    finally:
        release_first_load.set()
    first.join(timeout=10)
    second.join(timeout=10)
    assert not first.is_alive() and not second.is_alive()
    assert not failures
    assert len(results) == 2 and max_active_loads == 1


def test_source_read_exception_releases_admission_slot(original, monkeypatch):
    store, dataset, _ = original
    semaphore = threading.BoundedSemaphore(1)
    monkeypatch.setattr(source_inspection, '_SOURCE_READ_SEMAPHORE', semaphore)
    real_load = source_inspection.load_workbook

    def fail_open(*args, **kwargs):
        raise OSError('temporary workbook read failure')

    monkeypatch.setattr(source_inspection, 'load_workbook', fail_open)
    with pytest.raises(SourceInspectionReadError) as caught:
        inspect_source(store, dataset, sheet='月度', ranges=['B2'])
    assert caught.value.phase == 'workbook_open'

    monkeypatch.setattr(source_inspection, 'load_workbook', real_load)
    result = inspect_source(store, dataset, sheet='月度', ranges=['B2'])
    assert result['hash_verified'] and result['ranges'][0]['cells'][0]['value'] == 650
    assert semaphore.acquire(blocking=False)
    semaphore.release()
