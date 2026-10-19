import hashlib

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
import pytest

from structured_data.store import DataStore


def source_fixture(tmp_path):
    source = tmp_path / 'fixture.xlsx'
    book = Workbook()
    data = book.active
    data.title = 'Data'
    data.append(['Month', 'Widget total', 'Other total'])
    for month in range(1, 11):
        data.append([f'2025-{month:02}', 119 + month, 0.20 + month / 100])

    plot = book.create_sheet('Charts')
    for column, title, axis in [(2, 'Widget inventory', 'Quantity (units)'),
                                (3, 'Other metric', 'Rate (fraction)')]:
        chart = LineChart()
        chart.title = title
        chart.x_axis.title = 'Month'
        chart.y_axis.title = axis
        chart.add_data(Reference(data, min_col=column, max_col=column,
                                 min_row=1, max_row=11), titles_from_data=True)
        chart.set_categories(Reference(data, min_col=1, min_row=2, max_row=11))
        plot.add_chart(chart, f'A{1 if column == 2 else 18}')
    book.save(source)

    sha = hashlib.sha256(source.read_bytes()).hexdigest()
    store = DataStore(tmp_path / 'store')
    parsed = {
        'source_file': str(source),
        'sha256': sha,
        'rule_version': 'test',
        'series': [
            {'series_id': 'widgets', 'sheet': 'Data', 'metric': 'Widget total',
             'unit': None, 'unit_status': 'needs_review', 'dimensions': {},
             'header_cells': ['B1']},
            {'series_id': 'other', 'sheet': 'Data', 'metric': 'Other total',
             'unit': None, 'unit_status': 'needs_review', 'dimensions': {},
             'header_cells': ['C1']},
        ],
        'observations': [
            {'series_id': series_id, 'date': f'2025-{month:02}-01',
             'value': value, 'source_cell': f'{column}{month + 1}',
             'date_cell': f'A{month + 1}', 'raw_value': value,
             'value_status': 'valid', 'formula': None}
            for series_id, column, value_fn in [
                ('widgets', 'B', lambda month: 120 + month - 1),
                ('other', 'C', lambda month: 0.20 + month / 100),
            ] for month in range(1, 11) for value in [value_fn(month)]
        ],
        'sheets': [{'name': 'Data', 'parsed_range': 'A1:C11'}],
        'excluded_regions': [],
        'issues': [],
    }
    dataset = store.publish(parsed)['dataset_id']
    store.chart_index(dataset)
    folder = store.datasets / dataset
    return store, dataset, folder


def test_query_attaches_chart_evidence_without_assigning_unit(tmp_path):
    store, dataset, _ = source_fixture(tmp_path)

    result = store.query_series(dataset, 'widgets', '2025-01-01', '2025-01-01')

    assert result['status'] == 'ok'
    assert result['rows'][0]['source_cell'] == 'B2'
    assert result['rows'][0]['value'] == 120
    assert result['series']['unit'] is None
    evidence = result['source_evidence']
    assert evidence['status'] == 'ok'
    assert evidence['source']['hash_verified']
    assert evidence['scope']['requested_value_cells'] == ['B2']
    assert evidence['scope']['truncated'] is True
    assert evidence['scope']['chart_scope_complete'] is False
    assert 'candidate_index_source_sha_unbound' in evidence['scope']['chart_scope_reasons']
    assert evidence['cells'][0]['cell'] == 'B2'
    assert evidence['cells'][0]['value'] == 120
    assert len(evidence['chart_references']) == 1
    chart = evidence['chart_references'][0]
    assert chart['title'] == 'Widget inventory'
    assert any('Quantity (units)' in axis for axis in chart['axis_titles'])
    assert chart['value_ref'] == "'Data'!$B$2:$B$11"
    assert chart['intersecting_references']['value']['range']['column'] == 'B'
    assert 'unit' not in chart
    assert evidence['source']['file'] == 'fixture.xlsx'
    assert str(store.root) not in str(evidence)


def test_query_chart_evidence_does_not_cross_to_another_series_range(tmp_path):
    store, dataset, _ = source_fixture(tmp_path)

    result = store.query_series(dataset, 'other', '2025-01-01', '2025-01-01')

    evidence = result['source_evidence']
    assert evidence['status'] == 'ok'
    assert evidence['scope']['requested_value_cells'] == ['C2']
    assert len(evidence['chart_references']) == 1
    assert evidence['chart_references'][0]['title'] == 'Other metric'
    assert evidence['chart_references'][0]['intersecting_references']['value']['range']['column'] == 'C'
    assert all(chart['title'] != 'Widget inventory'
               for chart in evidence['chart_references'])


def test_query_source_context_caps_cells_and_marks_partial_scope(tmp_path):
    store, dataset, _ = source_fixture(tmp_path)

    result = store.query_series(dataset, 'widgets', limit=20)

    evidence = result['source_evidence']
    assert evidence['status'] == 'ok'
    assert len(evidence['cells']) == 8
    assert len(evidence['scope']['requested_value_cells']) == 8
    assert evidence['scope']['requested_value_cell_count'] == 10
    assert evidence['scope']['value_cells_truncated'] is True
    assert evidence['scope']['truncated'] is True
    assert len(evidence['chart_references']) <= 8
    assert evidence['scope']['query_rows_truncated'] is False
    assert '不代表全部聚合输入' in evidence['scope']['query_notice']


def test_query_source_evidence_does_not_make_missing_unit_comparable(tmp_path):
    store, dataset, _ = source_fixture(tmp_path)

    with pytest.raises(ValueError, match='单位不明确或不一致'):
        store.compare_series(dataset, 'widgets', 'other', '2025-01-01')


def test_missing_original_keeps_successful_query_values_and_marks_unavailable(tmp_path):
    store, dataset, folder = source_fixture(tmp_path)
    (folder / 'source.xlsx').unlink()

    result = store.query_series(dataset, 'widgets', '2025-01-01', '2025-01-01')

    assert result['status'] == 'ok'
    assert result['rows'][0]['value'] == 120
    assert result['rows'][0]['source_cell'] == 'B2'
    assert result['series']['unit'] is None
    evidence = result['source_evidence']
    assert evidence['status'] == 'unavailable'
    assert evidence['reason'] == 'original_source_unavailable_or_verification_failed'
    assert 'source.xlsx' not in str(evidence)
    assert str(folder) not in str(evidence)
    assert '不表示原件没有单位' in evidence['notice']


def test_full_analysis_inputs_keep_bounded_original_source_evidence(tmp_path):
    store, dataset, _ = source_fixture(tmp_path)

    result = store.query_series(dataset, 'widgets', _complete=True)

    assert len(result['rows']) == 10
    assert result['series']['unit'] is None
    evidence = result['source_evidence']
    assert len(evidence['cells']) == 8
    assert evidence['scope']['requested_value_cell_count'] == 10
    assert evidence['scope']['value_cells_truncated'] is True
    assert evidence['scope']['query_rows_returned_count'] == 10
    assert evidence['chart_references'][0]['title'] == 'Widget inventory'
