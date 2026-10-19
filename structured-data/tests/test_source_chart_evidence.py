import hashlib
import json

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel

from structured_data.source_inspection import inspect_source
from structured_data.store import DataStore, dump
from structured_data.native_charts import VERSION, extract


def chart_source(tmp_path):
    store = DataStore(tmp_path / 'store')
    book = Workbook()
    data = book.active
    data.title = 'Data'
    data.append(['Month'] + [f'Primary {i}' for i in range(1, 12)]
                + [f'Secondary {i}' for i in range(1, 11)])
    for row in range(2, 5):
        data.append([f'M{row-1}'] + [row * col for col in range(1, 22)])

    plot = book.create_sheet('Plot')
    primary = LineChart()
    primary.title = 'Operations by month'
    primary.x_axis.title = 'Month'
    primary.y_axis.title = 'Revenue (USD)'
    primary.add_data(Reference(data, min_col=2, max_col=12, min_row=1, max_row=4),
                     titles_from_data=True)
    for index, series in enumerate(primary.series, 1):
        series.tx = SeriesLabel(v=f'Primary {index}')
    primary.set_categories(Reference(data, min_col=1, min_row=2, max_row=4))

    secondary = LineChart()
    secondary.y_axis.title = 'Temperature (°C)'
    secondary.y_axis.axId = 200
    secondary.y_axis.crosses = 'max'
    secondary.x_axis.axId = primary.x_axis.axId
    secondary.add_data(Reference(data, min_col=13, max_col=22, min_row=1, max_row=4),
                       titles_from_data=True)
    for index, series in enumerate(secondary.series, 1):
        series.tx = SeriesLabel(v=f'Secondary {index}')
    secondary.set_categories(Reference(data, min_col=1, min_row=2, max_row=4))
    primary += secondary
    plot.add_chart(primary, 'A1')

    source = tmp_path / 'source.xlsx'
    book.save(source)
    sha = hashlib.sha256(source.read_bytes()).hexdigest()
    dataset = hashlib.sha256((sha + ':charts').encode()).hexdigest()[:24]
    folder = store.datasets / dataset
    folder.mkdir(parents=True)
    (folder / 'source.xlsx').write_bytes(source.read_bytes())
    dump(folder / 'manifest.json', {'dataset_id': dataset, 'source_file': source.name,
        'sha256': sha, 'sheets': [], 'excluded_regions': []})
    charts = extract(folder / 'source.xlsx')
    dump(folder / f'{VERSION}.json', {'version': VERSION, 'dataset_id': dataset,
        'source_sha256': sha, 'chart_count': len(charts), 'charts': charts})
    return store, dataset, sha


def test_explicit_source_range_returns_intersecting_native_chart_text(tmp_path):
    store, dataset, sha = chart_source(tmp_path)
    result = inspect_source(store, dataset, sheet='Data', ranges=['B2:B4'])

    assert result['hash_verified'] and result['sha256'] == sha
    assert len(result['chart_references']) == 1
    chart = result['chart_references'][0]
    assert chart['title'] == 'Operations by month'
    assert chart['curve_label'] == 'Primary 1'
    assert chart['value_ref'] == "'Data'!$B$2:$B$4"
    assert chart['source_references']['category']['range']['column'] == 'A'
    assert chart['part'].startswith('xl/charts/chart')
    assert chart['source_file'] == 'source.xlsx' and chart['sha256'] == sha
    assert 'unit' not in chart and 'Secondary 1' not in chart['curve_label']
    assert any('Revenue (USD)' in title for title in chart['axis_titles'])


def test_chart_evidence_is_bounded_and_marks_omitted_references(tmp_path):
    store, dataset, _ = chart_source(tmp_path)
    result = inspect_source(store, dataset, sheet='Data', ranges=['B2:V4'])

    assert len(result['chart_references']) == 20
    assert result['chart_references_truncated']
    assert result['truncated']
    assert 'chart_references_limit' in result['truncation_reasons']
    assert any('Temperature (°C)' in title
               for title in result['chart_references'][0]['axis_titles'])
    assert len(inspect_source(store, dataset, sheet='Data', ranges=['W2:W4'])
               ['chart_references']) == 0
