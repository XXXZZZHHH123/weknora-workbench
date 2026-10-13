import copy
from datetime import datetime
import json

from openpyxl import Workbook, load_workbook
import pytest

from structured_data.imports import ImportManager


def workbook(path, day=1, label='现货'):
    wb = Workbook()
    ws = wb.active
    ws.title = '简单表'
    ws.append(['日期', '价格'])
    ws.append([datetime(2023, 1, day), 100])
    ws.append([datetime(2023, 2, day), 120])
    ws = wb.create_sheet('复杂表')
    ws.append(['日期',label])
    ws.append([None,'品种甲'])
    ws.append([datetime(2023, 1, day), 10])
    ws.append([datetime(2023, 2, day), 12])
    wb.save(path)


def test_review_confirm_immutable_and_template_reuse(tmp_path):
    source = tmp_path/'source.xlsx'
    workbook(source)
    manager = ImportManager(tmp_path/'data')
    first = manager.ingest(source)
    assert first['summary'] == {'auto_ready':1,'needs_review':1}
    original_id = first['dataset_id']
    assert len(manager.store.list_series(original_id)['series']) == 1
    mapping = manager.sheet(first['import_id'],'复杂表')['mapping']
    preview = manager.preview(first['import_id'],mapping)
    assert preview['observation_count'] == 2
    assert len(manager.store.list_series(original_id)['series']) == 1
    confirmed = manager.confirm(first['import_id'], mapping)
    assert confirmed['dataset_id'] != original_id
    assert len(manager.store.list_series(confirmed['dataset_id'])['series']) == 2
    assert len(manager.store.list_series(original_id)['series']) == 1
    assert len(manager.list_templates()) == 1
    assert manager.ingest(source)['dataset_id'] == confirmed['dataset_id']
    fresh = tmp_path/'fresh.xlsx'
    workbook(fresh,day=2)
    second = manager.ingest(fresh)
    assert second['summary'] == {'auto_ready':1,'template_ready':1}
    changed = tmp_path/'changed.xlsx'
    workbook(changed,day=3,label='新口径')
    third = manager.ingest(changed)
    assert third['summary'] == {'auto_ready':1,'needs_review':1}
    assert manager.sheet(third['import_id'],'复杂表')['template_status'] == 'layout_changed'
    # Moving the data directory between host and container must not break source lookup.
    assert manager.source_path(manager.get(first['import_id'])).is_file()
    assert 'source_path' not in manager.get(first['import_id'])


def test_failed_confirm_preserves_published_data(tmp_path):
    source=tmp_path/'source.xlsx'
    workbook(source)
    manager=ImportManager(tmp_path/'data')
    first=manager.ingest(source)
    mapping=manager.sheet(first['import_id'],'复杂表')['mapping']
    mapping['blocks'][0]['date_mode']='year'
    with pytest.raises(ValueError):manager.confirm(first['import_id'],mapping)
    assert manager.get(first['import_id'])['dataset_id']==first['dataset_id']
    assert not manager.list_templates()


def test_suggestion_does_not_publish_or_persist_credentials(tmp_path, monkeypatch):
    import httpx
    source=tmp_path/'source.xlsx'
    workbook(source)
    manager=ImportManager(tmp_path/'data')
    result=manager.ingest(source)
    mapping=manager.sheet(result['import_id'],'复杂表')['mapping']
    config=tmp_path/'models.json'
    config.write_text(json.dumps({'chat':{'api_key':'private-test-key','base_url':'https://example.invalid/v1','model':'test'}}))
    monkeypatch.setenv('COTTON_MODEL_CONFIG',str(config))
    def fake_post(self,url,**kw):
        assert len(json.dumps(kw['json']))<20000
        return httpx.Response(200,json={'choices':[{'message':{'content':json.dumps({'mapping':mapping,'notice':'需确认'})}}]},request=httpx.Request('POST',url))
    monkeypatch.setattr(httpx.Client,'post',fake_post)
    suggested=manager.suggest(result['import_id'],'复杂表')
    assert suggested['mapping']['sheet']=='复杂表'
    assert manager.get(result['import_id'])['dataset_id']==result['dataset_id']
    assert manager.sheet(result['import_id'],'复杂表')['status']=='needs_review'
    assert not manager.list_templates()
    assert 'private-test-key' not in (manager.path(result['import_id'])/'review.json').read_text()


def test_pending_scope_never_falls_back_to_other_table(tmp_path):
    source=tmp_path/'source.xlsx'
    wb=Workbook();ws=wb.active;ws.title='国内价格'
    ws.append(['日期','产量']);ws.append([datetime(2023,1,1),100]);ws.append([datetime(2023,1,2),110])
    ws=wb.create_sheet('USDA棉花平衡表-全球')
    ws.append(['日期','产量']);ws.append([None,'预测年度']);ws.append([datetime(2023,1,1),200]);ws.append([datetime(2023,1,2),210])
    wb.save(source)
    manager=ImportManager(tmp_path/'data');r=manager.ingest(source)
    result=manager.store.list_series(r['dataset_id'],'USDA 全球 棉花 产量')
    assert result['series']==[]
    assert result['match_mode']=='pending_scope'
    assert result['pending_sheets'][0]['name']=='USDA棉花平衡表-全球'


def test_verified_sheet_edit_seed_uses_actual_fingerprint_bound_mapping(tmp_path, monkeypatch):
    from structured_data.profiles import reviewed
    import hashlib
    root=tmp_path/'data'
    root.mkdir()
    source=tmp_path/'source.xlsx'
    workbook(source)
    manager=ImportManager(root)
    imported=manager.ingest(source)
    doc=manager.get(imported['import_id'])
    scan=next(s for s in doc['sheets'] if s['name']=='简单表')
    scan['status']='trusted_existing'
    manager._save(doc)
    config_dir=tmp_path/'rules'
    (config_dir/'reviewed_mappings').mkdir(parents=True)
    monkeypatch.setattr(reviewed,'__file__',str(config_dir/'profiles'/'reviewed.py'))
    digest=hashlib.sha256(source.read_bytes()).hexdigest()
    rule={'sha256':digest,'sheets':[{'name':'简单表','start_row':2,
        'header_expectations':{'A1':'日期','B1':'价格'},
        'specs':[[2,1,'verified.price','已保存字段',{'来源':'实际核验规则'},'元',['B1']]]}]}
    (config_dir/'reviewed_mappings'/(digest+'.json')).write_text(json.dumps(rule))
    display=manager.sheet(imported['import_id'],'简单表')['mapping']
    assert not display.get('view_only')
    assert display['blocks'][0]['columns'][0]['metric']=='已保存字段'
    assert display['blocks'][0]['columns'][0]['series_id']=='verified.price'
    assert display['blocks'][0]['columns'][0]['header_cells']==['B1']
    assert display['blocks'][0]['data_start_row']==2
    assert manager.get(imported['import_id'])['sheets'][0]['mapping'] != display
    with manager.source_path(doc).open('ab') as stream:
        stream.write(b'changed')
    assert manager.sheet(imported['import_id'],'简单表')['mapping']['blocks']==[]
