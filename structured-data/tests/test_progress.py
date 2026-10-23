"""Progress reports completed work units without inventing overall duration."""
from datetime import datetime

from openpyxl import Workbook
import pytest

from structured_data.imports import ImportManager
from structured_data.progress import progress_value
from structured_data.review_worker import run_review_task, _monitor_worker
from structured_data.scanner import scan_workbook
from test_review_api import client, make_workbook, poll


def source_file(tmp_path):
    source = tmp_path / 'source.xlsx'
    book = Workbook()
    book.active.title = 'ready'
    book.active.append(['date', 'measure'])
    book.active.append([datetime(2024, 1, 1), 1])
    book.active.append([datetime(2024, 1, 2), 2])
    bad = book.create_sheet('bad')
    bad.append(['date', 'measure'])
    for _ in range(41):
        bad.append([datetime(2024, 1, 1), 3])
    book.create_sheet('empty')
    book.save(source)
    return source


def test_scanner_counts_every_completed_sheet_including_empty(tmp_path):
    events = []
    scan = scan_workbook(source_file(tmp_path), progress=events.append)
    measured = [event for event in events if 'total' in event]
    assert [(event['completed'], event['total']) for event in measured] == [
        (0, 3), (1, 3), (2, 3), (3, 3)]
    assert scan['sheet_count'] == 3
    assert all(event['phase'] == 'scan' for event in events)


def test_import_counts_failed_sheet_as_processed_and_leaves_publish_unquantified(tmp_path):
    events = []
    manager = ImportManager(tmp_path / 'data')
    result = manager.ingest(source_file(tmp_path), progress=events.append)
    parsed = [event for event in events if event['phase'] == 'parse']
    assert [(event['completed'], event['total']) for event in parsed] == [(0, 2), (1, 2), (2, 2)]
    assert manager.sheet(result['import_id'], 'bad')['import_state'] == 'failed'
    final = [event for event in events if event['phase'] in {'publish', 'index'}]
    assert [event['phase'] for event in final] == ['publish', 'index']
    assert all('total' not in event and 'completed' not in event for event in final)


def test_worker_progress_only_channel_does_not_require_unconfigured_publish_guard(tmp_path):
    events = []
    result = run_review_task(tmp_path / 'data', 'import', path=source_file(tmp_path),
                             progress_callback=events.append)
    assert result['dataset_id']
    assert {'scan', 'parse', 'publish', 'index'} <= {event['phase'] for event in events}


def test_worker_progress_keeps_final_guard_once_and_denial_unpublished(tmp_path):
    source = source_file(tmp_path)
    events, guards = [], []

    def deny():
        guards.append('denied')
        assert events[-1]['phase'] == 'publish'
        raise ValueError('权限已撤销')

    with pytest.raises(ValueError, match='权限已撤销'):
        run_review_task(tmp_path / 'data', 'import', path=source,
                        progress_callback=events.append, publish_guard=deny)
    assert guards == ['denied']
    assert not any(event['phase'] == 'index' for event in events)
    manager = ImportManager(tmp_path / 'data')
    assert manager.list() == [] and manager.store.list_datasets() == []


@pytest.mark.parametrize('value', [
    {'phase': 'arbitrary'}, {'phase': 'scan', 'completed': True, 'total': 2},
    {'phase': 'scan', 'completed': 3, 'total': 2},
    {'phase': 'scan', 'completed': 0, 'total': 0},
    {'phase': 'scan', 'completed': 1, 'total': 1_000_001},
    {'phase': 'scan', 'completed': 1}, {'phase': 'scan', 'source_path': 'private'},
])
def test_progress_rejects_invalid_counts_or_extra_metadata(value):
    with pytest.raises(ValueError):
        progress_value(value)


def test_progress_discards_arbitrary_child_stage_text():
    value = progress_value({'phase': 'scan', 'completed': 1, 'total': 2, 'stage': '/private/source'})
    assert value == {'phase': 'scan', 'completed': 1, 'total': 2, 'stage': '识别工作表结构'}


def test_progress_messages_cannot_repeat_or_replace_publish_authorization():
    class Process:
        pid = 123
        alive = True

        def is_alive(self):
            return self.alive

        def join(self, timeout=None):
            self.alive = False

    class Control:
        messages = [
            {'kind': 'progress', 'value': {'phase': 'parse', 'completed': 1, 'total': 1}},
            {'kind': 'authorize_publish'}, {'kind': 'authorize_publish'}]

        def __init__(self):
            self.replies = []

        def poll(self):
            return bool(self.messages)

        def recv(self):
            return self.messages.pop(0)

        def send(self, value):
            self.replies.append(value)

    guards, events, control = [], [], Control()
    _monitor_worker(Process(), float('inf'), 512, rss_reader=lambda pid: 0,
                    control=control, publish_guard=lambda: guards.append('once'),
                    progress_callback=events.append)
    assert guards == ['once'] and len(events) == 1
    assert control.replies[0] == {'ok': True}
    assert control.replies[1]['ok'] is False


@pytest.mark.parametrize('deadline,rss,expected', [
    (0, 0, '时间限制'), (float('inf'), 513, '内存限制'),
])
def test_continuous_progress_cannot_bypass_worker_watchdogs(deadline, rss, expected):
    class Process:
        pid = 123
        alive = True

        def is_alive(self):
            return self.alive

        def terminate(self):
            self.alive = False

        def join(self, timeout=None):
            pass

    class Control:
        def poll(self):
            return True

        def recv(self):
            return {'kind': 'progress', 'value': {'phase': 'scan', 'completed': 0, 'total': 1}}

    events = []
    process = Process()
    with pytest.raises(ValueError, match=expected):
        _monitor_worker(process, deadline, 512, rss_reader=lambda pid: rss,
                        control=Control(), progress_callback=events.append)
    assert not process.is_alive() and len(events) == 128


def test_http_import_and_batch_jobs_expose_truthful_phase_counts(client):
    response = client.post('/api/import', files={'file': ('source.xlsx', make_workbook())})
    imported = poll(client, response)
    completed = client.get('/api/jobs/' + response.json()['id']).json()
    assert completed['progress'] == {'phase': 'index'}
    url = '/api/imports/' + imported['import_id']
    mapping = client.get(url + '/sheet', params={'name': 'parallel'}).json()['mapping']
    response = client.post(url + '/confirm-all', json={'mappings': [mapping]})
    poll(client, response)
    completed = client.get('/api/jobs/' + response.json()['id']).json()
    assert completed['progress'] == {'phase': 'batch', 'completed': 1, 'total': 1}
