from copy import deepcopy
from datetime import datetime
from pathlib import Path
import json
import re
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from structured_data.scanner import scan_workbook, complete_date
from structured_data.generic_parser import parse_mapped_workbook, validate_mapping


def book(tmp_path,rows,name='prices'):
    p=tmp_path/'sample.xlsx';w=Workbook();s=w.active;s.title=name
    for r in rows:s.append(r)
    w.save(p);return p


def candidate(p):return scan_workbook(p)['sheets'][0]


def test_simple_auto_parse_provenance_blank_error(tmp_path):
    p=book(tmp_path,[['日期','价格','库存'],[datetime(2023,9,22),17110,None],[datetime(2023,9,23),'#N/A',3]])
    s=candidate(p);assert s['status']=='auto_ready'
    assert s['preview'][1][0]=={'cell':'A2','value':'2023-09-22T00:00:00'}
    parsed=parse_mapped_workbook(p,[s['mapping']]);obs=parsed['observations']
    assert [(o['source_cell'],o['value_status']) for o in obs]==[('B2','valid'),('C2','blank'),('B3','excel_error'),('C3','valid')]
    assert obs[0]['value']==17110 and obs[0]['date_cell']=='A2'
    json.dumps(parsed,ensure_ascii=False)


def test_unit_row_and_multi_header_require_review(tmp_path):
    p=book(tmp_path,[['日期','价格'],[None,'元/吨'],[datetime(2023,1,1),1],[datetime(2023,1,2),2]])
    s=candidate(p);assert s['status']=='auto_ready';assert s['mapping']['blocks'][0]['columns'][0]['unit']=='元/吨'
    w=load_workbook(p);w.active['B2']='主力';w.save(p)
    assert candidate(p)['status']=='needs_review'


def test_date_modes_do_not_guess(tmp_path):
    for v in ['2023','2023-09','2023/24',2023,45200,datetime(2023,9,1,12),True,'09/22/2023']:
        assert complete_date(v) is None
    assert complete_date('2023年9月22日')=='2023-09-22'
    assert complete_date('2023-02-30') is None
    p=book(tmp_path,[['年月','价格'],['2023-09',1],['2023-10',2]])
    s=candidate(p);assert s['status']=='needs_review';assert not s['mapping']['blocks']


def test_late_nondate_blocks_entire_import(tmp_path):
    rows=[['日期','价格']]+[[datetime(2023,1,d),d] for d in range(1,29)]
    rows += [[datetime(2023,2,d),d] for d in range(1,20)]
    rows += [['合计',999]]
    p=book(tmp_path,rows);s=candidate(p);assert s['status']=='auto_ready'
    with pytest.raises(ValueError,match='阻断整页'):parse_mapped_workbook(p,[s['mapping']])
    m=deepcopy(s['mapping']);m['blocks'][0]['data_end_row']=len(rows)-1
    assert len(parse_mapped_workbook(p,[m])['observations'])==47


def test_duplicate_dates_block_import(tmp_path):
    p=book(tmp_path,[['日期','价格'],[datetime(2023,1,1),1],[datetime(2023,1,1),2]])
    with pytest.raises(ValueError,match='重复日期'):parse_mapped_workbook(p,[candidate(p)['mapping']])


def test_parallel_blocks_explicit_mapping(tmp_path):
    p=book(tmp_path,[['日期','收盘',None,'日期','库存'],[datetime(2023,1,1),1,None,datetime(2023,2,1),11],[datetime(2023,1,2),2,None,datetime(2023,2,2),12]])
    s=candidate(p);assert s['status']=='auto_ready';assert len(s['mapping']['blocks'])==2
    result=parse_mapped_workbook(p,[s['mapping']]);assert len(result['observations'])==4
    assert result['observations'][1]['date_cell']=='D2';assert result['observations'][1]['date']=='2023-02-01'


def test_fingerprint_reuses_new_dates_but_rejects_header_color_merge(tmp_path):
    p=book(tmp_path,[['日期','价格'],[datetime(2023,1,1),1],[datetime(2023,1,2),2]])
    first=candidate(p);w=load_workbook(p);w.active['A2']=datetime(2024,1,1);w.active['A3']=datetime(2024,1,2);w.save(p)
    assert candidate(p)['fingerprint']==first['fingerprint']
    parse_mapped_workbook(p,[first['mapping']])
    w.active['B1'].fill=PatternFill('solid',fgColor='FF0000');w.save(p)
    with pytest.raises(ValueError,match='指纹'):parse_mapped_workbook(p,[first['mapping']])
    s=candidate(p);w.active.merge_cells('A1:B1');w.save(p)
    assert candidate(p)['fingerprint']!=s['fingerprint']


def test_formula_cache_and_missing_cache(tmp_path):
    p=book(tmp_path,[['日期','价格'],[datetime(2023,1,1),'=1+2'],[datetime(2023,1,2),'=3+4']])
    with ZipFile(p) as z:parts={n:z.read(n) for n in z.namelist()}
    xml=re.sub(r'<f>1\+2</f><v(?:\s*/>|></v>)', '<f>1+2</f><v>3</v>', parts['xl/worksheets/sheet1.xml'].decode())
    parts['xl/worksheets/sheet1.xml']=xml.encode()
    with ZipFile(p,'w') as z:
        for n,v in parts.items():z.writestr(n,v)
    result=parse_mapped_workbook(p,[candidate(p)['mapping']]);assert result['observations'][0]['value']==3
    assert result['observations'][0]['formula']=='=1+2'
    assert result['observations'][1]['value_status']=='formula_no_cache'
    assert any(i['code']=='cached_formula_values' for i in result['issues'])


def test_explicit_serial_respects_epoch_and_rejects_year(tmp_path):
    p=book(tmp_path,[['日期','价格'],[45000,1],[45001,2]])
    w=load_workbook(p);w.epoch=CALENDAR_MAC_1904;w.save(p)
    s=candidate(p);assert s['status']=='needs_review'
    m={'sheet':'prices','fingerprint':s['fingerprint'],'blocks':[{'date_column':'A','date_mode':'excel_serial','header_rows':[1],'data_start_row':2,'columns':[{'column':'B','metric':'价格'}]}]}
    result=parse_mapped_workbook(p,[m]);assert result['observations'][0]['date']=='2027-03-16'
    w.active['A2']=2023;w.save(p);m['fingerprint']=candidate(p)['fingerprint']
    with pytest.raises(ValueError,match='日期无效'):parse_mapped_workbook(p,[m])


def test_chart_only_not_empty_and_empty_is_empty(tmp_path):
    p=book(tmp_path,[]);w=load_workbook(p);chart=BarChart();chart.add_data(Reference(w.active,min_col=1,min_row=1,max_row=3));w.active.add_chart(chart,'C1');w.save(p)
    s=candidate(p);assert s['status']=='needs_review' and s['charts']['chart_count']==1
    with pytest.raises(ValueError,match='数据区块'):validate_mapping(s['mapping'],s)
    q=tmp_path/'empty.xlsx';Workbook().save(q);assert candidate(q)['status']=='empty'


def test_mapping_validation_no_empty_metric_or_overlap(tmp_path):
    p=book(tmp_path,[['日期','价格'],[datetime(2023,1,1),1],[datetime(2023,1,2),2]])
    s=candidate(p);m=deepcopy(s['mapping']);m['blocks'][0]['columns'][0]['metric']=' '
    with pytest.raises(ValueError,match='不能为空'):validate_mapping(m,s)
    m=deepcopy(s['mapping']);m['blocks'].append(deepcopy(m['blocks'][0]))
    with pytest.raises(ValueError,match='重叠'):validate_mapping(m,s)
    m=deepcopy(s['mapping']);m['blocks'][0]['date_mode']='year'
    with pytest.raises(ValueError,match='年度'):validate_mapping(m,s)


def test_sparse_group_titles_inherit_inside_date_block_only(tmp_path):
    p=book(tmp_path,[['日期','收盘价',None,None,'结算价',None],
        [None,'close',None,None,'settle',None],
        ['日期',0,0,'日期',0,0],['Date','CF','CF09','Date','CF','CF09'],
        [datetime(2023,1,1),1,2,datetime(2023,1,1),3,4],
        [datetime(2023,1,2),1,2,datetime(2023,1,2),3,4]])
    s=candidate(p);m=s['mapping'];assert s['status']=='needs_review'
    assert m['blocks'][0]['columns'][1]['metric']=='收盘价 / close / CF09'
    assert m['blocks'][0]['columns'][1]['header_cells']==['B1','B2','C3','C4']
    assert m['blocks'][1]['columns'][1]['metric']=='结算价 / settle / CF09'
    assert 'E1' in m['blocks'][1]['columns'][1]['header_cells']
    assert len(parse_mapped_workbook(p,[m])['observations'])==8


def test_same_source_metric_retains_column_identity_or_explicit_business_dimensions(tmp_path):
    p=book(tmp_path,[['日期','产量','产量'],[datetime(2023,1,1),1,2],[datetime(2023,1,2),3,4]])
    s=candidate(p);m=s['mapping'];assert s['status']=='needs_review'
    normalized=validate_mapping(m,s)
    assert [c['dimensions'] for c in normalized['blocks'][0]['columns']]==[{'原表列':'B'},{'原表列':'C'}]
    assert all('年度' not in c['dimensions'] for c in normalized['blocks'][0]['columns'])
    m['blocks'][0]['columns'][0]['dimensions']={'年度':'2022/23'}
    m['blocks'][0]['columns'][1]['dimensions']={'年度':'2023/24'}
    assert len(validate_mapping(m,s)['blocks'][0]['columns'])==2


@pytest.mark.parametrize('change,expected', [
    ({'header_rows': []}, ('表头行为空', '填写原表的表头行号')),
    ({'header_rows': [2]}, ('数据从第 2 行开始', '数据起始行移到表头之后')),
    ({'data_end_row': 99}, ('99', '第 2 至 3 行', '修改结束行')),
])
def test_review_errors_identify_block_and_repair(tmp_path, change, expected):
    p=book(tmp_path,[['日期','价格'],[datetime(2023,1,1),1],[datetime(2023,1,2),2]])
    sheet=candidate(p)
    mapping=deepcopy(sheet['mapping'])
    mapping['blocks'][0].update(name='日度记录', **change)
    with pytest.raises(ValueError) as caught:
        validate_mapping(mapping,sheet)
    message=str(caught.value)
    assert f'工作表「{sheet["name"]}」' in message
    assert '区块 1（日度记录）' in message
    for text in expected:
        assert text in message


def test_empty_metric_identifies_real_column(tmp_path):
    p=book(tmp_path,[['日期','价格'],[datetime(2023,1,1),1],[datetime(2023,1,2),2]])
    sheet=candidate(p)
    mapping=deepcopy(sheet['mapping'])
    mapping['blocks'][0]['columns'][0]['metric']=''
    with pytest.raises(ValueError) as caught:
        validate_mapping(mapping,sheet)
    assert 'B 列的指标名称不能为空' in str(caught.value)
    assert '对照左侧原表表头' in str(caught.value)


def test_invalid_date_locates_date_and_populated_value_cells(tmp_path):
    from structured_data.review_errors import ReviewValidationError
    p=book(tmp_path,[['日期','价格'],[datetime(2023,1,1),1],[datetime(2023,1,2),2]])
    sheet=candidate(p)
    wb=load_workbook(p);wb.active['A3']='not a date';wb.save(p)
    # Rescan keeps an explicit user mapping with a current source fingerprint.
    scan=scan_workbook(p)
    mapping=sheet['mapping'];mapping['fingerprint']=scan['sheets'][0]['fingerprint']
    with pytest.raises(ReviewValidationError) as caught:
        parse_mapped_workbook(p,[mapping],scan=scan)
    assert caught.value.review_target=={'sheet':sheet['name'],'block_index':0,'field':'date_column','cells':['A3','B3']}


def test_verified_metadata_and_source_zero_fill_warning_survive_generic_import(tmp_path):
    p=book(tmp_path,[['日期','价格'],[datetime(2023,1,1),'=WSD("TEST","close","2023-01-01","ShowBlank==0")'],[datetime(2023,1,2),10]])
    sheet=candidate(p)
    mapping=deepcopy(sheet['mapping'])
    mapping['blocks'][0]['columns'][0]['dimensions']={'contract':'合同甲','block':'close'}
    result=parse_mapped_workbook(p,[mapping])
    assert result['series'][0]['dimensions']['block']=='close'
    assert result['series'][0]['dimensions']['contract']=='合同甲'
    assert any(issue['code']=='source_zero_fill' and issue['cell']=='B2' for issue in result['issues'])
    assert result['observations'][0]['formula']=='=WSD("TEST","close","2023-01-01","ShowBlank==0")'
    assert result['observations'][0]['value'] is None
