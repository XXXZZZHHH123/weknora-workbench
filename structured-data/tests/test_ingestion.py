"""Native-upload continuation: signature, immutable bytes and authoritative source checks."""
import hashlib
import json
import time

import pytest
from structured_data import ingestion
from test_native_knowledge_import import native, workbook
from test_pdf_tables import bundle


def claim_for(data, envelope, kind='xlsx'):
    return {'attempt': 'attempt-a', 'tenant_id': 10000, 'knowledge_base_id': 'kb-a',
            'knowledge_id': 'source-a', 'user_id': 'user-a', 'file_name': 'source.'+kind,
            'file_type': kind, 'file_hash': hashlib.md5(data).hexdigest(), 'file_size': len(data),
            'source_sha256': hashlib.sha256(data).hexdigest(), 'bundle_sha256': hashlib.sha256(envelope).hexdigest()}


@pytest.fixture
def setup(native, monkeypatch):
    client, app, state, root = native
    key = root / 'ingestion-key'
    key.write_bytes(b'isolated-ingestion-key-01234567890123456789')
    monkeypatch.setenv('STRUCTURED_SERVICE_TOKEN_FILE', str(key))
    checks = []
    monkeypatch.setattr(ingestion, 'check_task_write', lambda claim: checks.append(dict(claim)))
    return client, root, checks


def test_continuation_cannot_use_read_tool_route_or_mutate_unsigned_source(setup):
    client, _, checks = setup
    data, envelope = workbook(), b'{}'
    claim = claim_for(data, envelope)
    headers = ingestion.signed_headers(claim, 'POST', ingestion.UPLOAD_PATH)
    assert client.post(ingestion.UPLOAD_PATH, files={'file': ('source.xlsx', data)}, data={'bundle': '{}'}).status_code == 401
    tampered = ingestion.signed_headers({**claim, 'knowledge_base_id': 'kb-b'}, 'POST', ingestion.UPLOAD_PATH)
    tampered['X-Ingestion-Signature'] = headers['X-Ingestion-Signature']
    assert client.post(ingestion.UPLOAD_PATH, files={'file': ('source.xlsx', data)}, data={'bundle': '{}'}, headers=tampered).status_code == 401
    assert client.post(ingestion.UPLOAD_PATH, files={'file': ('source.xlsx', data)}, data={'bundle': '{}'},
        headers=ingestion.signed_headers(claim, 'POST', ingestion.UPLOAD_PATH, now=time.time()-120)).status_code == 401
    assert not checks
    assert client.post(ingestion.UPLOAD_PATH, files={'file': ('source.xlsx', data+b'x')}, data={'bundle': '{}'}, headers=headers).status_code == 413
    assert client.post(ingestion.UPLOAD_PATH, files={'file': ('source.xlsx', data)}, data={'bundle': '{"extra":true}'}, headers=headers).status_code == 409


def finish(client, claim):
    until = time.monotonic()+20
    while time.monotonic()<until:
        response = client.get(ingestion.RESULT_PATH, headers=ingestion.signed_headers(claim, 'GET', ingestion.RESULT_PATH))
        assert response.status_code == 200, response.text
        result = response.json()
        if result['status'] != 'processing':
            return result
        time.sleep(.02)
    raise AssertionError('isolated continuation did not finish')


def test_native_excel_continuation_publishes_and_reuses_same_attempt(setup):
    client, root, checks = setup
    data, envelope = workbook(), b'{}'
    claim = claim_for(data, envelope)
    response = client.post(ingestion.UPLOAD_PATH, files={'file': ('source.xlsx', data)}, data={'bundle': '{}'},
        headers=ingestion.signed_headers(claim, 'POST', ingestion.UPLOAD_PATH))
    assert response.status_code == 200, response.text
    result = finish(client, claim)
    assert result['status'] == 'completed', result
    assert result['result']['attention_count'] == 0
    assert result['result']['sheets'] == [{'name':'records','status':'auto_ready'}]
    datasets = list((root/'knowledge-bases'/'kb-a'/'datasets').glob('*/manifest.json'))
    assert len(datasets) == 1
    again = client.post(ingestion.UPLOAD_PATH, files={'file': ('source.xlsx', data)}, data={'bundle': '{}'},
        headers=ingestion.signed_headers(claim, 'POST', ingestion.UPLOAD_PATH))
    assert again.status_code == 200 and again.json()['reused'] is True
    assert len(list((root/'knowledge-bases'/'kb-a'/'datasets').glob('*/manifest.json'))) == 1
    assert len(checks)>2, 'publication must recheck the admitted native source'


def test_changed_source_blocks_even_previously_completed_result(setup, monkeypatch):
    client, _, _ = setup
    data=workbook();claim=claim_for(data,b'{}')
    monkeypatch.setattr(ingestion,'check_task_write',lambda claim: (_ for _ in ()).throw(ValueError('source changed')))
    response=client.post(ingestion.UPLOAD_PATH,files={'file':('source.xlsx',data)},data={'bundle':'{}'},headers=ingestion.signed_headers(claim,'POST',ingestion.UPLOAD_PATH))
    assert response.status_code == 400 and response.json()['detail'] == 'source changed'


def test_pdf_uses_captured_parse_once_and_keeps_source_cells(setup, monkeypatch):
    from structured_data.pdf_tables import PDFTableStore
    client, root, checks = setup
    data = b'%PDF-1.7 isolated captured-source fixture'
    raw = bundle()
    raw['source']['sha256'] = hashlib.sha256(data).hexdigest()
    envelope = json.dumps({'engine':'remote_vlm','bundle':raw}).encode()
    claim = claim_for(data, envelope, 'pdf')
    async def forbidden(*a, **kw):
        raise AssertionError('captured PDF must not be parsed a second time')
    monkeypatch.setattr('structured_data.pdf_parser_bridge.parse_pdf', forbidden)
    monkeypatch.setattr('structured_data.pdf_automation.group_fragments', lambda *a: [
        {'fragment_ids':['table-1'],'reason':'one actual table'}])
    def suggest(draft, body, directory):
        return {'status':'ready','plan':{
            'bundle_sha256':draft['bundle_sha256'], 'fragment_ids':['table-1'],
            'header_rows':{'table-1':[{'row':0,'reason':'printed header'}]},
            'columns':[{'name':'item','type':'text','source_columns':{'table-1':0}},
                       {'name':'value','type':'decimal','source_columns':{'table-1':1}}]}}
    monkeypatch.setattr('structured_data.pdf_api.suggest_plan', suggest)
    response = client.post(ingestion.UPLOAD_PATH, files={'file':('source.pdf',data)},
        data={'bundle':envelope.decode()}, headers=ingestion.signed_headers(claim,'POST',ingestion.UPLOAD_PATH))
    assert response.status_code == 200, response.text
    result = finish(client, claim)
    assert result['status'] == 'completed', result
    assert result['result']['accepted_count'] == 1
    store = PDFTableStore(root/'knowledge-bases'/'kb-a'/'pdf-tables')
    table = store.read_table(store.list_tables()['tables'][0]['table_id'])
    assert table['rows'][0]['raw_values']['value'] == '-1.20'
    cell = table['rows'][0]['source_cells']['value'][0]
    assert cell['source_sha256'] == claim['source_sha256'] and cell['page'] == 1
    assert len(checks) > 2


@pytest.mark.parametrize('engine', ['mineru_v1', 'remote_vlm'])
def test_pdf_without_tables_still_completes_native_document(setup, engine):
    client, root, _ = setup
    data = b'%PDF-1.7 text-only isolated fixture'
    raw = bundle()
    raw['source']['sha256'] = hashlib.sha256(data).hexdigest()
    raw['fragments'] = []
    captured = ({'engine':engine,'bundle':raw} if engine == 'remote_vlm' else
                {'engine':engine,'response':{'structured':[{'type':'text','page_idx':0,'text':'Narrative source text'}]}})
    envelope = json.dumps(captured).encode()
    claim = claim_for(data, envelope, 'pdf')
    response = client.post(ingestion.UPLOAD_PATH, files={'file':('source.pdf',data)},
        data={'bundle':envelope.decode()}, headers=ingestion.signed_headers(claim,'POST',ingestion.UPLOAD_PATH))
    assert response.status_code == 200, response.text
    result = finish(client, claim)
    assert result['status'] == 'completed', result
    assert result['result']['table_count'] == 0
    assert result['result']['attention_count'] == 0
    assert not list(root.glob('**/drafts/*.json'))


def test_native_excel_followup_publishes_with_one_source_guard(setup):
    from io import BytesIO
    from openpyxl import load_workbook
    client, root, checks = setup
    book=load_workbook(BytesIO(workbook()));book.active['B1']='mass (kg)'
    # A drawing keeps a dated numeric block pending at scan time; follow-up
    # must confirm the source-backed mapping without duplicating its guard.
    from openpyxl.chart import BarChart, Reference
    chart=BarChart();chart.add_data(Reference(book.active,min_col=2,min_row=1,max_row=3),titles_from_data=True)
    book.active.add_chart(chart,'E1')
    stream=BytesIO();book.save(stream);data=stream.getvalue()
    claim=claim_for(data,b'{}')
    response=client.post(ingestion.UPLOAD_PATH,files={'file':('source.xlsx',data)},data={'bundle':'{}'},
        headers=ingestion.signed_headers(claim,'POST',ingestion.UPLOAD_PATH))
    assert response.status_code==200,response.text
    result=finish(client,claim)
    assert result['status']=='completed',result
    assert result['result']['attention_count']==0,result
    assert result['result']['dataset_id'],result
    assert result['result']['sheets']==[{'name':'records','status':'confirmed'}],result
    review=result['result']['automatic_review']
    assert review['sheets_imported']==['records'] and not review['sheets_failed'],review
    assert len(checks)>3
