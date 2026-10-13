"""Mixed table/chart sheets must preserve real headers through follow-up queries."""
import hashlib
import json

import pytest
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

from structured_data.field_resolver import FieldResolver
from structured_data.store import DataStore


@pytest.fixture
def mixed_store(tmp_path):
    path=tmp_path/'mixed.xlsx'
    wb=Workbook();ws=wb.active;ws.title='新疆棉入库和公检量'
    ws.append(['日期','当日公检量','累计公检量','当日入库量','在库量'])
    ws.append(['2022-11-11',10,100,20,30])
    ws.append(['2022-11-12',15,115,25,40])
    ws.append(['2022-11-13',None,115,None,None])
    chart=LineChart();chart.title='新疆棉公检量'
    chart.add_data(Reference(ws,min_col=2,max_col=3,min_row=1,max_row=4),titles_from_data=True)
    chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=4));ws.add_chart(chart,'G1')
    wb.save(path);wb.close()
    series=[];observations=[]
    for col,metric in [('B','当日公检量'),('C','累计公检量'),('D','当日入库量'),('E','在库量')]:
        series.append(dict(series_id=col,sheet=ws.title,metric=metric,unit=None,
                           dimensions={},header_cells=[col+'1'],date_mode='date'))
        for row,day in [(2,'2022-11-11'),(3,'2022-11-12'),(4,'2022-11-13')]:
            value=ws[f'{col}{row}'].value
            observations.append(dict(series_id=col,date=day,value=value,raw_value=value,
                value_status='valid' if value is not None else 'blank',source_cell=f'{col}{row}',date_cell=f'A{row}'))
    series.append(dict(series_id='other',sheet='其他表',metric='期末结存',dimensions={},date_mode='date'))
    store=DataStore(tmp_path/'data')
    ds=store.publish(dict(source_file=str(path),sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        rule_version='test',series=series,observations=observations,sheets=[],excluded_regions=[],issues=[]))['dataset_id']
    return store,ds,ws.title


def test_sheet_mention_does_not_hide_table_fields_or_restrict_other_sheets(mixed_store):
    store,ds,sheet=mixed_store
    question=f'在库量哪天最多？有多少货？（延续上文，指{sheet}相关数据中的在库量）'
    full=store.catalog(ds,question=question); scoped=store.catalog(ds,sheet=sheet,question=question)
    assert {s['series_id'] for s in full if not s.get('chart_id')}=={'B','C','D','E','other'}
    assert {s['series_id'] for s in scoped if not s.get('chart_id')}=={'B','C','D','E'}
    assert any(s.get('chart_id') for s in full) and any(s.get('chart_id') for s in scoped)
    assert store.catalog(ds,question='不要图表，只按原值比较其他表')==full


@pytest.mark.parametrize('explicit_sheet',[False,True])
@pytest.mark.parametrize('with_title',[False,True])
def test_explicit_chart_query_still_uses_only_linked_curve_ranges(mixed_store,explicit_sheet,with_title):
    store,ds,sheet=mixed_store
    views=store.catalog(ds,sheet=sheet if explicit_sheet else None,
                       question=f'{sheet}图中'+('新疆棉公检量' if with_title else '')+'哪天最多')
    assert {s['series_id'] for s in views if not s.get('chart_id')}==({'B','C','D','E'} if explicit_sheet else {'B','C','D','E','other'})
    charts=[s for s in views if s.get('chart_id')]
    assert {s['source_metric'] for s in charts}=={'当日公检量','累计公检量'}
    assert all(s['source_series_id']!='E' for s in charts)
    assert all('source_series_id' in s and s['chart_ranges'] for s in charts)


def test_followup_switches_to_real_inventory_header_and_global_max(mixed_store):
    store,ds,sheet=mixed_store
    def model(p):
        cols=[c for g in p['catalog'] for c in g['columns']]
        metric='在库量' if '在库量' in p['question'] else '累计公检量' if '累计公检量' in p['question'] else '当日入库量'
        matches=[c for c in cols if c['name']==metric]
        assert len(matches)==1
        return dict(intent='lookup',decision='resolved',fields=[dict(role='v1',label=metric,candidate_ids=[matches[0]['id']])],
                    query={'operation':'max' if metric=='在库量' else 'rows'},time_scope=({'kind':'unbounded','evidence':[],'explanation':'全历史极值'} if metric=='在库量' else {'kind':'dates','dates':['2022-11-12'],'evidence':['2022/11/12'],'explanation':'所问日期'}))
    resolver=FieldResolver(store,model)
    first=resolver.answer('2022/11/12当日入库量是多少',dataset_id=ds)
    assert first['answer']['sources']['v1']['rows'][0]['value']==25
    second=resolver.answer('2022/11/12累计公检量是多少',dataset_id=ds,context_sheet=sheet,prior_series_ids=['D'])
    assert second['answer']['sources']['v1']['rows'][0]['value']==115
    result=resolver.answer(f'在库量哪天最多？有多少货？（延续上文，指{sheet}相关数据中的在库量）',
        dataset_id=ds,context_sheet=sheet,prior_series_ids=['C'],date='2022-11-13')
    assert result['status']=='resolved' and result['date'] is None
    source=result['answer']['sources']['v1']
    assert source['aggregate']==40 and source['missing_count']==1
    assert source['extremum_rows'][0]['date']=='2022-11-12'
    assert source['extremum_rows'][0]['source_cell']=='E3'
    assert source['extremum_rows'][0]['date_cell']=='A3'


def test_out_of_catalogue_candidates_are_repaired_before_asking_user(mixed_store):
    store,ds,_=mixed_store;calls=[]
    def model(p):
        calls.append(p)
        cols={c['name']:c['id'] for g in p['catalog'] for c in g['columns']}
        ids=['not-a-real-id'] if len(calls)==1 else [cols['在库量']]
        return dict(intent='lookup',decision='clarify' if len(calls)==1 else 'resolved',
                    fields=[dict(role='v1',label='在库量',candidate_ids=ids)],query={'operation':'max'})
    result=FieldResolver(store,model).answer('请核对原表：在库量哪天最多？有多少货？',dataset_id=ds)
    assert result['status']=='resolved' and len(calls)==2
    assert '实际候选目录' in calls[1]['repair_instruction']
    assert result['answer']['sources']['v1']['aggregate']==40


def test_repeated_invalid_identifier_stops_without_business_clarification(mixed_store):
    store,ds,_=mixed_store
    def model(p):
        cols={c['name']:c['id'] for g in p['catalog'] for c in g['columns']}
        return dict(intent='lookup',decision='clarify',fields=[dict(role='v1',label='在库量',
                    candidate_ids=['not-a-real-id'])],query={'operation':'max'})
    result=FieldResolver(store,model).tool_answer('请核对原表：在库量哪天最多？',dataset_id=ds)
    assert result['kind']=='structured_tool_notice' and result['status']=='query_unavailable'
    assert 'answer' not in result


def test_missing_named_header_expands_to_full_catalog_before_clarifying(mixed_store,monkeypatch):
    from structured_data.field_index import FieldIndex
    store,ds,_=mixed_store
    path=store._path(ds)/'series.json';catalog=json.loads(path.read_text())
    catalog.extend(dict(series_id=f'noise-{i}',sheet='其他表',metric=f'别的指标{i}') for i in range(30))
    path.write_text(json.dumps(catalog))
    def missed_search(self,*args,**kwargs):
        return [s for s in self.catalog if s['series_id'] in ('B','C')],{'method':'test_missed_recall'}
    monkeypatch.setattr(FieldIndex,'search',missed_search)
    calls=[]
    def model(p):
        calls.append(p)
        cols={c['name']:c['id'] for g in p['catalog'] for c in g['columns']}
        sid=cols.get('在库量',cols['累计公检量'])
        return dict(intent='lookup',decision='resolved',fields=[dict(role='v1',label='在库量',candidate_ids=[sid])],query={'operation':'max'})
    resolver=FieldResolver(store,model,lambda p:{'search_terms':['库存']})
    result=resolver.answer('请核对原表：在库量哪天最多？',dataset_id=ds)
    assert len(calls)==1 and result['retrieval']['full_catalogue_checked']
    assert any(c['name']=='在库量' for g in calls[0]['catalog'] for c in g['columns']), 'a weak literal recall cannot hide actual fields from model understanding'
    assert result['status']=='resolved' and result['answer']['sources']['v1']['aggregate']==40


def test_same_header_with_real_dimension_ambiguity_still_requires_selection(mixed_store):
    store,ds,_=mixed_store
    path=store._path(ds)/'series.json';catalog=json.loads(path.read_text())
    catalog.append(dict(series_id='another-inventory',sheet='另一表',metric='在库量',dimensions={'产地':'另一产地'},date_mode='date'))
    path.write_text(json.dumps(catalog))
    def model(p):
        ids=[c['id'] for g in p['catalog'] for c in g['columns'] if c['name']=='在库量']
        return dict(intent='lookup',decision='clarify',fields=[dict(role='v1',label='在库量',candidate_ids=ids)],query={'operation':'max'})
    result=FieldResolver(store,model).answer('在库量哪天最多？',dataset_id=ds)
    assert result['status']=='needs_clarification' and 'answer' not in result
    assert len(result['fields'][0]['candidates'])==2
