import hashlib
import json
from copy import deepcopy
from datetime import datetime

import pytest
from openpyxl import Workbook

from structured_data.imports import ImportManager
from structured_data.streaming import DiskRows, read_json_array, write_json_array


def source_file(tmp_path, name='业务.xlsx'):
    path = tmp_path / name
    wb = Workbook()
    ws = wb.active
    ws.title = '已导入'
    ws.append(['日期', '库存'])
    ws.append([datetime(2023, 1, 1), 100])
    ws.append([datetime(2023, 1, 2), '=NA()'])
    ws = wb.create_sheet('新增表')
    ws.append(['日期', '价格'])
    ws.append([None, '品种甲'])
    ws.append([datetime(2023, 1, 1), 10])
    ws.append([datetime(2023, 1, 2), 12])
    wb.save(path)
    return path


def test_confirmation_only_parses_changed_sheet_and_preserves_all_sources(tmp_path, monkeypatch):
    from structured_data import generic_parser
    source = source_file(tmp_path)
    manager = ImportManager(tmp_path / 'data')
    first = manager.ingest(source)
    old = first['dataset_id']
    stock = manager.store.list_series(old)['series'][0]['series_id']
    old_rows = manager.store.query_series(old, stock)['rows']
    old_hash = hashlib.sha256((manager.store._path(old) / 'data.duckdb').read_bytes()).hexdigest()
    calls = []
    original = generic_parser.parse_mapped_workbook

    def tracked(path, mappings, **kwargs):
        calls.extend(m['sheet'] for m in mappings)
        assert isinstance(kwargs['observations'], DiskRows)
        return original(path, mappings, **kwargs)

    monkeypatch.setattr(generic_parser, 'parse_mapped_workbook', tracked)
    mapping = manager.sheet(first['import_id'], '新增表')['mapping']
    result = manager.confirm(first['import_id'], mapping)
    new = result['dataset_id']
    assert set(calls) == {'新增表'}
    assert manager.store.manifest(new)['changed_sheets'] == ['新增表']
    assert manager.store.manifest(new)['parent_dataset_id'] == old
    assert manager.store.manifest(new)['observation_count'] == 4
    assert manager.store.query_series(new, stock)['rows'] == old_rows
    assert old_rows[1]['formula'] == '=NA()' and old_rows[1]['value'] is None
    assert old_rows[1]['source_cell'] == 'B3' and old_rows[1]['date_cell'] == 'A3'
    assert hashlib.sha256((manager.store._path(old) / 'data.duckdb').read_bytes()).hexdigest() == old_hash
    assert (manager.store._path(new) / 'source.xlsx').read_bytes() == source.read_bytes()
    assert not list(manager.store.datasets.glob('.import-*'))
    # Rechecking identical mappings (including browser block IDs) creates no new version.
    mapping['blocks'][0]['_review_id'] = 'another-browser-session'
    repeat = manager.confirm(first['import_id'], mapping)
    assert repeat['dataset_id'] == new
    assert len(manager.store.list_datasets()) == 2


def test_changed_mapping_replaces_rows_instead_of_appending_and_failure_is_atomic(tmp_path, monkeypatch):
    from structured_data import generic_parser
    manager = ImportManager(tmp_path / 'data')
    first = manager.ingest(source_file(tmp_path))
    mapping = manager.sheet(first['import_id'], '新增表')['mapping']
    second = manager.confirm(first['import_id'], mapping)
    mapping['blocks'][0]['columns'][0]['metric'] = '更正指标'
    third = manager.confirm(first['import_id'], mapping)
    assert manager.store.manifest(third['dataset_id'])['observation_count'] == 4
    assert manager.store.manifest(second['dataset_id'])['observation_count'] == 4
    before = len(manager.store.list_datasets())
    original = generic_parser.parse_mapped_workbook
    calls = 0

    def fail_after_preview(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise ValueError('模拟写入前失败')
        return original(*args, **kwargs)

    monkeypatch.setattr(generic_parser, 'parse_mapped_workbook', fail_after_preview)
    with pytest.raises(ValueError, match='模拟'):
        manager.confirm(first['import_id'], mapping)
    assert manager.get(first['import_id'])['dataset_id'] == third['dataset_id']
    assert len(manager.store.list_datasets()) == before
    assert not list(manager.store.datasets.glob('.import-*'))


def test_consecutive_block_confirmation_preserves_earlier_values_and_versions(tmp_path):
    """The editor submits all selected blocks, including already confirmed ones."""
    path = tmp_path / 'three-independent-blocks.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '业务记录'
    sheet.append(['日期', '计数甲', '计数乙', '计数丙'])
    sheet.append([datetime(2024, 2, 1), 0, 21, 31])
    sheet.append([datetime(2024, 2, 2), 12, 22, 32])
    workbook.save(path)
    manager = ImportManager(tmp_path / 'data')
    scanned = manager.ingest(path, use_reviewed_mappings=False)
    candidate = manager.sheet(scanned['import_id'], sheet.title)['mapping']
    base = candidate['blocks'][0]
    assert len(base['columns']) == 3
    blocks = [{**deepcopy(base), 'name': f'区块 {index + 1}', 'columns': [deepcopy(column)]}
              for index, column in enumerate(base['columns'])]
    versions = []
    for size in (1, 2, 3):
        mapping = {**candidate, 'blocks': deepcopy(blocks[:size])}
        result = manager.confirm(scanned['import_id'], mapping, save_template=False)
        dataset = result['dataset_id']
        versions.append(dataset)
        catalog = manager.store.list_series(dataset)['series']
        assert len(catalog) == size
        assert manager.store.manifest(dataset)['observation_count'] == size * 2
        values = {entry['metric']: manager.store.query_series(dataset, entry['series_id'])['rows']
                  for entry in catalog}
        assert [row['value'] for row in values['计数甲']] == [0, 12]
        assert [row['source_cell'] for row in values['计数甲']] == ['B2', 'B3']
        if size > 1:
            assert [row['value'] for row in values['计数乙']] == [21, 22]
            assert [row['source_cell'] for row in values['计数乙']] == ['C2', 'C3']
    assert len(set(versions)) == 3
    assert [manager.store.manifest(dataset)['observation_count'] for dataset in versions] == [2, 4, 6]
    final_mapping = manager.sheet(scanned['import_id'], sheet.title)['mapping']
    assert len(final_mapping['blocks']) == 3
    assert manager.confirm(scanned['import_id'], final_mapping, save_template=False)['dataset_id'] == versions[-1]


def test_stale_review_pointer_preserves_newer_published_tables(tmp_path):
    from structured_data.imports import atomic_json
    path = source_file(tmp_path)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb.create_sheet('另一张待核对表')
    ws.append(['日期', '进口']); ws.append([None, '另一口径'])
    ws.append([datetime(2023, 1, 1), 30]); ws.append([datetime(2023, 1, 2), 40]); wb.save(path)
    manager = ImportManager(tmp_path / 'data')
    first = manager.ingest(path)
    stale_doc = manager.get(first['import_id'])
    mapping = manager.sheet(first['import_id'], '新增表')['mapping']
    published = manager.confirm(first['import_id'], mapping)
    atomic_json(manager.path(first['import_id']) / 'review.json', stale_doc)
    other = manager.sheet(first['import_id'], '另一张待核对表')['mapping']
    result = manager.confirm(first['import_id'], other)
    manifest = manager.store.manifest(result['dataset_id'])
    assert manifest['parent_dataset_id'] == published['dataset_id']
    assert manifest['observation_count'] == 6
    assert manifest['changed_sheets'] == ['另一张待核对表']
    assert manager.sheet(first['import_id'], '新增表')['status'] == 'confirmed'
    assert {s['sheet'] for s in manager.store.list_series(result['dataset_id'])['series']} == {
        '已导入', '新增表', '另一张待核对表'}


def test_explicit_acceptance_sources_are_excluded_from_default_query_and_review_lists(tmp_path):
    manager = ImportManager(tmp_path / 'data')
    business = manager.ingest(source_file(tmp_path, '设备验收业务.xlsx'))
    path = tmp_path / '中哲_范围核对验收_非行情.xlsx'
    wb = Workbook(); ws = wb.active
    ws.append(['日期', '验收值']); ws.append([datetime(2024, 1, 1), 1]); ws.append([datetime(2024, 1, 2), 2]); wb.save(path)
    sample = manager.ingest(path)
    import json, hashlib
    (manager.root / 'source-classifications.json').write_text(json.dumps({
        hashlib.sha256(path.read_bytes()).hexdigest(): {'dataset_kind':'sample'}}))
    assert manager.store.select_dataset() == business['dataset_id']
    assert [d['source_file'] for d in manager.store.list_datasets()] == ['设备验收业务.xlsx']
    assert len(manager.store.list_datasets(include_samples=True)) == 2
    assert [d['source_file'] for d in manager.list()] == ['设备验收业务.xlsx']
    assert len(manager.list(include_samples=True)) == 2
    assert manager.store.manifest(sample['dataset_id'])['observation_count'] == 2


def test_large_legacy_issue_array_can_be_replayed_without_loading_whole_file(tmp_path):
    path = tmp_path / 'issues.json'
    values = [{'sheet': '原表', 'message': '值' * 70000}, {'sheet': '新表', 'cell': 'B3'}]
    assert write_json_array(path, iter(values)) == 2
    assert list(read_json_array(path)) == values
    path.write_text('[{"sheet":"原表"}', encoding='utf-8')
    with pytest.raises(ValueError, match='Incomplete'):
        list(read_json_array(path))


def reviewed_source(tmp_path, monkeypatch):
    from structured_data.profiles import reviewed
    path=tmp_path/'verified.xlsx'
    wb=Workbook();ws=wb.active;ws.title='核验记录'
    ws.append(['日期','收盘','结算']);ws.append([datetime(2023,1,1),10,20]);ws.append([datetime(2023,1,2),11,21])
    ws=wb.create_sheet('保留记录');ws.append(['日期','数量']);ws.append([datetime(2023,1,1),30]);ws.append([datetime(2023,1,2),31]);wb.save(path)
    rules=tmp_path/'rules';(rules/'reviewed_mappings').mkdir(parents=True)
    monkeypatch.setattr(reviewed,'__file__',str(rules/'profiles'/'reviewed.py'))
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    mapping={'sha256':digest,'rule_version':'reviewed-fixture-v1','sheets':[
        {'name':'核验记录','start_row':2,'max_col':3,'header_expectations':{'A1':'日期','B1':'收盘','C1':'结算'},'specs':[
            [2,1,'fixture.close','收盘',{'contract':'合同甲','block':'close'},None,['B1']],
            [3,1,'fixture.settle','结算',{'contract':'合同甲','block':'settle'},None,['C1']]]},
        {'name':'保留记录','start_row':2,'max_col':2,'header_expectations':{'A1':'日期','B1':'数量'},'specs':[
            [2,1,'fixture.other','数量',{},'件',['B1']]]}]}
    (rules/'reviewed_mappings'/(digest+'.json')).write_text(json.dumps(mapping))
    manager=ImportManager(tmp_path/'data');result=manager.ingest(path)
    return manager,path,result


def test_edit_verified_sheet_replaces_only_approved_sheet_and_preserves_history(tmp_path, monkeypatch):
    from copy import deepcopy
    manager,source,first=reviewed_source(tmp_path,monkeypatch)
    original=first['dataset_id'];old_bytes=source.read_bytes()
    assert manager.sheet(first['import_id'],'核验记录')['status']=='trusted_existing'
    retained=manager.store.query_series(original,'fixture.other')['rows']
    seed=manager.sheet(first['import_id'],'核验记录')['mapping']
    mapping=deepcopy(seed);mapping['blocks'][0]['columns']=mapping['blocks'][0]['columns'][:1]
    mapping['blocks'][0]['columns'][0]['metric']='调整后的收盘'
    mapping['blocks'][0]['data_end_row']=2
    preview=manager.preview(first['import_id'],mapping)
    assert preview['series_count']==1 and preview['observation_count']==1
    assert preview['series'][0]['dimensions']['contract']=='合同甲'
    assert preview['series'][0]['dimensions']['block']=='close'
    assert preview['sample'][0]['source_cell']=='B2' and preview['sample'][0]['date_cell']=='A2'
    assert manager.get(first['import_id'])['dataset_id']==original
    result=manager.confirm(first['import_id'],mapping,save_template=False);new=result['dataset_id']
    assert new!=original
    assert manager.store.manifest(new)['changed_sheets']==['核验记录']
    assert manager.store.manifest(new)['observation_count']==3
    assert manager.sheet(first['import_id'],'核验记录')['status']=='confirmed'
    new_series=manager.store.list_series(new,sheet='核验记录')['series']
    assert [s['metric'] for s in new_series]==['调整后的收盘']
    assert manager.store.query_series(new,new_series[0]['series_id'])['rows'][0]['value']==10
    assert manager.store.query_series(new,'fixture.other')['rows']==retained
    assert manager.store.query_series(original,'fixture.close')['record_count']==2
    assert (manager.store._path(new)/'source.xlsx').read_bytes()==old_bytes
    assert manager.source_path(manager.get(first['import_id'])).read_bytes()==old_bytes
    assert manager.ingest(source)['dataset_id']==new
    assert manager.confirm(first['import_id'],mapping,save_template=False)['dataset_id']==new
    assert manager.store.manifest(new)['series_count']==2


def test_failed_verified_edit_preserves_saved_mapping_and_snapshot(tmp_path, monkeypatch):
    from structured_data import generic_parser
    from copy import deepcopy
    manager,source,first=reviewed_source(tmp_path,monkeypatch)
    before=manager.get(first['import_id']);source_bytes=source.read_bytes()
    mapping=deepcopy(manager.sheet(first['import_id'],'核验记录')['mapping'])
    mapping['blocks'][0]['header_rows']=[]
    with pytest.raises(ValueError,match='表头行为空'):
        manager.confirm(first['import_id'],mapping)
    assert manager.get(first['import_id'])==before
    valid=manager.sheet(first['import_id'],'核验记录')['mapping']
    parser=generic_parser.parse_mapped_workbook;calls=0
    def fail_write(*args,**kwargs):
        nonlocal calls
        calls+=1
        if calls==2:raise ValueError('模拟写入失败')
        return parser(*args,**kwargs)
    monkeypatch.setattr(generic_parser,'parse_mapped_workbook',fail_write)
    with pytest.raises(ValueError,match='模拟写入失败'):
        manager.confirm(first['import_id'],valid)
    assert manager.get(first['import_id'])==before
    assert len(manager.store.list_datasets())==1
    assert source.read_bytes()==source_bytes


def test_verified_override_is_not_reverted_when_rebuilding_without_base(tmp_path, monkeypatch):
    manager,source,first=reviewed_source(tmp_path,monkeypatch)
    mapping=manager.sheet(first['import_id'],'核验记录')['mapping']
    mapping['blocks'][0]['columns']=mapping['blocks'][0]['columns'][:1]
    mapping['blocks'][0]['data_end_row']=2
    mapping['blocks'][0]['columns'][0]['metric']='用户映射'
    result=manager.confirm(first['import_id'],mapping,save_template=False)
    doc=manager.get(first['import_id']);doc['dataset_id']=None
    manager._publish(doc)
    dataset=doc['dataset_id']
    assert [s['metric'] for s in manager.store.list_series(dataset,sheet='核验记录')['series']]==['用户映射']
    assert manager.store.manifest(dataset)['observation_count']==3
    assert next(s for s in doc['sheets'] if s['name']=='核验记录')['status']=='confirmed'
    assert manager.store.query_series(dataset,'fixture.other')['record_count']==2
