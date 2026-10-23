"""Exercise the browser review contract through HTTP, without model calls."""
from datetime import datetime
from io import BytesIO
import time

from fastapi.testclient import TestClient
from openpyxl import Workbook
import pytest

from structured_data.server import create_app


class FakeAuthClient:
    def __init__(self, **kwargs):
        pass

    async def __aenter__(self):
        return self

    async def __aexit__(self, *args):
        pass

    async def post(self, url, json):
        assert url.endswith('/api/v1/auth/login')
        return Reply({'token': 'test-token'})

    async def get(self, url, headers):
        assert url.endswith('/api/v1/auth/me')
        return Reply({'data': {'user': {'id':'test-user','tenant_id':10000},
                              'memberships':[{'tenant_id':10000,'role':'owner'}]}})


class Reply:
    status_code = 200

    def __init__(self, value):
        self.value = value

    def json(self):
        return self.value


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr('httpx.AsyncClient', FakeAuthClient)
    monkeypatch.setenv('COTTON_MODEL_CONFIG', str(tmp_path / 'missing-model.json'))
    monkeypatch.setenv('COTTON_TENANT_ID', '10000')
    app = create_app(tmp_path / 'data')
    with TestClient(app, headers={'Origin': 'http://testserver'}) as client:
        response = client.post('/login', json={'email': 'tester@example.com', 'password': 'test-only'})
        assert response.status_code == 200
        yield client


def make_workbook(year=2023):
    book = Workbook()
    auto = book.active
    auto.title = 'simple'
    auto.append(['date', 'price'])
    auto.append([datetime(year, 1, 1), 10])
    auto.append([datetime(year, 1, 2), 11])
    parallel = book.create_sheet('parallel')
    parallel.append(['date', 'price', None, 'date', 'inventory'])
    # Keep this workflow fixture genuinely ambiguous: plain independent date
    # blocks are now auto-importable and have dedicated scanner regressions.
    parallel.append([None, 'group A', None, None, 'group B'])
    parallel.append([datetime(year, 1, 1), 20, None, datetime(year, 1, 1), 30])
    parallel.append([datetime(year, 1, 2), 21, None, datetime(year, 1, 2), 31])
    book.create_sheet('empty')
    raw = BytesIO()
    book.save(raw)
    return raw.getvalue()


def poll(client, response, *, timeout=10):
    assert response.status_code == 200, response.text
    identifier = response.json()['id']
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        response = client.get('/api/jobs/' + identifier)
        assert response.status_code == 200
        job = response.json()
        if job['status'] != 'processing':
            assert job['status'] == 'completed', job
            return job['result']
        time.sleep(.01)
    raise AssertionError('Local fixture job did not finish')


def test_request_schemas_resolve_outside_create_app(tmp_path):
    spec = create_app(tmp_path).openapi()
    for action in ('preview', 'confirm', 'suggest'):
        endpoint = spec['paths']['/api/imports/{import_id}/' + action]['post']
        assert 'requestBody' in endpoint
        assert endpoint['requestBody']['required'] is True


def test_review_preview_confirm_and_reuse_through_http(client):
    assert client.get('/api/imports').json() == {'imports': []}
    uploaded = poll(client, client.post('/api/import', files={'file': ('sample.xlsx', make_workbook())}))
    import_id = uploaded['import_id']
    assert uploaded['dataset_id']
    doc = client.get('/api/imports/' + import_id).json()
    assert 'source_path' not in doc
    assert {s['name']: s['status'] for s in doc['sheets']} == {
        'simple': 'auto_ready', 'parallel': 'needs_review', 'empty': 'empty'}
    assert {s['name']: s['published_in_dataset'] for s in doc['sheets']} == {
        'simple': True, 'parallel': False, 'empty': False}
    summary = client.get('/api/imports/' + import_id, params={'summary_only': True}).json()
    assert {s['name']: s['published_in_dataset'] for s in summary['sheets']} == {
        'simple': True, 'parallel': False, 'empty': False}
    assert {s['name']: s['import_state'] for s in summary['sheets']} == {
        'simple': 'imported', 'parallel': 'pending', 'empty': 'pending'}
    assert next(s for s in summary['sheets'] if s['name'] == 'simple')['review_state'] == (
        next(s for s in doc['sheets'] if s['name'] == 'simple')['review_state'])
    assert all('published_mapping' not in s for s in summary['sheets'])
    before = client.get('/api/datasets/' + uploaded['dataset_id'] + '/series').json()
    assert {s['sheet'] for s in before['series']} == {'simple'}
    sheet = client.get('/api/imports/' + import_id + '/sheet', params={'name': 'parallel'}).json()
    assert sheet['published_in_dataset'] is False
    mapping = sheet['mapping']
    preview = client.post('/api/imports/' + import_id + '/preview', json={'mapping': mapping})
    assert preview.status_code == 200, preview.text
    assert preview.json()['series_count'] == 2
    assert preview.json()['observation_count'] == 4
    assert preview.json()['sample'][0]['source_cell'] == 'B3'
    assert client.get('/api/imports/' + import_id).json()['dataset_id'] == uploaded['dataset_id']
    confirmed = poll(client, client.post('/api/imports/' + import_id + '/confirm', json={'mapping': mapping, 'save_template': True}))
    assert confirmed['dataset_id'] != uploaded['dataset_id']
    after = client.get('/api/datasets/' + confirmed['dataset_id'] + '/series').json()
    assert {s['sheet'] for s in after['series']} == {'simple', 'parallel'}
    saved = client.get('/api/imports/' + import_id + '/sheet', params={'name': 'parallel'}).json()
    assert saved['published_in_dataset'] is True
    assert saved['published_dataset_id'] == confirmed['dataset_id']
    assert saved['published_mapping']['blocks']
    assert len(client.get('/api/templates').json()['templates']) == 1
    repeat = poll(client, client.post('/api/import', files={'file': ('sample.xlsx', make_workbook(2024))}))
    assert repeat['import_id'] != import_id
    reused_sheet = client.get('/api/imports/' + repeat['import_id'] + '/sheet', params={'name': 'parallel'}).json()
    assert reused_sheet['status'] == 'template_ready'
    assert reused_sheet['template_status'] == 'matched'


def test_request_validation_and_unconfigured_suggestion_leave_dataset_unchanged(client):
    uploaded = poll(client, client.post('/api/import', files={'file': ('sample.xlsx', make_workbook())}))
    import_id = uploaded['import_id']
    url = '/api/imports/' + import_id
    assert client.post(url + '/preview', json={}).status_code == 422
    assert client.post(url + '/suggest', json={}).status_code == 422
    response = client.post(url + '/suggest', json={'sheet': 'parallel'})
    assert response.status_code in {400, 503}
    assert '未配置' in response.json()['detail']
    assert client.get(url).json()['dataset_id'] == uploaded['dataset_id']


def test_rescan_reuses_source_and_preserves_confirmed_mapping(client):
    uploaded=poll(client,client.post('/api/import',files={'file':('sample.xlsx',make_workbook())}))
    url='/api/imports/'+uploaded['import_id']
    before=client.get(url).json()
    r=client.post(url+'/rescan')
    assert r.status_code==200 and r.json()['reused'] is True
    after=client.get(url).json()
    assert before['dataset_id']==after['dataset_id']
    assert before['sheets']==after['sheets']
