from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook
from structured_data.profiles.cotton import (
    parse_workbook, SHEET_NAMES, CONTRACTS, FUTURES_BLOCKS,
    WAREHOUSE_HEADERS, BALANCE_HEADERS,
)


def fixture(path: Path):
    wb = Workbook()
    wb.remove(wb.active)
    fs, wh, bal = [wb.create_sheet(n) for n in SHEET_NAMES]
    for key, metric, dc in FUTURES_BLOCKS:
        fs.cell(1, dc + 1, metric)
        fs.cell(2, dc + 1, key)
        fs.cell(3, dc, "日期")
        fs.cell(4, dc, "Date")
        fs.cell(5, dc, datetime(2023, 9, 22))
        for offset, contract in enumerate(CONTRACTS, 1):
            fs.cell(4, dc + offset, contract)
    fs["B5"] = 17110
    fs["K5"] = 17200
    fs["G5"] = 16690
    fs["P5"] = 16720
    fs["Q5"] = 17375
    fs["AC5"] = "#N/A"
    wh["A3"], wh["A4"] = "日期", "Date"
    for col, metric in enumerate(WAREHOUSE_HEADERS, 2):
        wh.cell(3, col, metric)
        if col >= 32:
            wh.cell(4, col, "万吨")
    wh["A5"] = datetime(2023, 9, 21)
    wh["B5"] = 5160
    wh["P5"] = 36
    wh["AF5"] = 20.784
    wh["A6"] = datetime(2023, 9, 25)
    for col, metric in enumerate(BALANCE_HEADERS, 2):
        bal.cell(1, col, metric)
    bal["A2"] = datetime(2023, 10, 31)
    bal["C2"] = 110
    bal["V2"] = "纱厂月度负荷"
    bal["D2"] = "=1+2"
    bal["E2"] = 0
    bal["AA1"], bal["AA2"] = "产量", 9999
    wb.save(path)
    wb.close()


class ParserTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.path = Path(self.tmp.name) / "source.xlsx"
        fixture(self.path)

    def tearDown(self):
        self.tmp.cleanup()

    def test_source_fidelity_and_distinct_contracts(self):
        before = self.path.read_bytes()
        result = parse_workbook(self.path)
        self.assertEqual(before, self.path.read_bytes())
        obs = {o["series_id"]: o for o in result["observations"]}
        self.assertEqual(obs["futures.settle.CF.CZC"]["value"] - obs["futures.close.CF.CZC"]["value"], 90)
        self.assertEqual(obs["futures.settle.CF09M.CZC"]["source_cell"], "P5")
        self.assertEqual(obs["futures.settle.CF09M.CZC"]["value"], 16720)
        self.assertEqual(obs["futures.settle.CF11M.CZC"]["value"], 17375)
        warehouse_values = [o for o in result["observations"]
                            if o["series_id"] == "warehouse.B" and o["value_status"] == "valid"]
        self.assertEqual(max(warehouse_values, key=lambda o: o["date"])["date"], "2023-09-21")
        self.assertEqual(obs["balance.primary.C"]["value"], 110)
        self.assertFalse(any(o["source_cell"] == "AA2" for o in result["observations"]))
        self.assertEqual(result["sha256"], parse_workbook(self.path)["sha256"])

    def test_missing_errors_formulas_and_units_are_not_invented(self):
        result = parse_workbook(self.path)
        obs = {o["series_id"]: o for o in result["observations"]}
        self.assertIsNone(obs["futures.volume.CF.CZC"]["value"])
        self.assertEqual(obs["futures.volume.CF.CZC"]["value_status"], "excel_error")
        self.assertEqual(obs["balance.primary.V"]["value_status"], "unexpected_text")
        self.assertEqual(obs["balance.primary.D"]["value_status"], "formula_no_cache")
        self.assertEqual(obs["balance.primary.D"]["formula"], "=1+2")
        self.assertEqual(obs["balance.primary.E"]["value_status"], "valid")
        self.assertEqual(obs["balance.primary.E"]["value"], 0)
        series = {s["series_id"]: s for s in result["series"]}
        self.assertEqual(series["warehouse.AF"]["unit"], "万吨")
        self.assertIsNone(series["warehouse.B"]["unit"])
        self.assertEqual(series["warehouse.B"]["unit_status"], "needs_review")

    def test_rejects_layout_drift(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.path)
        wb[SHEET_NAMES[0]]["P4"] = "CF11M.CZC"
        wb.save(self.path)
        wb.close()
        with self.assertRaisesRegex(ValueError, "P4"):
            parse_workbook(self.path)

    def test_dated_blank_retains_source_and_does_not_replace_latest_valid(self):
        result = parse_workbook(self.path)
        self.assertEqual(result["rule_version"], "cotton-three-sheets-v2")
        observations = [o for o in result["observations"] if o["series_id"] == "warehouse.B"]
        blank = next(o for o in observations if o["date"] == "2023-09-25")
        self.assertEqual(blank["value_status"], "blank")
        self.assertEqual(blank["source_cell"], "B6")
        self.assertEqual(blank["date_cell"], "A6")
        for field in ("value", "formula", "raw_value"):
            self.assertIsNone(blank[field])
        latest_valid = max((o for o in observations if o["value_status"] == "valid"), key=lambda o: o["date"])
        self.assertEqual(latest_valid["date"], "2023-09-21")
        self.assertEqual(latest_valid["value"], 5160)
        self.assertFalse(any(o["date"] == "2023-09-26" for o in observations))
        warehouse_issues = [i for i in result["issues"] if i["sheet"] == SHEET_NAMES[1]]
        self.assertTrue(any(i["code"] == "blank_cells" and i["cell"] == "B6" for i in warehouse_issues))

    def test_uses_each_block_date(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.path)
        wb[SHEET_NAMES[0]]["J5"] = datetime(2023, 9, 21)
        wb.save(self.path)
        wb.close()
        result = parse_workbook(self.path)
        obs = {o["series_id"]: o for o in result["observations"]}
        self.assertEqual(obs["futures.close.CF.CZC"]["date"], "2023-09-22")
        self.assertEqual(obs["futures.settle.CF.CZC"]["date"], "2023-09-21")


if __name__ == "__main__":
    unittest.main()
