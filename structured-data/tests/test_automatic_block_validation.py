"""Automatic checks depend on original structure, never sheet/business names."""
from datetime import datetime, timedelta
import json

from openpyxl import Workbook, load_workbook

from structured_data.imports import ImportManager
from structured_data.scanner import scan_workbook


def original(path, *, layered=False, broken=False, offset=False):
    book = Workbook()
    sheet = book.active
    sheet.title = 'synthetic independent blocks'
    prefix = [None, None] if offset else []
    sheet.append(prefix + ['date', 'measure one', None, 'date', 'measure two'])
    if layered:
        sheet.append(prefix + [None, 'source group one', None, None, 'source group two'])
    for index in range(42 if broken else 3):
        day = datetime(2024, 1, 1) + timedelta(days=index)
        second_day = datetime(2024, 1, 1) if broken and index == 41 else day
        sheet.append(prefix + [day, index, None, second_day, index + 1])
    book.save(path)
    return path


def test_parallel_explicit_date_blocks_auto_import_without_ai_or_acknowledgement(tmp_path):
    manager = ImportManager(tmp_path / 'data')
    result = manager.ingest(original(tmp_path / 'source.xlsx'), use_reviewed_mappings=False)
    sheet = manager.get(result['import_id'])['sheets'][0]
    assert sheet['status'] == 'auto_ready'
    assert sheet['import_state'] == 'imported'
    assert [item['status'] for item in sheet['review_state']['blocks']] == ['confirmed', 'confirmed']
    assert manager.store.manifest(result['dataset_id'])['observation_count'] == 6


def test_date_blocks_need_not_start_at_column_a(tmp_path):
    path = original(tmp_path / 'source.xlsx', offset=True)
    scanned = scan_workbook(path, use_reviewed_mappings=False)['sheets'][0]
    assert scanned['status'] == 'auto_ready'
    assert [block['date_column'] for block in scanned['mapping']['blocks']] == ['C', 'F']


def test_layered_business_headings_remain_unpublished_but_valid_blocks_are_checked(tmp_path):
    manager = ImportManager(tmp_path / 'data')
    result = manager.ingest(original(tmp_path / 'source.xlsx', layered=True), use_reviewed_mappings=False)
    sheet = manager.get(result['import_id'])['sheets'][0]
    assert sheet['status'] == 'needs_review'
    assert sheet['import_state'] == 'pending'
    assert sheet['review_state']['status'] == 'structure_checked'
    assert all(item['status'] == 'confirmed' for item in sheet['review_state']['blocks'])
    assert result['dataset_id'] is None
    assert '业务指标和口径' in sheet['review_state']['notice']


def test_error_after_scanner_sample_marks_only_failing_block(tmp_path):
    manager = ImportManager(tmp_path / 'data')
    result = manager.ingest(original(tmp_path / 'source.xlsx', broken=True), use_reviewed_mappings=False)
    sheet = manager.get(result['import_id'])['sheets'][0]
    assert sheet['status'] == 'needs_review'
    assert sheet['import_state'] == 'failed'
    assert [item['status'] for item in sheet['review_state']['blocks']] == ['confirmed', 'needs_review']
    assert sheet['review_state']['blocks'][1]['diagnostics'][0]['field'] == 'date_column'
    assert result['dataset_id'] is None


def test_initial_pending_review_is_reused_without_reparsing_on_open(tmp_path, monkeypatch):
    manager = ImportManager(tmp_path / 'data')
    path = original(tmp_path / 'source.xlsx', layered=True)
    first = manager.ingest(path, use_reviewed_mappings=False)
    def forbidden(*args, **kwargs):
        raise AssertionError('An unchanged completed import must not reparse its pending sheets')
    monkeypatch.setattr('structured_data.generic_parser.parse_mapped_workbook', forbidden)
    repeated = manager.ingest(path, use_reviewed_mappings=False)
    assert repeated['reused'] and repeated['dataset_id'] == first['dataset_id']


def test_valid_first_block_cannot_hide_an_all_missing_second_block(tmp_path):
    path = original(tmp_path / 'source.xlsx')
    book = load_workbook(path)
    for row in range(2, 5):
        book.active.cell(row, 5).value = None
    book.save(path)
    manager = ImportManager(tmp_path / 'data')
    result = manager.ingest(path, use_reviewed_mappings=False)
    sheet = manager.get(result['import_id'])['sheets'][0]
    assert result['dataset_id'] is None
    assert sheet['import_state'] == 'failed'
    assert [item['status'] for item in sheet['review_state']['blocks']] == ['confirmed', 'needs_review']
    assert sheet['review_state']['blocks'][1]['diagnostics'][0]['field'] == 'columns'


def test_readiness_upgrade_does_not_replace_a_different_saved_mapping(tmp_path):
    path = original(tmp_path / 'source.xlsx')
    manager = ImportManager(tmp_path / 'data')
    first = manager.ingest(path, use_reviewed_mappings=False)
    doc = manager.get(first['import_id'])
    sheet = doc['sheets'][0]
    sheet['status'] = 'needs_review'
    sheet['mapping']['blocks'][0]['columns'][0]['metric'] = 'user-reviewed title'
    doc.pop('automatic_review_revision', None)
    (manager.path(first['import_id']) / 'review.json').write_text(json.dumps(doc))
    repeated = manager.ingest(path, use_reviewed_mappings=False)
    saved = manager.get(first['import_id'])['sheets'][0]
    assert saved['mapping']['blocks'][0]['columns'][0]['metric'] == 'user-reviewed title'
    assert saved['status'] == 'needs_review'
    assert repeated['dataset_id'] == first['dataset_id']
