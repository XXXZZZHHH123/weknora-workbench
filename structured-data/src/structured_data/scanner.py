"""Conservative layout discovery. A candidate is not a business-approved schema."""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime, time
from hashlib import sha256
import json
from pathlib import Path
import posixpath
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

RULE_VERSION = 'generic-layout-v1'
_UNITS = {'元/吨','元/件','元/米','元/公斤','美元/公斤','美元/磅','美元/吨','美分/磅','吨','万吨','千吨','包','万包','公顷','万公顷','亩','万亩','%','％','亿元','万元','美元','万美元','亿美元','万件','件','万米','米','天','摄氏度','华氏度','卢比/公担'}


def complete_date(value):
    """Never guess a year, year/month, Excel serial, US date or intraday timestamp."""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else None
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        match = re.fullmatch(r'(\d{4})[-/年](\d{1,2})[-/月](\d{1,2})日?', value.strip())
        if match:
            try:
                return date(*map(int, match.groups())).isoformat()
            except ValueError:
                pass
    return None


def _json_value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def _relations(z, part):
    ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
    relpart = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part)+'.rels')
    if relpart not in z.namelist():
        return {}
    return {r.attrib['Id']: {'type': r.attrib.get('Type','').rsplit('/',1)[-1],
        'external': r.attrib.get('TargetMode') == 'External',
        'target': r.attrib['Target'].lstrip('/') if r.attrib['Target'].startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(part), r.attrib['Target']))}
        for r in ET.fromstring(z.read(relpart)).findall(f'{{{ns}}}Relationship')}


def _drawings(path):
    main = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    office = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    result = {}
    with ZipFile(path) as z:
        rels = _relations(z, 'xl/workbook.xml')
        for sh in ET.fromstring(z.read('xl/workbook.xml')).findall(f'{{{main}}}sheets/{{{main}}}sheet'):
            sr = rels[sh.attrib[f'{{{office}}}id']]
            charts, images = [], []
            for relation in _relations(z, sr['target']).values():
                if relation['external'] or relation['type'] != 'drawing':
                    continue
                for obj in _relations(z, relation['target']).values():
                    if not obj['external'] and obj['type'] in {'chart','image'}:
                        (charts if obj['type'] == 'chart' else images).append(obj['target'])
            merged = re.findall(rb'<mergeCell\s+ref="([^"]+)"', z.read(sr['target']))
            result[sh.attrib['name']] = {'merged_ranges':[r.decode() for r in merged], 'chart_count': len(charts), 'image_count': len(images),
                'chart_parts': charts, 'image_parts': images, 'visual_content_parsed': False}
    return result


def scan_workbook(path, sample_rows=40, *, use_reviewed_mappings=True, progress=None):
    """Read bounded top-row samples of every sheet; report partial coverage honestly."""
    if type(use_reviewed_mappings) is not bool:
        raise ValueError('原件核对模式必须明确为布尔值。')
    from .progress import report_progress
    report_progress(progress, 'scan')
    path = Path(path).expanduser().resolve()
    if not 10 <= sample_rows <= 200:
        raise ValueError('sample_rows must be between 10 and 200')
    digest = sha256()
    with path.open('rb') as f:
        for block in iter(lambda: f.read(1024*1024), b''):
            digest.update(block)
    visuals = _drawings(path)
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    sheets = []
    try:
        total = len(wb.sheetnames)
        if total:
            report_progress(progress, 'scan', completed=0, total=total)
        for sheet_index, ws in enumerate(wb):
            width = min(ws.max_column or 1, 512)
            cell_rows = list(ws.iter_rows(max_row=min(ws.max_row or 1, sample_rows), max_col=width))
            rows = [tuple(c.value for c in row) for row in cell_rows]
            merged_headers = [range_boundaries(x) for x in visuals.get(ws.title, {}).get('merged_ranges', [])]
            dates = []
            for c in range(width):
                hits = [(r+1, complete_date(row[c])) for r,row in enumerate(rows) if complete_date(row[c])]
                if len(hits) >= 2:
                    dates.append({'column':get_column_letter(c+1), 'column_index':c+1,
                        'first_data_row':hits[0][0], 'date_mode':'date', 'sample_count':len(hits), 'examples':[h[1] for h in hits[:3]]})
            first = min((d['first_data_row'] for d in dates), default=min(len(rows)+1,13))
            head = [[_json_value(v) for v in row] for row in rows[:first-1]]
            def color_signature(color):
                if color is None:return None
                return [color.type, str(color.value), color.tint]
            styles = [[{'format':cell.number_format, 'fill_type':cell.fill.patternType if cell.fill else None,
                'fill':color_signature(cell.fill.fgColor if cell.fill else None), 'font':color_signature(cell.font.color if cell.font else None), 'bold':cell.font.bold if cell.font else None}
                for cell in row] for row in cell_rows[:first-1]]
            merges=[r for r in visuals.get(ws.title,{}).get('merged_ranges',[]) if range_boundaries(r)[1]<first]
            fingerprint = sha256(json.dumps({'width':ws.max_column,'headers':head,'styles':styles,'header_merges':merges,
                'date_columns':[d['column'] for d in dates]},ensure_ascii=False,sort_keys=True,default=str).encode()).hexdigest()
            blocks=[]
            for index,d in enumerate(dates):
                dc=d['column_index']; end=dates[index+1]['column_index']-1 if index+1<len(dates) else width
                start=d['first_data_row']; header_rows=list(range(1,start))
                columns=[]
                for c in range(dc+1,end+1):
                    direct=[(r,rows[r-1][c-1],f'{get_column_letter(c)}{r}') for r in header_rows if rows[r-1][c-1] is not None]
                    # A sparse upper title may label several columns inside this date block.
                    # Inheritance is a review candidate and never crosses a date column.
                    entries=[]
                    for r in header_rows:
                        value=rows[r-1][c-1];source=c
                        if value is None:
                            for left, top, right, bottom in merged_headers:
                                if top == r == bottom and dc < left <= c <= right and right <= end:
                                    value=rows[r-1][left-1];source=left;break
                        if value is None and r<header_rows[-1] and len(header_rows)>2:
                            for left in range(c-1,dc,-1):
                                prior=rows[r-1][left-1]
                                if prior is not None:
                                    if isinstance(prior,str) and prior.strip() not in _UNITS:
                                        value=prior;source=left
                                    break
                        if value is not None:entries.append((r,value,f'{get_column_letter(source)}{r}'))
                    sample=[row[c-1] for row in rows[start-1:]]
                    if not any(v is not None for v in sample) and not direct:
                        continue
                    labels=[str(v).strip() for _,v,_ in entries if isinstance(v,str) and v.strip() and v.strip() not in _UNITS]
                    units=[str(v).strip() for _,v,_ in entries if str(v).strip() in _UNITS]
                    for label in labels:
                        units.extend(u for u in re.findall(r'[（(]([^()（）]+)[）)]',label) if u in _UNITS)
                    if not any(v is not None for v in sample) and (not labels or any(len(label)>100 or label.startswith(('备注','说明')) for label in labels)):
                        # Annotation-only columns are not metrics. Named empty series remain candidates.
                        continue
                    columns.append({'column':get_column_letter(c),'metric':' / '.join(dict.fromkeys(labels)),
                        'unit':units[0] if len(set(units))==1 else None,
                        'header_cells':[coord for _,_,coord in entries]})
                blocks.append({'date_column':d['column'],'date_mode':'date','header_rows':header_rows,
                    'data_start_row':start,'data_end_row':None,'columns':columns})
            visual=visuals.get(ws.title, {'chart_count':0,'image_count':0,'visual_content_parsed':False})
            has_visual=bool(visual['chart_count'] or visual['image_count'])
            nonempty=any(v is not None for row in rows for v in row)
            reasons=[]
            status='needs_review';layout='unrecognized'
            if not dates:
                period_candidates=[]
                for c in range(width):
                    hits=[r+1 for r,row in enumerate(rows) if isinstance(row[c],str) and re.fullmatch(r'(?:\d{4}[-/年]\d{1,2}月?|\d{4}(?:年|/\d{2,4})?)', row[c].strip())]
                    if len(hits)>=2:period_candidates.append({'column':get_column_letter(c+1),'rows':hits[:5],'date_mode':'period_needs_review'})
                dates=period_candidates
                if has_visual:
                    layout='chart_or_image'; reasons.append('含图表或图片，视觉内容未解析；没有识别到完整日期表。')
                elif not nonempty and (ws.max_row or 1)<=sample_rows and (ws.max_column or 1)<=width:
                    status='empty';layout='empty'
                else:
                    layout='period_or_matrix' if dates else 'unrecognized'
                    reasons.append('未识别到完整日期纵向表；年份、年月、跨年年度不会自动转换成某一天。')
            else:
                layout='parallel_blocks' if len(blocks)>1 else 'multi_header'
                if len(blocks)>1:reasons.append('存在多个日期区块，须确认每块边界及指标。')
                b=blocks[0]
                simple = len(blocks)==1 and b['date_column']=='A' and b['data_start_row'] in (2,3) and bool(b['columns'])
                if simple:
                    metrics=[c['metric'] for c in b['columns']]
                    simple=all(metrics) and len(metrics)==len(set(metrics))
                    if b['data_start_row']==3:
                        # Only accept a dedicated unit/blank second row, not a second header layer.
                        simple=simple and all(rows[1][ordcol(c['column'])-1] is None or str(rows[1][ordcol(c['column'])-1]).strip() in _UNITS for c in b['columns'])
                    # Date-contiguous top sample and numeric/error/blank cells only.
                    simple=simple and all(complete_date(row[0]) for row in rows[b['data_start_row']-1:])
                    simple=simple and all(v is None or isinstance(v,(int,float)) and not isinstance(v,bool) or isinstance(v,str) and v.startswith('#') for row in rows[b['data_start_row']-1:] for c in b['columns'] for v in [row[ordcol(c['column'])-1]])
                if simple and not has_visual:
                    status='auto_ready';layout='single_header' if b['data_start_row']==2 else 'header_and_units'
                    reasons.append('样本符合单日期列、明确指标表头及数字/空值/错误值结构；全量导入仍执行逐行校验。')
                else:
                    reasons.append('复杂表头、混合内容或图形需要审核字段映射。')
            if (ws.max_column or 1)>width:
                status='needs_review';reasons.append('超过512列，右侧未扫描；不可自动导入整页。')
            candidate_revision='merged-headers-v2' if merged_headers and blocks else None
            regions=[]
            if blocks and status=='needs_review':
                from .row_sections import discover_row_sections, REVISION
                sections=discover_row_sections(ws,blocks,path)
                if sections:
                    blocks=sections['blocks'];regions=sections['regions'];layout='mixed_regions'
                    candidate_revision=REVISION
                    reasons=['发现每日数据下方有独立统计行。已建议拆分区块；统计行保留原始口径，不推测日期。请预览并确认，趋势展示不参与数值导入。']
            if not blocks and nonempty:
                from .matrix import discover_matrices
                matrices=discover_matrices(ws)
                if matrices:
                    blocks=matrices;layout='horizontal_matrix';status='needs_review'
                    candidate_revision='horizontal-matrix-v1'
                    reasons=['发现横向年份/期间矩阵：列标签保留原文，指标从行读取。不同国家、版本及情景必须分别核对；未纳入区块的区域不会导入。颜色不能单独决定年份或情景含义。']
            mapping={'sheet':ws.title,'fingerprint':fingerprint,'blocks':blocks}
            sheets.append({'name':ws.title,'rows':ws.max_row,'columns':ws.max_column,'status':status,
                'layout':layout,'candidate_revision':candidate_revision,'regions':regions,'fingerprint':fingerprint,'preview':[[{'cell':f'{get_column_letter(ci)}{ri}', 'value':_json_value(v)} for ci,v in enumerate(r,1)] for ri,r in enumerate(rows[:8],1)],
                'sample_rows':len(rows),'date_candidates':dates,'mapping':mapping,'reasons':reasons,'charts':visual})
            report_progress(progress, 'scan', completed=sheet_index + 1, total=total)
    finally:
        wb.close()
    # Only an identical reviewed original may supply migrated region metadata.
    configuration = None
    if use_reviewed_mappings:
        from .profiles.reviewed import reviewed_mapping
        from .regions import reviewed_regions, apply_profile
        configuration = reviewed_mapping(path)
    if configuration and configuration.get('region_profiles'):
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in sheets:
                if sheet['name'] in configuration['region_profiles']:
                    profile = reviewed_regions(path, book[sheet['name']])
                    if profile:
                        apply_profile(sheet, profile)
        finally:
            book.close()
    return {'rule_version':RULE_VERSION,'source_file':str(path),'sha256':digest.hexdigest(),
        'use_reviewed_mappings':use_reviewed_mappings,
        'sheets':sheets,'summary':dict(Counter(s['status'] for s in sheets)), 'sheet_count':len(sheets),
        'scope':'top_rows_only','sample_rows':sample_rows}


def ordcol(column):
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(column)
