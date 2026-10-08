"""Propose explicit boundaries for short dated tables with labelled summaries."""
from copy import deepcopy
from numbers import Real

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

REVISION = 'row-summary-sections-v1'
SUMMARY_LABELS = {'日变化', '日变动', '涨跌', '合计', '总计', '平均值'}
VISUAL_LABELS = {'趋势', '走势图'}


def discover_row_sections(ws, blocks, path):
    """Never truncate an uninspected tail or silently skip an invalid date.

    Only short, single-block review candidates qualify. Every remaining row in
    the mapped columns must be a known summary, an empty visual row, or blank.
    Formula cells are inspected too, even when Excel has no cached result.
    """
    from .scanner import complete_date, _json_value
    if len(blocks) != 1 or (ws.max_row or 1) > 200 or (ws.max_column or 1) > 512:
        return None
    block = blocks[0]
    if block.get('date_mode') != 'date' or not block.get('columns'):
        return None
    dc = column_index_from_string(block['date_column'])
    cols = [column_index_from_string(c['column']) for c in block['columns']]
    width = max(dc, *cols)
    rows = list(ws.iter_rows(max_row=ws.max_row, max_col=width, values_only=True))
    start = block['data_start_row']
    boundary = start
    while boundary <= len(rows) and complete_date(rows[boundary-1][dc-1]):
        boundary += 1
    if boundary - start < 2 or boundary > len(rows):
        return None
    if str(rows[boundary-1][dc-1]).strip() not in SUMMARY_LABELS:
        return None
    rawbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        rawrows = list(rawbook[ws.title].iter_rows(max_row=ws.max_row, max_col=width, values_only=True))
    finally:
        rawbook.close()
    dated = deepcopy(block)
    dated.update(name=block.get('name') or '每日数据', data_end_row=boundary-1)
    result = [dated]
    regions = []

    def region(first, last, status, reason):
        regions.append({'range':f'{block["date_column"]}{first}:{get_column_letter(width)}{last}',
            'status':status, 'reason':reason,
            'preview':[[{'cell':f'{get_column_letter(c)}{r}', 'value':_json_value(rows[r-1][c-1])}
                        for c in [dc]+cols[:7]] for r in range(first, min(last, first+2)+1)]})

    region(start, boundary-1, 'included', '每日数据；使用同一行的完整日期。')
    for rn in range(boundary, len(rows)+1):
        raw = rawrows[rn-1]
        label = str(raw[dc-1]).strip() if raw[dc-1] is not None else ''
        values = [raw[c-1] for c in cols]
        has_values = any(v is not None for v in values)
        if not label and not has_values:
            continue
        if label in VISUAL_LABELS and not has_values:
            region(rn, rn, 'needs_review', '趋势展示行，没有可读取的数值；不参与数值导入，图形未解析。')
            continue
        numeric = lambda v: v is None or isinstance(v, Real) and not isinstance(v, bool) or isinstance(v, str) and v.startswith(('=', '#'))
        if label not in SUMMARY_LABELS or not has_values or not all(numeric(v) for v in values):
            return None
        summary = deepcopy(block)
        summary.pop('date_column', None)
        summary.update(name=label, date_mode='none', category_column=block['date_column'],
                       category_name='统计口径', data_start_row=rn, data_end_row=rn)
        for metric in summary['columns']:
            metric['metric'] = f'{label} / {metric["metric"]}'
        result.append(summary)
        region(rn, rn, 'included', f'独立的“{label}”统计结果，保留原始值和公式；不推测所属日期，不混入每日数值。')
    return {'blocks':result, 'regions':regions}
