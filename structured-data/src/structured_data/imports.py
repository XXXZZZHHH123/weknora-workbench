"""Durable workbook review, opt-in schema suggestions, and reusable approved mappings."""
from __future__ import annotations

from collections import Counter, deque
from datetime import datetime, timezone
import hashlib
import json
import logging
import os
from pathlib import Path
import re
import shutil
import threading
import tempfile
import time
from contextlib import ExitStack, contextmanager

from .store import DataStore, dump
from .progress import report_progress

READY = {'auto_ready', 'template_ready', 'confirmed'}
logger = logging.getLogger(__name__)


def now():
    return datetime.now(timezone.utc).isoformat()


def atomic_json(path, value):
    temp = path.with_suffix('.tmp')
    dump(temp, value)
    os.replace(temp, path)


class ImportManager:
    def __init__(self, root):
        self.root = Path(root)
        self.directory = self.root / 'imports'
        self.templates = self.root / 'templates'
        self.directory.mkdir(parents=True, exist_ok=True)
        self.templates.mkdir(parents=True, exist_ok=True)
        self.store = DataStore(self.root)
        self.lock = threading.RLock()

    def path(self, import_id):
        if not re.fullmatch(r'[0-9a-f]{24}', import_id):
            raise ValueError('导入编号无效。')
        path = self.directory / import_id
        if not (path / 'review.json').exists():
            raise ValueError('导入记录不存在。')
        return path

    def get(self, import_id):
        return json.loads((self.path(import_id) / 'review.json').read_text())

    def list(self, *, include_samples=False):
        from .dataset_kind import dataset_kind
        records = [json.loads(p.read_text()) for p in self.directory.glob('*/review.json')]
        return sorted([{'dataset_kind': dataset_kind(d, self.root),
                        'use_reviewed_mappings': d.get('use_reviewed_mappings', True),
                        'native_source_bindings': d.get('native_source_bindings', []), **{k: d.get(k) for k in
                       ('import_id', 'source_file', 'updated_at', 'dataset_id', 'summary')}}
                       for d in records if include_samples or dataset_kind(d, self.root) != 'sample'],
                      key=lambda d: d['updated_at'], reverse=True)

    def sheet(self, import_id, name):
        doc = self.get(import_id)
        for sheet in doc['sheets']:
            if sheet['name'] == name:
                if sheet['status'] == 'trusted_existing' and not doc.get('use_reviewed_mappings', True):
                    raise ValueError('本次原件须重新核对，不能使用既有专用映射。')
                if sheet['status'] == 'trusted_existing':
                    from .profiles.reviewed import reviewed_mapping, reviewed_sheet_mapping
                    reviewed = reviewed_mapping(self.source_path(doc))
                    configured = next((s for s in (reviewed or {}).get('sheets', []) if s['name'] == name), None)
                    # Seed edits from actual verified metadata, never from the scanner guess.
                    sheet['mapping'] = reviewed_sheet_mapping(configured, sheet) if configured else {'sheet': name, 'fingerprint': sheet['fingerprint'], 'blocks': []}
                return sheet
        raise ValueError('工作表不存在。')

    def list_templates(self):
        return [json.loads(p.read_text()) for p in self.templates.glob('*.json')]

    def source_path(self, doc):
        return self.directory / doc['import_id'] / doc['source_file']

    def bind_native_source(self, import_id, source, *, before_commit=None):
        with self.lock:
            doc = self.get(import_id)
            bindings = list(doc.get('native_source_bindings', []))
            if source not in bindings:
                bindings.append(dict(source))
                doc['native_source_bindings'] = bindings
                # A knowledge-base source binding is durable authorization
                # metadata. Recheck access immediately before writing it, just
                # as snapshot publication does at its atomic commit boundary.
                if before_commit is not None:
                    before_commit()
                self._save(doc)
            return self._result(doc, reused=True)

    def find_native_source(self, source, *, use_reviewed_mappings=True):
        with self.lock:
            for record in self.directory.glob('*/review.json'):
                if not re.fullmatch(r'[0-9a-f]{24}', record.parent.name):
                    continue
                doc = json.loads(record.read_text())
                if doc.get('use_reviewed_mappings', True) != use_reviewed_mappings:
                    continue
                if source not in doc.get('native_source_bindings', []):
                    continue
                original = self.source_path(doc)
                if not original.is_file() or original.stat().st_size != source['file_size']:
                    raise ValueError('既有结构化原件丢失或已改变，请核对已导入文件。')
                with original.open('rb') as stream:
                    if hashlib.file_digest(stream, 'md5').hexdigest() != source['file_hash']:
                        raise ValueError('既有结构化原件校验失败，请核对已导入文件。')
                return self._result(doc, reused=True)
        return None

    @staticmethod
    def template_key(sheet):
        return hashlib.sha256((sheet['name'] + ':' + sheet['fingerprint']).encode()).hexdigest()

    def _save(self, doc):
        doc['summary'] = dict(Counter(s['status'] for s in doc['sheets']))
        doc['updated_at'] = now()
        atomic_json(self.directory / doc['import_id'] / 'review.json', doc)

    def ingest(self, path, *, use_reviewed_mappings=True, before_publish=None,
               native_source=None, progress=None):
        from .scanner import scan_workbook
        with self.lock:
            scan = scan_workbook(path, use_reviewed_mappings=use_reviewed_mappings,
                                 **({'progress': progress} if progress is not None else {}))
            identity = scan['sha256'] + ':' + scan['rule_version']
            if not use_reviewed_mappings:
                identity += ':fresh-original-v1'
            import_id = hashlib.sha256(identity.encode()).hexdigest()[:24]
            folder = self.directory / import_id
            if (folder / 'review.json').exists():
                doc = self.get(import_id)
                changed = False
                binding_changed = False
                if native_source is not None:
                    bindings = list(doc.get('native_source_bindings', []))
                    if native_source not in bindings:
                        bindings.append(dict(native_source))
                        doc['native_source_bindings'] = bindings
                        binding_changed = True
                for updated in scan['sheets']:
                    prior=next((x for x in doc['sheets'] if x['name']==updated['name']),None)
                    if prior and prior['status']=='needs_review' and ((updated.get('region_profile') and updated.get('region_profile')!=prior.get('region_profile')) or (updated.get('candidate_revision') and updated.get('candidate_revision')!=prior.get('candidate_revision'))):
                        prior.update(updated)
                        changed = True
                if doc.get('publish_status') != 'completed':
                    self._publish(doc, before_publish=before_publish, progress=progress)
                    doc['publish_status'] = 'completed'
                    self._save(doc)
                elif changed or binding_changed:
                    # A completed import can still receive updated scanner
                    # diagnostics. They are durable review metadata, so do not
                    # save them after a long rescan unless write access still
                    # holds at this final boundary.
                    if before_publish is not None:
                        before_publish()
                    self._save(doc)
                return self._result(doc, reused=True)
            folder.mkdir(exist_ok=True)
            doc = None
            try:
                saved = folder / Path(path).name
                shutil.copyfile(path, saved)
                with saved.open('rb') as original:
                    saved_digest = hashlib.file_digest(original, 'sha256').hexdigest()
                if saved_digest != scan['sha256']:
                    raise ValueError('扫描期间源文件改变，请重新上传。')
                doc = {**scan, 'import_id': import_id, 'source_file': Path(path).name,
                       'dataset_id': None, 'publish_status':'pending', 'updated_at': now()}
                if native_source is not None:
                    doc['native_source_bindings'] = [dict(native_source)]
                for sheet in doc['sheets']:
                    prefix = '' if use_reviewed_mappings else 'fresh-'
                    template_file = self.templates / (prefix + self.template_key(sheet) + '.json')
                    if template_file.exists():
                        template = json.loads(template_file.read_text())
                        # Sheet names can be semantic dimensions, so require them as well as layout.
                        if template['sheet'] == sheet['name'] and (not template.get('source_sha256') or template['source_sha256']==scan['sha256']):
                            sheet['mapping'] = template['mapping']
                            sheet['status'] = 'template_ready'
                            sheet['template_status'] = 'matched'
                        elif template.get('source_sha256') != scan['sha256']:
                            sheet['template_status'] = 'source_changed'
                            sheet['status'] = 'needs_review'
                            sheet['reasons'].append('独立小表的源文件已改变，请重新核对区块范围、日期来源和类别。')
                    elif any(t['sheet'] == sheet['name'] and t.get('use_reviewed_mappings', True) == use_reviewed_mappings
                             for t in self.list_templates()):
                        sheet['template_status'] = 'layout_changed'
                        sheet['status'] = 'needs_review'
                        sheet['reasons'].append('已保存模板与本次表头指纹不同，必须重新确认。')
                # Expected per-sheet parse failures are recorded in ``doc`` by
                # _publish. Save the source and review only after its final
                # authorization guard succeeds, so a revoked first import does
                # not leave a discoverable half-created record behind.
                self._publish(doc, before_publish=before_publish, progress=progress)
                doc['publish_status'] = 'completed'
                self._save(doc)
                return self._result(doc)
            except Exception:
                # Before a snapshot is committed, this directory is only
                # staging for the current request. Removing it also removes the
                # copied source and any premature review metadata on denial.
                if doc is None or not doc.get('dataset_id'):
                    shutil.rmtree(folder, ignore_errors=True)
                raise

    @staticmethod
    def _result(doc, reused=False):
        result = {'import_id': doc['import_id'], 'dataset_id': doc.get('dataset_id'),
                  'summary': doc['summary'], 'reused': reused}
        preparation = doc.get('chart_catalog_preparation')
        if isinstance(preparation, dict):
            result['chart_catalog_preparation'] = dict(preparation)
        return result

    def _prepare_chart_catalog(self, dataset_id):
        """Warm disposable chart metadata after publication without affecting it."""
        started = time.perf_counter()
        try:
            self.store.chart_index(dataset_id)
        except Exception as exc:
            elapsed_ms = round((time.perf_counter() - started) * 1000)
            result = {'status': 'failed', 'purpose': 'derived_catalog_only',
                      'elapsed_ms': elapsed_ms, 'error_type': type(exc).__name__}
            logger.warning('published_chart_catalog_preparation_failed dataset_id=%s elapsed_ms=%d error_type=%s',
                           dataset_id, elapsed_ms, type(exc).__name__)
            return result
        elapsed_ms = round((time.perf_counter() - started) * 1000)
        logger.info('published_chart_catalog_preparation_ready dataset_id=%s elapsed_ms=%d',
                    dataset_id, elapsed_ms)
        return {'status': 'ready', 'purpose': 'derived_catalog_only',
                'elapsed_ms': elapsed_ms}

    @contextmanager
    def _records(self):
        from .streaming import DiskRows
        with tempfile.TemporaryDirectory(prefix='.import-parse-', dir=self.store.datasets) as temporary:
            observations = DiskRows(Path(temporary) / 'observations.jsonl')
            issues = DiskRows(Path(temporary) / 'issues.jsonl')
            try:
                yield observations, issues
            finally:
                observations.close()
                issues.close()

    @staticmethod
    def _mapping_signature(mapping):
        # Review selection identifiers are browser state, never data versions.
        def clean(value):
            if isinstance(value, dict):
                return {k: clean(v) for k, v in value.items() if not k.startswith('_review')}
            if isinstance(value, list):
                return [clean(v) for v in value]
            return value
        return json.dumps(clean(mapping), ensure_ascii=False, sort_keys=True)

    def _publish(self, doc, changed_sheet=None, before_publish=None, progress=None):
        from .profiles.reviewed import parse_workbook, reviewed_mapping
        from .generic_parser import parse_mapped_workbook, validate_mapping
        from .snapshots import publish_parts
        path = self.source_path(doc)
        base_id = doc.get('dataset_id')
        base = self.store.manifest(base_id) if base_id else None
        # Older pilot tools sometimes published a newer snapshot without moving
        # review.json. Preserve that snapshot's other tables when confirming one.
        if base and not doc.get('use_reviewed_mappings', True):
            if base.get('use_reviewed_mappings', True) or base.get('import_id') != doc['import_id']:
                raise ValueError('本次原件没有同一核对范围的快照，未复用其他解析结果。')
        if base and doc.get('use_reviewed_mappings', True):
            for item in self.store.list_datasets(include_samples=True):
                if item['sha256'] != doc['sha256']:
                    continue
                candidate = self.store.manifest(item['dataset_id'])
                if candidate.get('import_id') == doc['import_id']:
                    if candidate['imported_at'] > base['imported_at']:
                        base_id, base = item['dataset_id'], candidate
                    break
        base_sheets = {s['name']: s for s in base['sheets']} if base else {}
        if changed_sheet and base_id != doc.get('dataset_id'):
            published_coverage = {s['name']: s for s in base.get('coverage', [])}
            for sheet in doc['sheets']:
                prior = published_coverage.get(sheet['name'])
                if sheet['name'] != changed_sheet and sheet['name'] in base_sheets and prior:
                    sheet.update({k: prior[k] for k in ('status', 'reasons')})
                    if base_sheets[sheet['name']].get('mapping'):
                        sheet['mapping'] = base_sheets[sheet['name']]['mapping']
        legacy = reviewed_mapping(path) if doc.get('use_reviewed_mappings', True) else None
        if legacy:
            # Explicit user/template mappings supersede initial verified defaults.
            # Other verified tables keep their saved records and original rules.
            overrides = {s['name'] for s in doc['sheets'] if s['status'] in {'confirmed', 'template_ready'}}
            if changed_sheet:
                overrides.add(changed_sheet)
            legacy = {**legacy, 'sheets': [s for s in legacy['sheets'] if s['name'] not in overrides]}
        legacy_names = {s['name'] for s in legacy['sheets']} if legacy else set()
        has_legacy = bool(legacy_names) and legacy_names.issubset({s['name'] for s in doc['sheets']})
        reuse_legacy = has_legacy and legacy_names.issubset(base_sheets)
        candidates = [s for s in doc['sheets'] if s['status'] in READY and
                      (s['name'] == changed_sheet if changed_sheet else True)]
        metadata = {'coverage': [], 'import_id': doc['import_id'],
                    'use_reviewed_mappings': doc.get('use_reviewed_mappings', True),
                    'charts': {'status': 'inventory_only', 'visual_content_parsed': False,
                               'sheets': [{'name': s['name'], **s['charts']} for s in doc['sheets']]}}
        failed = set()
        imported = []
        candidate_names = {sheet['name'] for sheet in candidates
                           if sheet['status'] != 'trusted_existing'}
        pending_legacy = {sheet['name'] for sheet in doc['sheets']
                          if sheet['name'] in legacy_names and not reuse_legacy}
        total = len(candidate_names | pending_legacy)
        completed_names = set()

        def parsed_sheet(sheet):
            completed_names.add(sheet['name'])
            report_progress(progress, 'parse', completed=len(completed_names), total=total)

        if total:
            report_progress(progress, 'parse', completed=0, total=total)
        with ExitStack() as temporary:
            parts = []
            for sheet in doc['sheets']:
                if sheet['name'] not in legacy_names or reuse_legacy:
                    continue
                observations, issues = temporary.enter_context(self._records())
                try:
                    configured = next(s for s in legacy['sheets'] if s['name'] == sheet['name'])
                    parsed = parse_workbook(path, mapping={**legacy, 'sheets': [configured]},
                                            observations=observations, issues=issues)
                except ValueError as exc:
                    sheet['status'] = 'needs_review'
                    sheet['reasons'] = ['原专用规则校验失败：' + str(exc)]
                    sheet['review_state'] = self._review_failed(doc, sheet, sheet.get('mapping'), exc)
                    failed.add(sheet['name'])
                else:
                    sheet['status'] = 'trusted_existing'
                    sheet['reasons'] = ['沿用绑定此原文件指纹的已核对映射；仅导入已核对的区域。']
                    sheet['review_state'] = self._review_success(sheet.get('mapping', {}))
                    parts.append(parsed)
                    imported.append(sheet['name'])
                finally:
                    if sheet['status'] == 'trusted_existing' or sheet['name'] not in candidate_names:
                        parsed_sheet(sheet)
            for sheet in candidates:
                if sheet['status'] == 'trusted_existing':
                    continue
                observations, issues = temporary.enter_context(self._records())
                try:
                    mapping = validate_mapping(sheet['mapping'], sheet)
                    if (not changed_sheet and sheet['name'] in base_sheets and
                            self._mapping_signature(mapping) ==
                            self._mapping_signature(base_sheets[sheet['name']].get('mapping'))):
                        sheet['review_state'] = self._review_success(mapping)
                        continue
                    parsed = parse_mapped_workbook(path, [mapping], scan=doc,
                                                   observations=observations, issues=issues)
                    if not parsed['series'] or not len(observations):
                        raise ValueError('映射没有产生可查询记录，不能导入。')
                except ValueError as exc:
                    review = self._review_failed(doc, sheet, sheet.get('mapping'), exc)
                    if sheet['name'] == changed_sheet:
                        exc.review_state = review
                        raise
                    sheet['status'] = 'needs_review'
                    sheet['reasons'] = ['全量校验未通过，未导入此表：' + str(exc)]
                    sheet['review_state'] = review
                    failed.add(sheet['name'])
                else:
                    sheet['review_state'] = self._review_success(mapping)
                    parts.append(parsed)
                    imported.append(sheet['name'])
                finally:
                    parsed_sheet(sheet)
            coverage = {s['name']: s for s in base.get('coverage', [])} if base else {}
            for sheet in doc['sheets']:
                # A failed replacement leaves the prior published sheet intact.
                if sheet['name'] in failed and sheet['name'] in base_sheets:
                    continue
                coverage[sheet['name']] = {k: sheet[k] for k in ('name', 'status', 'reasons')}
            metadata['coverage'] = [coverage[s['name']] for s in doc['sheets']]
            report_progress(progress, 'publish')
            if not parts and base and metadata['coverage'] == base.get('coverage', []):
                # Even a no-op snapshot can update review metadata below.
                if before_publish is not None:
                    before_publish()
                doc['dataset_id'] = base_id
                result = None
            else:
                result = publish_parts(self.store, path, doc['sha256'], parts, metadata,
                                       base_id=base_id, replace_sheets=imported,
                                       before_commit=before_publish)
            if result:
                doc['dataset_id'] = result['dataset_id']
                report_progress(progress, 'index')
                doc['chart_catalog_preparation'] = self._prepare_chart_catalog(result['dataset_id'])
        published = set(base_sheets) | set(imported)
        for sheet in doc['sheets']:
            sheet['import_state'] = ('imported' if sheet['name'] in published else
                                     'failed' if sheet['name'] in failed else 'pending')
        return {'sheets_imported': imported,
                'sheets_failed': [{'sheet': s['name'], 'message': s['reasons'][0],
                                   'review_state': s['review_state']}
                                  for s in doc['sheets'] if s['name'] in failed],
                'sheets_pending': [s['name'] for s in doc['sheets'] if s['status'] not in READY
                                   and s['status'] != 'trusted_existing']}

    @staticmethod
    def _review_success(mapping):
        return {'status': 'confirmed', 'diagnostics': [],
                'blocks': [{'block_index': index, 'block_id': f'block-{index}',
                            'status': 'confirmed', 'diagnostics': []}
                           for index, _ in enumerate(mapping.get('blocks', []))]}

    @staticmethod
    def _diagnostic(exc, sheet, block_index=None):
        from .review_errors import clean_target
        target = clean_target(getattr(exc, 'review_target', {}))
        if block_index is not None:
            target['block_index'] = block_index
        target['sheet'] = sheet
        if 'block_index' in target:
            target['block_id'] = f'block-{target["block_index"]}'
        return {'code': 'mapping_validation_failed', 'message': str(exc)[:2000], **target}

    def _review_failed(self, doc, sheet, mapping, exc):
        from .generic_parser import parse_mapped_workbook
        from .review_errors import clean_target
        diagnostic = self._diagnostic(exc, sheet['name'])
        blocks = mapping.get('blocks', []) if isinstance(mapping, dict) else []
        review = {'status': 'needs_review', 'blocks': [], 'diagnostics': [diagnostic]}
        first_target = clean_target(getattr(exc, 'review_target', {})).get('block_index')
        for index, block in enumerate(blocks):
            item = {'block_index': index, 'block_id': f'block-{index}',
                    'status': 'confirmed', 'diagnostics': []}
            if index == first_target:
                item.update(status='needs_review', diagnostics=[diagnostic])
            else:
                # This slower path is only used after a sheet fails. Full data,
                # rather than scanner samples, determine each remaining block.
                with self._records() as (observations, issues):
                    try:
                        parsed = parse_mapped_workbook(self.source_path(doc),
                                                       [{**mapping, 'blocks': [block]}], scan=doc,
                                                       observations=observations, issues=issues)
                        if not parsed['series'] or not len(observations):
                            raise ValueError('区块没有产生可查询记录。')
                    except ValueError as block_error:
                        problem = self._diagnostic(block_error, sheet['name'], index)
                        item.update(status='needs_review', diagnostics=[problem])
                        review['diagnostics'].append(problem)
            review['blocks'].append(item)
        # A sheet-level error has no block location. Preserve blocks that pass
        # their own full parse when another block was isolated as the cause.
        # Only mark every block pending when the cross-block failure cannot be
        # localized at all (for example, a conflict between otherwise valid
        # blocks).
        if first_target is None and not any(
                item['status'] == 'needs_review' for item in review['blocks']):
            for item in review['blocks']:
                if item['status'] == 'confirmed':
                    item.update(status='needs_review', diagnostics=[diagnostic])
        return review

    def review_mapping(self, import_id, mapping):
        """Check every selected block without publishing data or changing review."""
        doc = self.get(import_id)
        sheet = self.sheet(import_id, mapping.get('sheet'))
        try:
            self.preview(import_id, mapping)
        except ValueError as exc:
            return self._review_failed(doc, sheet, mapping, exc)
        return self._review_success(mapping)

    def confirm_all(self, import_id, save_template=True, mappings=None):
        """Import recognized sheets and explicit user-confirmed drafts independently."""
        from .review_errors import error_location
        with self.lock:
            doc = self.get(import_id)
            drafts = [] if mappings is None else mappings
            if not isinstance(drafts, list) or len(drafts) > 256:
                raise ValueError('批量确认须为不超过 256 个工作表映射的列表。')
            names = []
            for mapping in drafts:
                if not isinstance(mapping, dict) or not isinstance(mapping.get('sheet'), str):
                    raise ValueError('批量确认中的工作表映射无效。')
                name = mapping['sheet']
                if name in names:
                    raise ValueError('批量确认不能重复提交同一工作表。')
                self.sheet(import_id, name)
                names.append(name)
            result = self._publish(doc)
            doc['publish_status'] = 'completed'
            self._save(doc)
            imported = list(result['sheets_imported'])
            failures = {item['sheet']: item for item in result['sheets_failed']}
            for mapping in drafts:
                name = mapping['sheet']
                try:
                    self.confirm(import_id, mapping, save_template=save_template)
                except ValueError as exc:
                    failures[name] = {'sheet': name, 'message': str(exc), **error_location(exc)}
                else:
                    failures.pop(name, None)
                    if name not in imported:
                        imported.append(name)
            doc = self.get(import_id)
            pending = [s['name'] for s in doc['sheets'] if s['status'] not in READY
                       and s['status'] != 'trusted_existing']
            return {**self._result(doc), 'sheets_imported': imported,
                    'sheets_failed': list(failures.values()), 'sheets_pending': pending}

    def preview(self, import_id, mapping):
        from .generic_parser import parse_mapped_workbook, validate_mapping
        doc = self.get(import_id)
        sheet = self.sheet(import_id, mapping.get('sheet'))
        mapping = validate_mapping(mapping, sheet)
        with self._records() as (observations, issues):
            parsed = parse_mapped_workbook(self.source_path(doc), [mapping], scan=doc,
                                           observations=observations, issues=issues)
            blocks = []
            summaries = {i: {'first': [], 'last': deque(maxlen=10), 'count': 0}
                         for i in range(len(mapping['blocks']))}
            matrix_samples = {}
            sample = []
            for row in observations:
                if len(sample) < 20:
                    sample.append(row)
                bi = row.get('block_index', 0)
                summary = summaries[bi]
                summary['count'] += 1
                if len(summary['first']) < 10:
                    summary['first'].append(row)
                summary['last'].append(row)
                if mapping['blocks'][bi].get('layout') == 'matrix':
                    item = matrix_samples.setdefault(row['series_id'], {'first': row, 'last': deque(maxlen=2), 'count': 0})
                    item['last'].append(row)
                    item['count'] += 1
            bounded_issues = issues[:50]
            for bi, block in enumerate(mapping['blocks']):
                series = [x for x in parsed['series'] if x.get('block_index') == bi]
                summary = summaries[bi]
                if block.get('layout') == 'matrix':
                    block_sample = []
                    for item in series[:20]:
                        rows = matrix_samples.get(item['series_id'])
                        if rows:
                            block_sample.extend(([rows['first']] if rows['count'] > 2 else []) + list(rows['last']))
                else:
                    block_sample = summary['first'] + list(summary['last'])[max(0, 20-summary['count']):] if summary['count'] > 10 else summary['first']
                blocks.append({'name': block.get('name', f'区块{bi+1}'), 'series_count': len(series),
                               'observation_count': summary['count'], 'series': series[:20],
                               'sample': block_sample, 'issues': bounded_issues[:20]})
            return {'blocks': blocks, 'excluded_regions': parsed.get('excluded_regions', []),
                    'mapping': mapping, 'series_count': len(parsed['series']),
                    'observation_count': len(observations), 'series': parsed['series'][:20],
                    'sample': sample, 'issues': bounded_issues,
                    'notice': ('同名指标按原表列分别保存，未合并或推测品种、年度。' if any('原表列' in c.get('dimensions', {}) for b in mapping['blocks'] for c in b['columns']) else '') + '预览已执行全量结构校验，但业务指标、单位和年度口径仍需人工确认。'}

    def confirm(self, import_id, mapping, save_template=True, before_publish=None, progress=None):
        with self.lock:
            report_progress(progress, 'preview')
            try:
                checked = self.preview(import_id, mapping)
                if not checked['series_count'] or not checked['observation_count']:
                    raise ValueError('映射没有产生可查询记录，不能确认。')
            except ValueError as exc:
                doc = self.get(import_id)
                sheet = self.sheet(import_id, mapping.get('sheet'))
                exc.review_state = self._review_failed(doc, sheet, mapping, exc)
                raise
            doc = self.get(import_id)
            sheet = next(s for s in doc['sheets'] if s['name'] == mapping['sheet'])
            sheet['mapping'] = checked['mapping']
            # Local headers inside the sheet must be rechecked on template reuse too.
            from openpyxl import load_workbook
            from openpyxl.utils import column_index_from_string
            workbook=load_workbook(self.source_path(doc),read_only=True,data_only=True,keep_links=False)
            try:
                ws=workbook[sheet['name']]
                refs={cell for b in sheet['mapping']['blocks'] for c in b['columns'] for cell in c['header_cells']}
                maxrow=max(int(re.search(r'[0-9]+',cell).group()) for cell in refs) if refs else 1
                maxcol=max((column_index_from_string(re.match(r'[A-Z]+', cell).group()) for cell in refs), default=1)
                values={cell.coordinate:str(cell.value) for row in ws.iter_rows(max_row=maxrow, max_col=maxcol) for cell in row if getattr(cell,'coordinate',None) in refs}
                for b in sheet['mapping']['blocks']:
                    b['header_expectations']={cell:values[cell] for c in b['columns'] for cell in c['header_cells']}
            finally:workbook.close()
            sheet['status'] = 'confirmed'
            sheet['reasons'] = ['用户确认映射；保留原文件及单元格来源。']
            self._publish(doc, changed_sheet=sheet['name'], before_publish=before_publish,
                          progress=progress)
            if sheet['status'] != 'confirmed':
                raise ValueError('确认未通过全量校验，原版本保持有效。')
            if save_template:
                prefix = '' if doc.get('use_reviewed_mappings', True) else 'fresh-'
                atomic_json(self.templates / (prefix + self.template_key(sheet) + '.json'),
                            {'fingerprint':sheet['fingerprint'],'sheet':sheet['name'],
                             'mapping':sheet['mapping'],'created_at':now(), 'scope':'matching_sheet_and_header_fingerprint',
                             'use_reviewed_mappings':doc.get('use_reviewed_mappings', True),
                             'source_sha256':doc['sha256'] if sheet.get('region_profile') or any(b['date_mode'] in ('shared_cell','none','period') for b in sheet['mapping']['blocks']) else None})
                sheet['template_status'] = 'saved'
            self._save(doc)
            return self._result(doc)

    def suggest(self, import_id, sheet_name, *, mapping=None, review_errors=None):
        """Return source-grounded model advice; publication is an explicit operation."""
        from .ai_review import suggest_mapping
        return suggest_mapping(self, import_id, sheet_name, mapping=mapping,
                               review_errors=review_errors)
