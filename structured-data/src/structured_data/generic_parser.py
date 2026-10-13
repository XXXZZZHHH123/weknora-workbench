"""Parse explicit, fingerprint-checked numeric table mappings without model calls."""
from __future__ import annotations

from collections import Counter
from copy import deepcopy
from datetime import date, datetime
from hashlib import sha256
from itertools import zip_longest
from math import isfinite
from pathlib import Path
import re
import json

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.datetime import from_excel

from .scanner import complete_date, scan_workbook

RULE_VERSION = 'generic-mapped-v3'


def has_period_number(value):
    return isinstance(value,(int,float)) and not isinstance(value,bool) and isfinite(value) and value==int(value)


def _column(value, width):
    if not isinstance(value,str) or not re.fullmatch('[A-Z]{1,3}',value):
        raise ValueError('列号必须使用大写Excel列名，例如 A、AB')
    col=column_index_from_string(value)
    if not 1<=col<=width:
        raise ValueError(f'列 {value} 超出工作表范围')
    return col


def validate_mapping(mapping: dict, sheet_scan: dict) -> dict:
    """Validate explicit bounds/semantics; return a copied normalized mapping."""
    if not isinstance(mapping,dict):raise ValueError('映射必须是对象')
    m=deepcopy(mapping)
    if m.get('sheet') != sheet_scan['name']:raise ValueError('工作表名称不匹配')
    if m.get('fingerprint') != sheet_scan['fingerprint']:
        raise ValueError('工作表结构指纹变化，旧规则不可直接复用，请重新确认')
    blocks=m.get('blocks')
    if not isinstance(blocks,list) or not blocks:raise ValueError('至少需要一个数据区块')
    for block in blocks:
        if isinstance(block, dict):
            for key in list(block):
                if key.startswith('_review'):
                    del block[key]
    used=set()
    meanings=set()
    for b in blocks:
        if not isinstance(b,dict):raise ValueError('区块必须是对象')
        if b.get('layout')=='matrix':
            from .matrix import validate_matrix
            validate_matrix(b,sheet_scan)
            for c in b['columns']:
                ci=_column(c['column'],sheet_scan['columns']);start=b['data_start_row'];end=b['data_end_row']
                if any(ci==x[0] and start<=x[2] and end>=x[1] for x in used):raise ValueError('矩阵与其他区块范围重叠')
                used.add((ci,start,end))
            continue
        shared=b.get('date_mode')=='shared_cell'
        undated=b.get('date_mode')=='none'
        dc=None if shared or undated else _column(b.get('date_column'),sheet_scan['columns'])
        if shared:
            cell=b.get('date_cell','')
            match=re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]*)',cell)
            if not match or int(match.group(2))>sheet_scan['rows']:raise ValueError('共用日期必须填写原表内的日期单元格，例如 AG15')
            _column(match.group(1),sheet_scan['columns'])
            if b.get('data_end_row') is None:raise ValueError('共用日期区块必须明确结束行，不能延伸到整页末尾')
        if b.get('category_column'):
            _column(b['category_column'],sheet_scan['columns'])
            if not isinstance(b.get('category_name','类别'),str) or not b.get('category_name','类别').strip():raise ValueError('类别名称不能为空')
        if b.get('date_mode') not in ('date','excel_serial','shared_cell','custom','none','period'):
            raise ValueError('请选择日期、自定义日期格式、期间标签或无时间字段；年度用期间标签，不转换成某一天')
        if b.get('date_mode')=='custom' or shared and b.get('date_format'):_format_pattern(b.get('date_format',''))
        if b.get('date_mode') in ('none','period') and b.get('data_end_row') is None:raise ValueError('无日期或期间标签表必须明确结束行')
        start=b.get('data_start_row');end=b.get('data_end_row')
        if type(start) is not int or not 1<=start<=sheet_scan['rows']:
            raise ValueError('数据起始行超出范围')
        if end is not None and (type(end) is not int or not start<=end<=sheet_scan['rows']):
            raise ValueError('数据结束行超出范围')
        headers=b.get('header_rows',[])
        if not isinstance(headers,list) or not headers or any(type(r) is not int or not 1<=r<start for r in headers):
            raise ValueError('表头行必须在数据起始行之前')
        expectations=b.get('header_expectations',{})
        if not isinstance(expectations,dict):raise ValueError('表头核对信息必须是对象')
        for cell,expected in expectations.items():
            match=re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]*)',cell) if isinstance(cell,str) else None
            if not match or int(match.group(2)) not in headers or not isinstance(expected,str):raise ValueError('表头核对信息必须指向声明的表头行')
            _column(match.group(1),sheet_scan['columns'])
        columns=b.get('columns')
        if not isinstance(columns,list) or not columns:raise ValueError('区块没有指标列')
        for c in columns:
            if not isinstance(c,dict):raise ValueError('指标列必须是对象')
            ci=_column(c.get('column'),sheet_scan['columns'])
            if b.get('category_column')==c.get('column'):raise ValueError('类别列不能同时作为数值指标列')
            if ci==dc:raise ValueError('指标列不能与日期列相同')
            key=(ci,start,end or sheet_scan['rows'])
            if any(ci==prior[0] and start<=prior[2] and key[2]>=prior[1] for prior in used):
                raise ValueError('同一指标列的数据范围重复或重叠')
            used.add(key)
            metric=c.get('metric')
            if not isinstance(metric,str) or not metric.strip():raise ValueError('指标名称不能为空')
            c['metric']=metric.strip()
            if c.get('unit') is not None and (not isinstance(c['unit'],str) or not c['unit'].strip()):
                raise ValueError('单位必须为空或非空字符串')
            dims=c.get('dimensions',{})
            if not isinstance(dims,dict) or any(not isinstance(k,str) or not isinstance(v,(str,int,float,bool)) for k,v in dims.items()):
                raise ValueError('dimensions 必须是简单键值对象')
            meaning=(c['metric'],json.dumps(dims,ensure_ascii=False,sort_keys=True), b.get('name',''), b.get('category_column',''))
            if meaning in meanings:
                raise ValueError('同一工作表多个指标列的名称和维度完全相同，请添加年度、品种或字段维度后确认')
            meanings.add(meaning)
            c['dimensions']=dims
            expected=[f"{c['column']}{r}" for r in headers]
            cells=c.get('header_cells',expected)
            if not isinstance(cells,list) or any(not isinstance(cell,str) or not re.fullmatch('[A-Z]{1,3}[1-9][0-9]*',cell) for cell in cells):
                raise ValueError('表头来源位置格式无效')
            # A submitted rule cannot invent provenance outside its declared header.
            for cell in cells:
                mo=re.fullmatch('([A-Z]+)([0-9]+)',cell)
                _column(mo.group(1),sheet_scan['columns'])
                if int(mo.group(2)) not in headers:raise ValueError('表头来源位置不在声明的表头行内')
            c['header_cells']=cells
    return m


def _format_pattern(value):
    if not isinstance(value,str) or len(value)>80:raise ValueError('请填写日期格式，例如 YYYYMMDD 或 DD/MM/YYYY')
    tokens=re.findall(r'YYYY|MM|DD',value.upper())
    if sorted(tokens)!=['DD','MM','YYYY']:raise ValueError('日期格式必须各包含一次 YYYY、MM、DD；只有年月或年度请选期间标签')
    parts=re.split(r'(YYYY|MM|DD)',value.upper())
    if any(re.search('[A-Z%]',x) for x in parts if x not in ('YYYY','MM','DD')):raise ValueError('日期格式仅支持 YYYY、MM、DD 和分隔文字或符号')
    patterns={'YYYY':r'(?P<year>\d{4})','MM':r'(?P<month>\d{1,2})','DD':r'(?P<day>\d{1,2})'}
    if re.fullmatch(r'(YYYY|MM|DD){3}',value.upper()):patterns.update(MM=r'(?P<month>\d{2})',DD=r'(?P<day>\d{2})')
    return ''.join(patterns.get(x,re.escape(x)) for x in parts)


def _date(value, mode, epoch, date_format=None):
    if mode=='custom':
        pattern=_format_pattern(date_format)
        if isinstance(value,(date,datetime)):return complete_date(value)
        text=str(int(value)) if isinstance(value,(int,float)) and not isinstance(value,bool) and isfinite(value) and value==int(value) else str(value).strip()
        match=re.fullmatch(pattern,text)
        if not match:return None
        try:return date(**{k:int(v) for k,v in match.groupdict().items()}).isoformat()
        except ValueError:return None
    text=complete_date(value)
    if text is not None:return text
    if mode=='excel_serial' and isinstance(value,(int,float)) and not isinstance(value,bool) and isfinite(value):
        if 1900<=value<=2200 or value<=0 or value!=int(value):return None
        try:return complete_date(from_excel(value,epoch=epoch))
        except (OverflowError,ValueError):pass
    return None


def parse_mapped_workbook(path, mappings, scan=None, *, observations=None, issues=None):
    """Parse bounded regions with row dates or a source-cell date shared by categories."""
    path=Path(path).expanduser().resolve()
    current=scan_workbook(path)
    if scan is not None and scan.get('sha256')!=current['sha256']:
        raise ValueError('源文件已变化，请重新扫描后导入')
    byname={s['name']:s for s in current['sheets']}
    if not isinstance(mappings,list) or not mappings:raise ValueError('至少选择一个工作表映射')
    if len({m.get('sheet') for m in mappings if isinstance(m,dict)})!=len(mappings):raise ValueError('工作表映射重复或无效')
    checked=[]
    for m in mappings:
        if not isinstance(m,dict) or m.get('sheet') not in byname:raise ValueError('映射工作表不存在')
        checked.append(validate_mapping(m,byname[m['sheet']]))
    result={'source_file':str(path),'sha256':current['sha256'],'rule_version':RULE_VERSION,
        'sheets':[],'series':[],'observations':[] if observations is None else observations,
        'issues':[] if issues is None else issues,'excluded_regions':[],
        'charts':{'status':'inventory_only','visual_content_parsed':False,
            'sheets':[dict(sheet=m['sheet'],**byname[m['sheet']]['charts']) for m in checked]}}
    vw=load_workbook(path,read_only=True,data_only=True,keep_links=False)
    fw=load_workbook(path,read_only=True,data_only=False,keep_links=False)
    try:
        for m in checked:
            name=m['sheet'];ws=vw[name];fs=fw[name]
            from .regions import auction_profile, validate_region_mapping
            profile=auction_profile(ws)
            validate_region_mapping(m,profile)
            counts=Counter();firsts={};formula_count=0;obs_start=len(result['observations']);series_start=len(result['series'])
            rowblocks=[(bi,b) for bi,b in enumerate(m['blocks']) if b.get('layout')!='matrix']
            for bi,b in enumerate(m['blocks']):
                if b.get('layout')=='matrix':
                    from .matrix import parse_matrix
                    parse_matrix(ws,fs,b,bi,name,result)
            start=min((b['data_start_row'] for bi,b in rowblocks),default=1)
            end=max((b.get('data_end_row') or ws.max_row for bi,b in rowblocks),default=1)
            maxcol=max([column_index_from_string(c['column']) for bi,b in rowblocks for c in b['columns']]+
                       [column_index_from_string(b[k]) for bi,b in rowblocks for k in ('date_column','category_column') if b.get(k) and (k!='date_column' or b['date_mode'] not in ('shared_cell','none'))]+[1])
            # Read declared header/date anchors once; templates verify local headers too.
            refs={cell for bi,b in rowblocks for c in b['columns'] for cell in c['header_cells']}
            refs.update(b['date_cell'] for bi,b in rowblocks if b['date_mode']=='shared_cell')
            refs.update(cell for bi,b in rowblocks for cell in b.get('header_expectations',{}))
            maxref=max([int(re.search(r'[0-9]+',x).group()) for x in refs] or [1])
            refwidth=max([column_index_from_string(re.match(r'[A-Z]+',x).group()) for x in refs] or [1])
            refvalues={f'{get_column_letter(ci)}{rn}':cell.value for rn,row in enumerate(ws.iter_rows(max_row=maxref,max_col=refwidth),1) for ci,cell in enumerate(row,1) if f'{get_column_letter(ci)}{rn}' in refs}
            anchor_dates={}
            for bi,b in rowblocks:
                for cell,expected in b.get('header_expectations',{}).items():
                    if str(refvalues.get(cell))!=expected:raise ValueError(f'{name}!{cell} 区块表头已变化，模板不能直接复用，请重新确认')
                if b['date_mode']=='shared_cell':
                    value=_date(refvalues.get(b['date_cell']),'custom' if b.get('date_format') else 'date',vw.epoch,b.get('date_format'))
                    if not value:raise ValueError(f'{name}!{b["date_cell"]} 不是完整日期，不能作为整个小表的日期')
                    anchor_dates[bi]=value
            seen={};series_map={};last_dates={};block_counts=Counter()
            def add_series(bi,b,c,category,category_cell):
                key=(bi,c['column'],category)
                if key not in series_map:
                    stable=f"{name}|{bi}|{c['column']}|{c['metric']}"+(f'|{category}' if category is not None else '')
                    sid='generic.'+sha256(stable.encode()).hexdigest()[:20]
                    dims=dict(c['dimensions']);dims['block']=f'mapped_{bi+1}'
                    if b.get('name'):dims['data_region']=b['name']
                    if category is not None:
                        catname=b.get('category_name','类别')
                        if catname in dims:raise ValueError('类别名称与已有维度重名，请修改类别名称')
                        dims[catname]=category
                    date_note=f"共用日期来自{b['date_cell']}" if b['date_mode']=='shared_cell' else ('无时间字段；按原表记录查询' if b['date_mode']=='none' else f"时间标签取{b['date_column']}列")
                    item={'series_id':sid,'sheet':name,'metric':c['metric'],'dimensions':dims,
                        'unit':c.get('unit'),'unit_status':'explicit' if c.get('unit') else 'needs_review',
                        'header_cells':c['header_cells'],'date_mode':b['date_mode'],
                        'region_name':b.get('name',f'区块{bi+1}'),'block_index':bi,
                        'description':f"{name} / {c['metric']}；{date_note}；"+('单位：'+c['unit'] if c.get('unit') else '单位待确认')}
                    if category is not None:item['category_provenance']={}
                    series_map[key]=item;seen[sid]=set();result['series'].append(item)
                return series_map[key]
            # Preserve empty dated series in the catalogue, as before.
            for bi,b in rowblocks:
                if not b.get('category_column'):
                    for c in b['columns']:add_series(bi,b,c,None,None)
            for rn,(row,frow) in enumerate(zip_longest(ws.iter_rows(min_row=start,max_row=end,max_col=maxcol),fs.iter_rows(min_row=start,max_row=end,max_col=maxcol)) if rowblocks else [],start):
                if row is None or frow is None:raise ValueError('公式与缓存行数不一致')
                for bi,b in rowblocks:
                    if not b['data_start_row']<=rn<=(b.get('data_end_row') or ws.max_row):continue
                    shared=b['date_mode']=='shared_cell'
                    undated=b['date_mode']=='none';period_mode=b['date_mode']=='period'
                    dc=column_index_from_string(b['date_column']) if not shared and not undated else None
                    datecell=b['date_cell'] if shared else (None if undated or period_mode else f"{b['date_column']}{rn}")
                    dateval=anchor_dates[bi] if shared else (None if undated or period_mode else _date(row[dc-1].value,b['date_mode'],vw.epoch,b.get('date_format')))
                    period=str(row[dc-1].value).strip() if period_mode and row[dc-1].value is not None else None
                    if period_mode and has_period_number(row[dc-1].value):period=str(int(row[dc-1].value))
                    cc=column_index_from_string(b['category_column']) if b.get('category_column') else None
                    category=row[cc-1].value if cc else None
                    category_cell=f"{b['category_column']}{rn}" if cc else None
                    has_data=any(row[column_index_from_string(c['column'])-1].value is not None or frow[column_index_from_string(c['column'])-1].data_type=='f' for c in b['columns'])
                    if undated and not has_data and category is None:continue
                    if cc:
                        if isinstance(category,(int,float)) and not isinstance(category,bool):category=str(category)
                        if category is None and not has_data:continue
                        if not isinstance(category,str) or not category.strip():raise ValueError(f'{name}!{category_cell} 类别为空或无效，不能把不同类别的数据合并')
                        category=category.strip()
                    for c in b['columns']:
                        ci=column_index_from_string(c['column']);cell,fc=row[ci-1],frow[ci-1];raw=cell.value
                        formula=fc.value if fc.data_type=='f' else None
                        coord=f"{c['column']}{rn}"
                        if not undated and dateval is None and not (period_mode and period):
                            if raw is not None or formula is not None:
                                previous=last_dates.get(bi)
                                note=f'上一条有效日期位于{previous}。' if previous else ''
                                raise ValueError(f'{name}!{datecell} 日期无效，但同一行{coord}仍有内容{str(raw)[:100]}，该行不是整行空白。{note}已阻断整页导入，请核对区块范围。')
                            continue
                        last_dates[bi]=datecell
                        item=add_series(bi,b,c,category,category_cell);sid=item['series_id']
                        record_key=period if period_mode else (coord if undated else dateval)
                        if record_key in seen[sid]:raise ValueError(f'{name}!{datecell} 指标 {c["metric"]} / {category or ""} 出现重复日期 {dateval}，已阻断整页导入')
                        seen[sid].add(record_key)
                        if cc:item['category_provenance'][coord]=category_cell
                        value=None
                        if formula:formula_count+=1
                        if raw is None:status='formula_no_cache' if formula else 'blank'
                        elif cell.data_type=='e':status='excel_error'
                        elif isinstance(raw,(int,float)) and not isinstance(raw,bool) and isfinite(raw):status='valid';value=raw
                        else:status='unexpected_text' if isinstance(raw,str) else 'invalid_value'
                        if status!='valid':counts[status]+=1;firsts.setdefault(status,coord)
                        result['observations'].append({'series_id':sid,'date':dateval,'period':period,'period_cell':f"{b['date_column']}{rn}" if period_mode else None,'value':value,'source_cell':coord,
                            'date_cell':datecell,'category_cell':category_cell,'category':category,'block_index':bi,
                            'formula':formula,'value_status':status,'raw_value':raw if raw is None or isinstance(raw,(str,int,float,bool)) else str(raw)})
                        block_counts[bi]+=1
            for code,count in counts.items():
                label={'blank':'空白','excel_error':'Excel错误','unexpected_text':'非数值文本'}.get(code,code)
                result['issues'].append({'sheet':name,'cell':firsts[code],'severity':'warning','code':code,'count':count,'message':f'{label} 共{count}处，保留来源；不填零、不重算。'})
            if formula_count:result['issues'].append({'sheet':name,'cell':None,'severity':'warning','code':'cached_formula_values','message':f'{formula_count}个公式使用原文件缓存，未重算。'})
            for bi,b in rowblocks:
                if not block_counts[bi]:raise ValueError(f'{name} 区块{bi+1}没有任何可导入记录')
            ranges=[]
            for b in m['blocks']:
                cols=[column_index_from_string(c['column']) for c in b['columns']]
                ranges.append(f"{get_column_letter(min(cols))}{b['data_start_row']}:{get_column_letter(max(cols))}{b.get('data_end_row') or ws.max_row}")
            result['sheets'].append({'name':name,'rows':ws.max_row,'columns':ws.max_column,
                'observations':len(result['observations'])-obs_start,'series_count':len(result['series'])-series_start,
                'blank_cells':counts['blank'],'non_valid_counts':dict(counts),'formula_cells':formula_count,
                'parsed_range':'; '.join(ranges),'fingerprint':m['fingerprint'],'mapping':m})
            result['excluded_regions'].append({'sheet':name,'range':'outside_explicit_mapping','reason':'仅导入明确区块范围；独立小表使用自己的日期来源，区域外内容不参与查询。'})
            if profile:result['excluded_regions'] += [{'sheet':name,'range':r['range'],'reason':r['reason']} for r in profile['regions'] if r['status']=='needs_review']
    finally:
        vw.close();fw.close()
    return result
