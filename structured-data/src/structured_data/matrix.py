"""Generic, review-only discovery and parsing of horizontal period matrices."""
import re
from hashlib import sha256
from math import isfinite
from openpyxl.utils import get_column_letter, column_index_from_string


def period_label(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return str(int(value)) if value == int(value) and 1900 <= value <= 2200 else None
    text = str(value).strip()
    return text if re.fullmatch(r'(?:\d{4}年?|\d{2,4}[/\-]\d{2,4}(?:[\-／].{1,16})?|\d{4}年\d{1,2}月)', text) else None


def discover_matrices(ws, max_rows=200):
    """Infer candidate bounds only from explicit local headers, independent of sheet name."""
    rows = list(ws.iter_rows(max_row=min(ws.max_row, max_rows), max_col=min(ws.max_column, 512)))
    blocks=[]
    for ri,row in enumerate(rows):
        ci=1
        while ci < len(row):
            if not period_label(row[ci].value) or period_label(row[ci-1].value) or not isinstance(row[ci-1].value,str):
                ci+=1;continue
            end=ci
            while end<len(row) and period_label(row[end].value):end+=1
            if end-ci<3:ci=end;continue
            last=ri+1; rules=[]
            for rj in range(ri+1,len(rows)):
                label=rows[rj][ci-1].value
                vals=[c.value for c in rows[rj][ci:end]]
                if not isinstance(label,str) or not label.strip() or all(period_label(v) for v in vals):break
                if not any(isinstance(v,(int,float)) and not isinstance(v,bool) or isinstance(v,str) and v.startswith('#') for v in vals):break
                # Only parse a unit actually written in the metric label.
                match=re.search(r'[（(]([^（）()]+)[）)]',label)
                unit=match.group(1) if match and match.group(1) in {'吨','万吨','千吨','元','元/吨','%','千英亩','公顷','千公顷','公斤/公顷'} else None
                rules.append({'row':rj+1,'metric':label.strip(),'unit':unit})
                last=rj+1
            if rules:
                title=str(row[ci-1].value).strip()
                blocks.append({'layout':'matrix','name':title,'date_mode':'period',
                    'row_label_column':get_column_letter(ci),'header_rows':[ri+1],
                    'data_start_row':ri+2,'data_end_row':last,'row_rules':rules,
                    'dimensions':{'区域':title},
                    'columns':[{'column':get_column_letter(k+1),'period':period_label(row[k].value),
                                'header_cells':[f'{get_column_letter(k+1)}{ri+1}']} for k in range(ci,end)]})
            ci=end
    return blocks


def validate_matrix(b, sheet):
    from .generic_parser import _column
    rc=_column(b.get('row_label_column'),sheet['columns'])
    start,end=b.get('data_start_row'),b.get('data_end_row')
    if type(start) is not int or type(end) is not int or not 1<=start<=end<=sheet['rows']:
        raise ValueError('矩阵须填写明确且有效的数据起止行')
    headers=b.get('header_rows',[])
    if not headers or any(type(r) is not int or not 1<=r<start for r in headers):raise ValueError('矩阵表头行必须在数据之前')
    if not isinstance(b.get('columns'),list) or not b['columns']:raise ValueError('矩阵至少需要一列期间或版本')
    seen=set();cols=set()
    for c in b['columns']:
        col=_column(c.get('column'),sheet['columns'])
        if col==rc or col in cols:raise ValueError('矩阵数值列重复或与指标名称列重叠')
        cols.add(col)
        p=c.get('period')
        if not isinstance(p,str) or not p.strip() or p.strip() in seen:raise ValueError('矩阵期间/版本标签必须填写且不能重复；相同年度的不同预测版本请分别命名')
        c['period']=p.strip();seen.add(c['period'])
        c.setdefault('header_cells',[f"{c['column']}{headers[-1]}"])
        if not c['header_cells']:raise ValueError('矩阵须保留列标签来源')
        for cell in c['header_cells']:
            mo=re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]*)',cell)
            if not mo or int(mo.group(2)) not in headers:raise ValueError('矩阵列标签来源必须在表头行')
            _column(mo.group(1),sheet['columns'])
    dims=b.setdefault('dimensions',{})
    if not isinstance(dims,dict) or any(not isinstance(k,str) or not isinstance(v,(str,int,float,bool)) for k,v in dims.items()):raise ValueError('矩阵维度必须是简单键值对象')
    expectations=b.get('header_expectations',{})
    if not isinstance(expectations,dict):raise ValueError('矩阵表头核对信息无效')
    for cell,expected in expectations.items():
        mo=re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]*)',cell) if isinstance(cell,str) else None
        if not mo or int(mo.group(2)) not in headers or not isinstance(expected,str):raise ValueError('矩阵表头核对信息必须位于表头行')
        _column(mo.group(1),sheet['columns'])
    rules=b.get('row_rules',[])
    if not isinstance(rules,list):raise ValueError('矩阵行配置必须为列表')
    seen=set()
    for rule in rules:
        r=rule.get('row')
        if type(r) is not int or not start<=r<=end or r in seen:raise ValueError('矩阵指标行重复或超出数据范围')
        seen.add(r)
        if rule.get('metric') is not None and not isinstance(rule['metric'],str):raise ValueError('矩阵指标名称必须是文字')
        if rule.get('unit') is not None and (not isinstance(rule['unit'],str) or not rule['unit'].strip()):raise ValueError('矩阵单位必须为空或文字')
    b['date_mode']='period'


def parse_matrix(ws,fs,b,bi,sheet,result):
    width=max(column_index_from_string(c['column']) for c in b['columns'])
    width=max(width,column_index_from_string(b['row_label_column']))
    end=b['data_end_row']
    rows=list(ws.iter_rows(max_row=end,max_col=width));frows=list(fs.iter_rows(max_row=end,max_col=width))
    for cell,expected in b.get('header_expectations',{}).items():
        # Headers may be sparse; always verify exact coordinates on reuse.
        c=column_index_from_string(re.match('[A-Z]+',cell).group());r=int(re.search('[0-9]+',cell).group())
        if str(rows[r-1][c-1].value)!=expected:raise ValueError(f'{sheet}!{cell} 区块表头已变化，模板不能直接复用')
    rules={r['row']:r for r in b.get('row_rules',[])};seen=set();n=0
    for rn in range(b['data_start_row'],end+1):
        labelcell=f"{b['row_label_column']}{rn}"
        rawlabel=rows[rn-1][column_index_from_string(b['row_label_column'])-1].value
        values=[rows[rn-1][column_index_from_string(c['column'])-1].value for c in b['columns']]
        if rawlabel is None and not any(v is not None for v in values):continue
        if not isinstance(rawlabel,str) or not rawlabel.strip():raise ValueError(f'{sheet}!{labelcell} 指标名称为空或无效，请缩小矩阵范围')
        rule=rules.get(rn,{})
        metric=rule.get('metric') or rawlabel.strip()
        if metric in seen:raise ValueError(f'{sheet} 矩阵指标名称重复：{metric}；请拆分区块或明确指标名称')
        seen.add(metric)
        sid='generic.'+sha256(f'{sheet}|matrix|{bi}|{rn}|{metric}'.encode()).hexdigest()[:20]
        dims={**b.get('dimensions',{}),'block':f'mapped_{bi+1}'}
        if b.get('name'):dims['data_region']=b['name']
        unit=rule.get('unit')
        result['series'].append({'series_id':sid,'sheet':sheet,'metric':metric,'unit':unit,
            'unit_status':'explicit' if unit else 'needs_review','dimensions':dims,'date_mode':'period',
            'block_index':bi,'header_cells':[labelcell]+[p for c in b['columns'] for p in c['header_cells']],
            'periods':[c['period'] for c in b['columns']], 'region_name':b.get('name',''),
            'description':f'{sheet} / {metric}；按原表期间或版本标签查询，不转换成具体日期。'})
        for c in b['columns']:
            ci=column_index_from_string(c['column']);cell=rows[rn-1][ci-1];fc=frows[rn-1][ci-1]
            raw=cell.value;formula=fc.value if fc.data_type=='f' else None
            if raw is None:status='formula_no_cache' if formula else 'blank';value=None
            elif cell.data_type=='e':status='excel_error';value=None
            elif isinstance(raw,(int,float)) and not isinstance(raw,bool) and isfinite(raw):status='valid';value=raw
            else:status='unexpected_text';value=None
            result['observations'].append({'series_id':sid,'date':None,'period':c['period'],
                'period_cell':c['header_cells'][-1],'metric_cell':labelcell,'date_cell':None,
                'source_cell':f"{c['column']}{rn}",'value':value,'formula':formula,'value_status':status,
                'raw_value':str(raw) if raw is not None else None,'block_index':bi})
            n+=1
            if status!='valid':result['issues'].append({'sheet':sheet,'cell':f"{c['column']}{rn}",'code':status,'severity':'warning','message':'原单元格为空、错误或非数值，保留来源；不填零。'})
    if not n:raise ValueError('矩阵没有可导入记录')
    return n
