"""File-bound reviewed region metadata and generic mapping safety checks."""
from copy import deepcopy
from openpyxl.utils import get_column_letter, range_boundaries


def reviewed_regions(path, ws):
    from .profiles.reviewed import reviewed_mapping
    config = reviewed_mapping(path)
    profile = deepcopy(config.get('region_profiles', {}).get(ws.title)) if config else None
    if profile is None:
        return None
    last_row = max(range_boundaries(r['range'])[3] for r in profile['regions'])
    last_col = max(range_boundaries(r['range'])[2] for r in profile['regions'])
    rows = list(ws.iter_rows(max_row=last_row, max_col=last_col, values_only=True))
    for region in profile['regions']:
        c1, r1, c2, r2 = range_boundaries(region['range'])
        region['preview'] = [[{'cell': f'{get_column_letter(c)}{r}',
            'value': value.isoformat() if hasattr(value, 'isoformat') else value}
            for c in range(c1, min(c2, c1 + 7) + 1)
            for value in [rows[r-1][c-1] if r <= len(rows) and c <= len(rows[r-1]) else None]]
            for r in range(r1, min(r2, r1 + 5) + 1)]
    return profile


def apply_profile(sheet, profile):
    sheet['layout'] = 'mixed_regions'
    sheet['status'] = 'needs_review'
    sheet['mapping'] = {'sheet': sheet['name'], 'fingerprint': sheet['fingerprint'], 'blocks': deepcopy(profile['blocks'])}
    sheet['regions'] = profile['regions']
    sheet['region_profile'] = profile['version']
    sheet['reasons'] = ['原文件指纹匹配已核对的区域配置。各区块保留独立时间来源；待核对区域不参与导入。']


def validate_region_mapping(mapping, profile):
    """Protect stored reviewed ranges/date origins, without sheet/column rules."""
    if not profile:
        return
    for block in mapping['blocks']:
        for column in block['columns']:
            allowed = any(
                column['column'] in {c['column'] for c in source['columns']}
                and block['data_start_row'] >= source['data_start_row']
                and block.get('data_end_row') is not None
                and block['data_end_row'] <= source['data_end_row']
                and all(block.get(key) == source.get(key) for key in
                        ('date_mode', 'date_column', 'date_cell', 'category_column'))
                for source in profile['blocks'])
            if not allowed:
                raise ValueError(f'已核对配置包含独立小表，{column["column"]}列的范围或时间来源超出已核对区域；请对照各区块重新确认，不能整列套用其他区域的日期。')
