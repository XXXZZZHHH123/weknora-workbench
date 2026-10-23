"""Serial, bounded AI-assisted import; source data and authorization stay authoritative."""
from __future__ import annotations

from copy import deepcopy
from datetime import date, datetime
import hashlib
import math
import time


NO_DATA_NOTICE = '已检查的原表没有可导入指标数值；保留原件和预览，可手动选择映射后重试。'


def inspect_sheet_data(manager, import_id, name, *, max_cells=1_000_000, max_seconds=30):
    """Only prove absence; numbers, formulas or incomplete scans never approve a sheet.

    Run in a resource-limited review worker. Read formulas as well as values so
    an uncached formula cannot be mistaken for an empty page. No business names,
    numeric column heuristics or model guesses can produce a positive approval.
    """
    from openpyxl import load_workbook

    doc = manager.get(import_id)
    manager.sheet(import_id, name)
    path = manager.source_path(doc)
    with path.open('rb') as stream:
        if hashlib.file_digest(stream, 'sha256').hexdigest() != doc['sha256']:
            raise ValueError('原文件已改变，请重新扫描。')
    started = time.monotonic()
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    count = 0
    try:
        ws = book[name]
        # The stored dimension is only a hint and may understate the real XML.
        ws.reset_dimensions()
        for row in ws.iter_rows():
            for cell in row:
                count += 1
                if count > max_cells or time.monotonic() - started > max_seconds:
                    return {'status': 'uncertain', 'cells_checked': count - 1,
                            'message': '完整原表检查达到资源上限，保留待确认；不据局部空白跳过。'}
                value = cell.value
                if getattr(cell, 'data_type', None) == 'f':
                    return {'status': 'uncertain', 'cells_checked': count,
                            'message': '原表含公式，须核对已有缓存与实际指标；不能当作空白页。'}
                if value is None or isinstance(value, (bool, date, datetime)):
                    continue
                if isinstance(value, (int, float)) and math.isfinite(value):
                    return {'status': 'uncertain', 'cells_checked': count,
                            'message': '原表存在数值，仍须核对真实指标和范围；编号不等于指标。'}
                if isinstance(value, str) and any(char.isdigit() for char in value):
                    return {'status': 'uncertain', 'cells_checked': count,
                            'message': '原表含数字文本，须确认属于指标、日期或编号；不自动排除。'}
        return {'status': 'no_indicator_data', 'cells_checked': count, 'message': NO_DATA_NOTICE}
    finally:
        book.close()


def mark_no_data(manager, import_id, mapping, *, before_write):
    """Persist only re-verified absence, with the same final guard as publication."""
    with manager.lock:
        name = mapping.get('sheet') if isinstance(mapping, dict) else None
        doc = manager.get(import_id)
        sheet = next((row for row in doc['sheets'] if row['name'] == name), None)
        if sheet is None or mapping.get('fingerprint') != sheet['fingerprint']:
            raise ValueError('工作表结构指纹已改变，未记录自动跳过状态。')
        published = manager.publication_view(doc)
        if any(row['name'] == name and row.get('published_in_dataset') for row in published['sheets']):
            raise ValueError('工作表已在当前快照中，不修改已有导入数据或跳过状态。')
        if mapping.get('blocks'):
            checked = manager.preview(import_id, mapping)
            if checked.get('valid_observation_count'):
                raise ValueError('原映射存在实际数值，不能记录为空数据跳过。')
            message = NO_DATA_NOTICE
        else:
            evidence = inspect_sheet_data(manager, import_id, name)
            if evidence.get('status') != 'no_indicator_data':
                raise ValueError('原表可能有数据或证据不完整，保留待确认，不记录自动跳过。')
            message = evidence['message']
        before_write()
        sheet['automatic_skip'] = {'code': 'no_indicator_data', 'message': message,
            'verified_source_sha256': doc['sha256'], 'verified_fingerprint': sheet['fingerprint']}
        sheet['import_state'] = 'skipped'
        manager._save(doc)
        return {'sheet': name, **sheet['automatic_skip']}


class _Stopped(Exception):
    def __init__(self, code, message):
        super().__init__(message)
        self.code, self.message = code, message


def import_all_with_ai(manager, import_id, *, mappings=(), save_template=True,
                       total_max_requests=256, per_sheet_max_requests=12,
                       task, check_write, cancelled, progress=None, sheet_result=None,
                       max_seconds=3600, assist=None):
    """Import the current workbook in source order, one limited worker at a time.

    The caller holds the scoped import lock and propagates its captured request
    identity. ``check_write`` refreshes it before each sheet and at publication.
    Completed sheets are durable; every stop returns a truthful partial result.
    An explicit draft is never applied to an already-published sheet in a bulk
    operation: replacement remains a separate single-sheet decision.
    """
    from .ai_review import AIRequestBudget, assist_mapping
    from .review_errors import error_location

    if (type(total_max_requests) is not int or not 1 <= total_max_requests <= 512
            or type(per_sheet_max_requests) is not int or not 1 <= per_sheet_max_requests <= 32):
        raise ValueError('批量 AI 总预算须为 1 至 512 次，单表预算须为 1 至 32 次。')
    doc = manager.publication_view(manager.get(import_id))
    drafts = {}
    for mapping in mappings:
        name = mapping.get('sheet') if isinstance(mapping, dict) else None
        if not isinstance(name, str) or name in drafts:
            raise ValueError('批量 AI 的工作表草稿无效或重复。')
        manager.sheet(import_id, name)
        drafts[name] = deepcopy(mapping)
    sheets = doc['sheets']
    deadline = time.monotonic() + max_seconds
    used = 0
    details = []
    imported, existing, failed, skipped, pending = [], [], [], [], []
    stopped = None
    notice = '逐表检查；仅来源证据明确且通过全量校验的表会导入，其他表保留待确认。'
    assist = assist or assist_mapping
    skip_source_verified = False

    def guard():
        if cancelled():
            raise _Stopped('cancelled', '已取消；已完成工作表保留，后续工作表未导入。')
        if time.monotonic() >= deadline:
            raise _Stopped('timeout', '批量任务达到时间上限；已导入数据保留，其余工作表待确认。')
        try:
            check_write()
        except Exception:
            raise _Stopped('permission_denied', '编辑权限或知识库归属已改变；停止后续导入，原有数据保留。') from None

    def call(action, mapping, **kwargs):
        guard()
        return task(action, import_id=import_id, mapping=mapping,
                    timeout=min(300, max(.001, deadline - time.monotonic())),
                    # Read-only children can stop immediately. A publication
                    # uses its final guard instead: killing between the atomic
                    # snapshot and review pointer writes could hide a success.
                    cancel_check=(None if action in {'confirm', 'mark_skip'} else
                                  lambda: cancelled() or time.monotonic() >= deadline), **kwargs)

    def record(name, status, message='', requests_used=0, **extra):
        row = {'sheet': name, 'status': status, 'message': str(message)[:2000],
               'requests_used': requests_used, **extra}
        details.append(row)
        if sheet_result:
            sheet_result(deepcopy(row))

    def notify(index, name, stage, **extra):
        if progress:
            progress({'completed': index, 'total': len(sheets), 'sheet': name,
                      'stage': stage, 'requests_used': used,
                      'total_max_requests': total_max_requests, **extra})

    def skip(name, mapping, requests_used=0):
        checked = call('mark_skip', mapping, publish_guard=guard)
        message = checked.get('message') or NO_DATA_NOTICE
        skipped.append({'sheet': name, 'code': 'no_indicator_data', 'message': message})
        record(name, 'skipped_no_data', message, requests_used)

    try:
        guard()
        for index, source in enumerate(sheets):
            name = source['name']
            guard()
            notify(index, name, '检查当前工作表')
            if source.get('published_in_dataset'):
                existing.append(name)
                record(name, 'existing', '已存在于当前已校验快照，批量操作不重复导入或覆盖。')
                notify(index + 1, name, '保留已导入工作表')
                continue
            saved_skip = source.get('automatic_skip') or {}
            if (name not in drafts and saved_skip.get('code') == 'no_indicator_data'
                    and saved_skip.get('verified_source_sha256') == doc.get('sha256')
                    and saved_skip.get('verified_fingerprint') == source.get('fingerprint')
                    and isinstance(doc.get('sha256'), str)):
                # The source is immutable through this API. Hash it once before
                # reusing persisted evidence; a changed file never inherits a
                # previous skip and explicit user drafts always recheck.
                if not skip_source_verified:
                    with manager.source_path(doc).open('rb') as stream:
                        if hashlib.file_digest(stream, 'sha256').hexdigest() != doc['sha256']:
                            raise _Stopped('source_changed', '原文件已改变，停止复用跳过记录；请重新扫描。')
                    skip_source_verified = True
                skipped.append({'sheet': name, 'code': 'no_indicator_data', 'message': saved_skip['message']})
                record(name, 'skipped_no_data', saved_skip['message'])
                notify(index + 1, name, '沿用原件校验一致的无数据记录')
                continue
            mapping = deepcopy(drafts.get(name) or source.get('mapping'))
            if not isinstance(mapping, dict):
                mapping = {'sheet': name, 'fingerprint': source.get('fingerprint'), 'blocks': []}
            sheet_used = 0
            try:
                if not mapping.get('blocks'):
                    evidence = call('evidence', {'sheet': name})
                    if evidence.get('status') == 'no_indicator_data':
                        skip(name, mapping)
                        continue

                # A successful syntax/structure review can still contain no
                # observations. Preview counts are the decisive data evidence.
                preview = None
                try:
                    preview = call('preview', mapping)
                except ValueError:
                    pass
                if preview is not None:
                    if not preview.get('series_count') or not preview.get('valid_observation_count'):
                        skip(name, mapping)
                        continue
                    # Scanner-ready or explicitly confirmed drafts may publish
                    # without a model. Pending mappings still require the same
                    # evidence-backed assistance as the per-sheet operation.
                    from .imports import READY
                    state = call('review', mapping)
                    confirmed = (state.get('status') == 'confirmed'
                                 and (source.get('status') in READY or name in drafts))
                else:
                    confirmed = False

                if not confirmed:
                    if used >= total_max_requests:
                        pending.append(name)
                        record(name, 'needs_review', '总 AI 请求预算已用完，本表保留待确认。')
                        continue
                    budget = AIRequestBudget(min(per_sheet_max_requests, total_max_requests - used))

                    def ai_progress(value):
                        guard()
                        notify(index, name, str(value.get('stage') or '逐个辅助核对区块'),
                               block_index=value.get('block_index'),
                               completed_blocks=value.get('completed_blocks'),
                               total_blocks=value.get('total_blocks'))

                    try:
                        proposal = assist(manager, import_id, name, mapping=mapping, budget=budget,
                            cancelled=lambda: cancelled() or time.monotonic() >= deadline,
                            progress=ai_progress,
                            review=lambda candidate: call('review', candidate))
                    finally:
                        sheet_used = budget.used
                        used += budget.used
                    guard()
                    mapping = proposal.get('mapping')
                    state = proposal.get('review_state') or {}
                    if state.get('status') != 'confirmed' or not isinstance(mapping, dict):
                        pending.append(name)
                        record(name, 'needs_review', proposal.get('notice') or '缺少可验证的映射依据，保留待确认。',
                               sheet_used, review_state=state)
                        continue
                    # Never trust a model/helper's claimed confirmation. The
                    # resource-limited real parser rechecks values and bounds.
                    preview = call('preview', mapping)
                    if not preview.get('series_count') or not preview.get('valid_observation_count'):
                        skip(name, mapping, sheet_used)
                        continue
                    state = call('review', mapping)
                    if state.get('status') != 'confirmed':
                        pending.append(name)
                        record(name, 'needs_review', 'AI 建议未通过逐区块真实数据检查，请补齐或排除无数据区块。',
                               sheet_used, review_state=state)
                        continue

                guard()
                call('confirm', mapping, save_template=save_template, publish_guard=guard)
                imported.append(name)
                record(name, 'imported', '已通过全量校验并导入。', sheet_used)
            except _Stopped:
                raise
            except Exception as exc:
                # A worker error must not conceal a permission loss/cancel and
                # let the loop continue with the next sheet.
                guard()
                location = error_location(exc)
                if 'review_state' not in location:
                    try:
                        location['review_state'] = call('review', mapping)
                    except ValueError:
                        pass
                failed.append({'sheet': name, 'message': str(exc)[:2000], **location})
                pending.append(name)
                record(name, 'failed', str(exc), sheet_used, **location)
            finally:
                completed = index + 1 if any(row['sheet'] == name for row in details) else index
                notify(completed, name, '完成当前工作表检查')
    except _Stopped as exc:
        stopped, notice = exc.code, exc.message

    done = {row['sheet'] for row in details}
    for source in sheets:
        name = source['name']
        if name not in done:
            if source.get('published_in_dataset'):
                existing.append(name)
                record(name, 'existing', '原有已导入工作表保持有效。')
            else:
                pending.append(name)
                record(name, 'not_started', notice)
    if used >= total_max_requests and pending and not stopped:
        notice = '已达到总 AI 请求预算；已导入数据保留，其余工作表仍待确认。'
    current = manager.get(import_id)
    return {**manager._result(current), 'sheets_imported': imported, 'sheets_existing': existing,
            'sheets_pending': pending, 'sheets_failed': failed, 'sheets_skipped': skipped,
            'requests_used': used, 'total_max_requests': total_max_requests,
            'details': details, 'cancelled': stopped == 'cancelled', 'stopped': stopped,
            'notice': notice}
