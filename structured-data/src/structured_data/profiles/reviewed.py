"""Execute source-fingerprint-bound reviewed mappings without business branches."""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from hashlib import sha256
from itertools import zip_longest
from math import isfinite
from pathlib import Path
import posixpath
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

def reviewed_mapping(path):
    import hashlib, json
    with Path(path).open('rb') as stream:
        digest = hashlib.file_digest(stream, 'sha256').hexdigest()
    config = Path(__file__).parent.parent / 'reviewed_mappings' / (digest + '.json')
    if not config.is_file():
        return None
    mapping = json.loads(config.read_text())
    if mapping.get('sha256') != digest:
        raise ValueError('已核对映射的文件指纹不匹配')
    return mapping


def _chart_inventory(path: Path, sheet_names) -> dict:
    """Count actual sheet drawing relations, without evaluating or rendering charts."""
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    office = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    relns = "http://schemas.openxmlformats.org/package/2006/relationships"

    def relations(z: ZipFile, part: str) -> dict:
        relpart = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
        if relpart not in z.namelist():
            return {}
        return {
            r.attrib["Id"]: {
                "type": r.attrib.get("Type", "").rsplit("/", 1)[-1],
                "external": r.attrib.get("TargetMode") == "External",
                "target": (r.attrib["Target"].lstrip("/") if r.attrib["Target"].startswith("/")
                           else posixpath.normpath(posixpath.join(posixpath.dirname(part), r.attrib["Target"]))),
            }
            for r in ET.fromstring(z.read(relpart)).findall(f"{{{relns}}}Relationship")
        }

    result = {"status": "inventory_only", "visual_content_parsed": False, "sheets": []}
    with ZipFile(path) as z:
        names = set(z.namelist())
        result["workbook_chart_parts"] = sum(n.startswith("xl/charts/chart") and n.endswith(".xml") and "/_rels/" not in n for n in names)
        rels = relations(z, "xl/workbook.xml")
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        for sheet in wb.findall(f"{{{main}}}sheets/{{{main}}}sheet"):
            if sheet.attrib["name"] not in sheet_names:
                continue
            part = rels[sheet.attrib[f"{{{office}}}id"]]["target"]
            links = relations(z, part)
            chart_parts, image_parts, drawings = set(), set(), []
            for draw in ET.fromstring(z.read(part)).findall(f"{{{main}}}drawing"):
                link = links.get(draw.attrib[f"{{{office}}}id"])
                if not link or link["external"]:
                    continue
                drawings.append(link["target"])
                for item in relations(z, link["target"]).values():
                    if item["external"]:
                        continue
                    if item["type"] == "chart":
                        chart_parts.add(item["target"])
                    elif item["type"] == "image":
                        image_parts.add(item["target"])
            result["sheets"].append({"sheet": sheet.attrib["name"], "drawing_parts": drawings,
                                      "chart_parts": sorted(chart_parts), "image_parts": sorted(image_parts)})
    return result


def parse_workbook(path: Path | str, *, observations=None, issues=None, mapping=None) -> dict:
    """Read known layouts; fail on header drift; never recalculate or embed data.

    Observations retain errors and dated blank cells as null values with source
    coordinates. Formula cache freshness cannot be established and is explicitly
    reported. Queries for latest data must filter value_status='valid'.
    """
    path = Path(path).expanduser().resolve()
    digest = sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(block)
    mapping = reviewed_mapping(path) if mapping is None else mapping
    if not isinstance(mapping, dict) or mapping.get('sha256') != digest.hexdigest():
        raise ValueError('没有绑定此文件指纹的已核对映射，请使用通用导入核对流程')
    sheet_names = [s['name'] for s in mapping['sheets']]
    result: dict[str, Any] = {"source_file": str(path), "sha256": digest.hexdigest(),
        "rule_version": mapping["rule_version"], "sheets": [], "series": [],
        "observations": [] if observations is None else observations,
        "issues": [] if issues is None else issues, "excluded_regions": [], "charts": _chart_inventory(path, sheet_names)}

    def issue(sheet, cell, code, message, severity="warning"):
        result["issues"].append(dict(sheet=sheet, cell=cell, severity=severity, code=code, message=message))

    values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        missing = set(sheet_names) - set(values.sheetnames)
        if missing:
            raise ValueError("缺少必要工作表：" + ", ".join(sorted(missing)))
        for configured in mapping['sheets']:
            name = configured['name']
            ws, fs = values[name], formulas[name]
            import re
            header_rows = [int(re.search(r'[0-9]+$', cell).group()) for cell in configured['header_expectations']]
            heads = list(ws.iter_rows(min_row=1, max_row=max(header_rows), max_col=configured['max_col'], values_only=True))

            def header(row, col):
                return heads[row - 1][col - 1] if col <= len(heads[row - 1]) else None

            def require(row, col, expected):
                actual = header(row, col)
                if actual != expected:
                    raise ValueError(f"表头不匹配 {name}!{get_column_letter(col)}{row}：期望 {expected!r}，实际 {actual!r}；请审核解析规则")

            import re
            from copy import deepcopy
            for coordinate, expected in configured['header_expectations'].items():
                match = re.fullmatch(r'([A-Z]+)([0-9]+)', coordinate)
                from openpyxl.utils import column_index_from_string
                require(int(match.group(2)), column_index_from_string(match.group(1)), expected)
            specs = deepcopy(configured['specs'])
            start_row, max_col = configured['start_row'], configured['max_col']
            if configured.get('excluded_tail'):
                tail = configured['excluded_tail']
                result['excluded_regions'].append({'sheet': name,
                    'range': f"{get_column_letter(tail['start_column'])}1:{get_column_letter(max(tail['start_column'], ws.max_column))}{ws.max_row}",
                    'reason': tail['reason']})
            for col, dc, sid, metric, dimensions, unit, cells in specs:
                result["series"].append({"series_id": sid, "sheet": name, "metric": metric,
                    "dimensions": dimensions, "unit": unit, "unit_status": "explicit" if unit else "needs_review",
                    "header_cells": cells, "description": f"{name} / {metric}；日期取{get_column_letter(dc)}列；" +
                    (f"源表明确单位：{unit}" if unit else "源表未明确单位，尚待确认")})
            blanks, blank_first = Counter(), {}
            errors = Counter()
            obs_start = len(result["observations"])
            formula_count = 0
            zero_fill_reported = False
            dates_seen: dict[str, set[str]] = {}
            for rownum, (row, frow) in enumerate(zip_longest(
                    ws.iter_rows(min_row=start_row, max_col=max_col),
                    fs.iter_rows(min_row=start_row, max_col=max_col)), start_row):
                if row is None or frow is None:
                    raise ValueError(f"{name} 公式与缓存行数不一致")
                for col, dc, sid, metric, dims, unit, hcells in specs:
                    cell, fc = row[col - 1], frow[col - 1]
                    raw, rawdate = cell.value, row[dc - 1].value
                    coord = f"{get_column_letter(col)}{rownum}"
                    formula = fc.value if fc.data_type == "f" else None
                    if isinstance(formula, str) and "ShowBlank==0" in formula and not zero_fill_reported:
                        issue(name, coord, "source_zero_fill", "源工作簿 Wind 查询使用 ShowBlank==0，已有零值可能包含上游缺失填充，无法逆向判别。")
                        zero_fill_reported = True
                    if not isinstance(rawdate, (date, datetime)):
                        if raw is not None or formula is not None:
                            issue(name, coord, "invalid_date", f"数据日期列{get_column_letter(dc)}{rownum}不是日期，未导入此值。")
                        continue
                    date_text = rawdate.date().isoformat() if isinstance(rawdate, datetime) else rawdate.isoformat()
                    if raw is None and formula is None:
                        blanks[sid] += 1
                        blank_first.setdefault(sid, coord)
                    value = None
                    if formula:
                        formula_count += 1
                    if raw is None and formula is None:
                        status = "blank"
                    elif raw is None:
                        status = "formula_no_cache"
                    elif cell.data_type == "e":
                        status = "excel_error"
                    elif isinstance(raw, (int, float)) and not isinstance(raw, bool) and isfinite(raw):
                        status, value = "valid", raw
                    else:
                        status = "unexpected_text" if isinstance(raw, str) else "invalid_value"
                    if status != "valid":
                        errors[status] += 1
                        if status != "blank":
                            issue(name, coord, status, f"原值{str(raw)[:100]!r}；未转换为数字。")
                    seen = dates_seen.setdefault(sid, set())
                    if date_text in seen:
                        issue(name, coord, "duplicate_date", f"系列 {sid} 重复日期 {date_text}，保留来源记录，查询须避免重复汇总。", "error")
                    seen.add(date_text)
                    result["observations"].append({"series_id": sid, "date": date_text, "value": value,
                        "source_cell": coord, "date_cell": f"{get_column_letter(dc)}{rownum}",
                        "formula": formula, "value_status": status,
                        "raw_value": raw if isinstance(raw, (str, int, float, bool)) or raw is None else str(raw)})
            for sid, count in blanks.items():
                issue(name, blank_first[sid], "blank_cells", f"系列 {sid} 有 {count} 个日期对应空单元格，保留空值和来源位置，未填零。")
            if formula_count:
                issue(name, None, "cached_formula_values", f"导入 {formula_count} 个公式单元格的已保存缓存；未重算，缓存新鲜度未知。")
            result["sheets"].append({"name": name, "rows": ws.max_row, "columns": ws.max_column,
                "observations": len(result["observations"]) - obs_start, "series_count": len(specs),
                "blank_cells": sum(blanks.values()), "non_valid_counts": dict(errors), "formula_cells": formula_count,
                "parsed_range": f"A{start_row}:{get_column_letter(max_col)}{ws.max_row}"})
    finally:
        values.close()
        formulas.close()
    return result
