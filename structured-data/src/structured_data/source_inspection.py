"""Bounded evidence from the verified original, including unparsed regions."""
from __future__ import annotations

import datetime as dt
from decimal import Decimal, InvalidOperation, localcontext
import hashlib
import json
import logging
import math
import os
import re
import threading
import time
from contextlib import contextmanager

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from .chart_source_cache import SourceEvidenceReadBudgetExceeded

MAX_SCAN_CELLS = 250_000
MAX_SCAN_ROWS = 50_000
MAX_SCAN_SECONDS = 8
MAX_READ_ADMISSION_WAIT_SECONDS = 8
MAX_RESULT_BYTES = 512 * 1024
CELL = r'[A-Z]{1,3}[1-9][0-9]{0,6}'
EXCEL_CELL = r'\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6}'
logger = logging.getLogger(__name__)


def _configured_source_read_concurrency():
    raw = os.environ.get('STRUCTURED_SOURCE_READ_CONCURRENCY', '1')
    if not re.fullmatch(r'[1-4]', raw):
        raise ValueError('STRUCTURED_SOURCE_READ_CONCURRENCY must be an integer from 1 to 4')
    return int(raw)


SOURCE_READ_CONCURRENCY = _configured_source_read_concurrency()
_SOURCE_READ_SEMAPHORE = threading.BoundedSemaphore(SOURCE_READ_CONCURRENCY)


@contextmanager
def _source_read_admission(scope):
    if not _SOURCE_READ_SEMAPHORE.acquire(timeout=MAX_READ_ADMISSION_WAIT_SECONDS):
        raise SourceEvidenceReadBudgetExceeded(
            'original-source read admission exceeded its wait budget',
            phase='read_admission', scope=scope)
    try:
        yield
    finally:
        _SOURCE_READ_SEMAPHORE.release()


class SourceInspectionReadError(RuntimeError):
    """Sanitized source-read failure with a stable stage and exception class."""

    def __init__(self, phase, cause):
        self.phase = phase
        self.exception_type = type(cause).__name__
        super().__init__('original-source inspection failed')


class SourceVerificationError(ValueError):
    """A source identity/hash check failed without disclosing storage details."""

    phase = 'source_verification'
    exception_type = 'SourceVerificationError'


def _integer(value, low, high, name):
    if type(value) is not int or not low <= value <= high:
        raise ValueError(f'{name} 必须为 {low} 至 {high} 的整数')


def _items(values, name, maximum=8):
    if values is None:
        return []
    if not isinstance(values, list) or len(values) > maximum:
        raise ValueError(f'{name} 最多 {maximum} 项')
    return values


def _bounds(ref):
    if not isinstance(ref, str) or not re.fullmatch(CELL + r'(?:\:' + CELL + r')?', ref):
        raise ValueError('范围必须是 A1 坐标，如 AF77:AI90；不接受公式')
    bounds = range_boundaries(ref)
    left, top, right, bottom = bounds
    if left > right or top > bottom or right > 16_384 or bottom > 1_048_576:
        raise ValueError('Excel 范围无效或超出工作表边界')
    return bounds


def _reference(ref, single=False):
    """Normalize Excel syntax only; a prefix never selects another source file."""
    pattern = (r"(?:(?P<sheet>'(?:[^']|'')*'|[^!']+)!)?"
               + r'(?P<ref>' + EXCEL_CELL + ('' if single else r'(?:\:' + EXCEL_CELL + r')?') + r')')
    matched = re.fullmatch(pattern, ref.strip()) if isinstance(ref, str) else None
    if not matched:
        raise ValueError('坐标必须是 Excel A1 引用，可带同工作表前缀及 $；不接受公式或外部文件')
    prefix = matched.group('sheet')
    if prefix is not None:
        if prefix.startswith("'"):
            prefix = prefix[1:-1].replace("''", "'")
        if not prefix or len(prefix) > 200 or any(char in prefix for char in '[]:*?/\\'):
            raise ValueError('工作表前缀无效；不能引用外部文件或路径')
    normalized = matched.group('ref').replace('$', '').upper()
    _bounds(normalized)
    return normalized, prefix


def _json_value(value):
    if isinstance(value, (dt.date, dt.datetime, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def _chart_evidence(original, sheet, requested_bounds, source_file, sha256):
    """Return fresh chart XML evidence for ranges recalled by a disposable index."""
    if not requested_bounds:
        return [], False, {'complete': False, 'reasons': ['chart_scope_not_requested']}
    from .chart_source_cache import candidate_chart_parts, load_native_chart_snapshot

    evidence = []
    omitted = False
    index_scope = candidate_chart_parts(original, sha256, sheet, requested_bounds)
    snapshot = load_native_chart_snapshot(original, sha256,
                                          parts=index_scope['candidate_parts'])
    charts = snapshot['charts']
    for chart_index, curve_index in snapshot['curves_by_source_sheet'].get(sheet, []):
        chart = charts[chart_index]
        curve = chart.get('curves', [])[curve_index]
        parsed_references = snapshot['parsed_references'][chart_index][curve_index]
        matched = []
        refs = {}
        source_refs = {}
        for role, raw_ref in (('value', curve.get('value_ref')),
                              ('category', curve.get('category_ref'))):
            parsed = parsed_references.get(role)
            if not parsed:
                continue
            source_refs[role] = {'reference': raw_ref, 'range': parsed}
            if parsed['sheet'] != sheet:
                continue
            ref_bounds = (range_boundaries(
                f"{parsed['column']}{parsed['start_row']}:{parsed['column']}{parsed['end_row']}"))
            hits = [name for name, bounds in requested_bounds
                    if max(bounds[0], ref_bounds[0]) <= min(bounds[2], ref_bounds[2])
                    and max(bounds[1], ref_bounds[1]) <= min(bounds[3], ref_bounds[3])]
            if hits:
                refs[role] = {'reference': raw_ref, 'range': parsed}
                matched.extend({'role': role, 'requested_range': name} for name in hits)
        if not matched:
            continue
        if len(evidence) >= 20:
            omitted = True
            continue
        evidence.append({'chart_id': chart['chart_id'], 'chart_sheet': chart['sheet'],
                         'title': chart['title'], 'axis_titles': chart['axis_titles'],
                         'axes': chart.get('axes', []), 'axis_ids': curve.get('axis_ids', []),
                         'part': chart['part'], 'curve_label': curve['label'],
                         'value_ref': curve.get('value_ref'),
                         'category_ref': curve.get('category_ref'),
                         'source_references': source_refs,
                         'intersecting_references': refs,
                         'intersections': matched, 'source_file': source_file,
                         'sha256': sha256,
                         'evidence_note': '原生图表的原始文本与真实引用；程序未确认字段单位或业务含义，须结合系列范围、轴对应及来源上下文判断适用性。'})
    scope_reasons = list(dict.fromkeys(index_scope['reasons'] + snapshot.get('scope_reasons', [])))
    scope = {'complete': bool(index_scope['complete'] and snapshot.get('scope_complete', False)
                              and not scope_reasons),
             'reasons': scope_reasons,
             'candidate_parts': len(index_scope['candidate_parts']),
             'parsed_parts': len(snapshot['charts'])}
    return evidence, omitted, scope


def inspect_query_source_context(store, dataset_id, sheet, records, *, max_cells=8,
                                 max_chart_references=8, query_scope=None):
    """Attach bounded original-source evidence for returned value cells.

    This deliberately reports evidence only. It does not assign units, resolve
    business meaning, or treat empty/truncated chart results as proof that the
    source has no relevant evidence.
    """
    _integer(max_cells, 1, 8, 'max_cells')
    _integer(max_chart_references, 1, 8, 'max_chart_references')
    if not isinstance(sheet, str) or not sheet or not isinstance(records, list):
        return {'kind': 'structured_query_source_evidence', 'status': 'unavailable',
                'reason': 'invalid_query_source_scope', 'truncated': False,
                'notice': '未能建立原件证据范围；查询值仍保留，不能据此认定原件没有单位或图表证据。'}

    cells = []
    seen_cells = set()
    invalid_source_cells = 0
    for record in records:
        if not isinstance(record, dict) or record.get('value') is None:
            continue
        source_cell = record.get('source_cell')
        if not isinstance(source_cell, str) or not re.fullmatch(CELL, source_cell, re.IGNORECASE):
            invalid_source_cells += 1
            continue
        normalized = source_cell.upper()
        if normalized not in seen_cells:
            seen_cells.add(normalized)
            if len(cells) < max_cells:
                cells.append(normalized)
    if not cells:
        return {'kind': 'structured_query_source_evidence', 'status': 'unavailable',
                'reason': 'invalid_or_missing_returned_value_cell',
                'truncated': bool(invalid_source_cells),
                'scope': {'sheet': sheet, 'requested_value_cell_count': 0,
                          'invalid_source_cell_count': invalid_source_cells},
                'notice': '没有可读取的返回值单元格；查询值仍保留，空结果不能证明原件没有证据。'}

    selected = cells[:max_cells]
    value_cells_truncated = len(seen_cells) > len(selected) or bool(invalid_source_cells)
    try:
        inspected = inspect_source(store, dataset_id, sheet=sheet,
                                   ranges=selected, max_cells=len(selected),
                                   max_matches=1, padding_rows=0, padding_cols=0)
    except SourceEvidenceReadBudgetExceeded as exc:
        return {'kind': 'structured_query_source_evidence', 'status': 'unavailable',
                'reason': 'source_read_budget_exceeded',
                'phase': exc.phase,
                'scope': {'sheet': sheet, 'requested_value_cells': selected,
                          'requested_value_cell_count': len(seen_cells),
                          'value_cells_truncated': value_cells_truncated,
                          'failed_source_scope': exc.scope,
                          'chart_scope_complete': False},
                'truncated': True,
                'notice': '原件证据读取超过明确预算；查询值仍保留，相关来源范围未完成核验。此状态不表示原件没有单位或图表证据。'}
    except SourceVerificationError as exc:
        return {'kind': 'structured_query_source_evidence', 'status': 'unavailable',
                'reason': 'original_source_unavailable_or_verification_failed',
                'phase': exc.phase, 'exception_type': exc.exception_type,
                'scope': {'sheet': sheet, 'requested_value_cells': selected,
                          'requested_value_cell_count': len(seen_cells),
                          'value_cells_truncated': value_cells_truncated},
                'truncated': value_cells_truncated,
                'notice': '原件身份或完整性核验未通过；查询值仍保留。此状态不表示原件没有单位或图表证据。'}
    except SourceInspectionReadError as exc:
        logger.exception('Query source evidence failed during %s (cause type %s)',
                         exc.phase, exc.exception_type)
        return {'kind': 'structured_query_source_evidence', 'status': 'unavailable',
                'reason': 'original_source_unavailable_or_verification_failed',
                'phase': exc.phase, 'exception_type': exc.exception_type,
                'scope': {'sheet': sheet,
                          'requested_value_cells': selected,
                          'requested_value_cell_count': len(seen_cells),
                          'invalid_source_cell_count': invalid_source_cells,
                          'value_cells_truncated': value_cells_truncated},
                'truncated': value_cells_truncated,
                'notice': '原件读取未能完成；查询值仍保留。此状态不表示原件没有单位或图表证据。'}
    except Exception as exc:
        # Source inspection is supplementary. Do not leak host paths or let an
        # unavailable/mismatched original replace a successful stored query.
        logger.exception('Unexpected query source evidence failure')
        return {'kind': 'structured_query_source_evidence', 'status': 'unavailable',
                'reason': 'original_source_unavailable_or_verification_failed',
                'phase': 'source_inspection', 'exception_type': type(exc).__name__,
                'scope': {'sheet': sheet,
                          'requested_value_cells': selected,
                          'requested_value_cell_count': len(seen_cells),
                          'invalid_source_cell_count': invalid_source_cells,
                          'value_cells_truncated': value_cells_truncated},
                'truncated': value_cells_truncated,
                'notice': '原件不可用或核验未完成；查询值仍保留。此状态不表示原件没有单位或图表证据。'}

    # A matching category/date range is not unit evidence for the returned
    # value cell; only retain charts whose value series intersects that cell.
    charts = [chart for chart in inspected.get('chart_references', [])
              if 'value' in chart.get('intersecting_references', {})]
    charts_truncated = inspected.get('chart_references_truncated', False) or len(charts) > max_chart_references
    chart_scope = inspected.get('chart_reference_scope', {})
    chart_scope_complete = chart_scope.get('complete') is True
    query_scope = query_scope if isinstance(query_scope, dict) else {}
    truncated = bool(inspected.get('truncated') or value_cells_truncated or charts_truncated
                     or query_scope.get('truncated') or not chart_scope_complete)
    return {
        'kind': 'structured_query_source_evidence',
        'status': 'ok',
        'source': {'file': inspected['source_file'], 'sha256': inspected['sha256'],
                   'sheet': sheet, 'hash_verified': inspected['hash_verified']},
        'scope': {'requested_value_cells': selected,
                  'requested_value_cell_count': len(seen_cells),
                  'invalid_source_cell_count': invalid_source_cells,
                  'value_cells_truncated': value_cells_truncated,
                  'query_rows_are': query_scope.get('rows_are'),
                  'query_record_count': query_scope.get('record_count'),
                  'query_rows_returned_count': len(records),
                  'query_rows_truncated': bool(query_scope.get('truncated')),
                  'query_filters': query_scope.get('filters'),
                  'query_notice': '仅对应query_series返回的行；聚合操作的这些行可能是来源示例，不代表全部聚合输入。',
                  'chart_scope': '仅列出原生图表系列引用与这些实际返回值单元格直接相交的证据',
                  'chart_scope_complete': chart_scope_complete,
                  'chart_scope_reasons': chart_scope.get('reasons', []),
                  'chart_references_truncated': charts_truncated,
                  'truncated': truncated,
                  'original_inspection_truncated': bool(inspected.get('truncated'))},
        'cells': [cell for item in inspected.get('ranges', []) for cell in item.get('cells', [])],
        'chart_references': charts[:max_chart_references],
        'notice': ('原件证据仅补充来源文字、单元格和相交图表原始引用；不自动认定单位或业务含义。'
                   'chart_scope_complete=false、空图表列表或截断结果不能证明原件其他区域/未返回图表中不存在相关证据。')
    }


def _validate_inspection_request(sheet=None, searches=None, ranges=None, sums=None,
                                 max_matches=10, max_cells=200,
                                 padding_rows=2, padding_cols=2):
    """Validate and normalize the caller request before taking a read slot."""
    _integer(max_matches, 1, 30, 'max_matches')
    _integer(max_cells, 1, 400, 'max_cells')
    _integer(padding_rows, 0, 3, 'padding_rows')
    _integer(padding_cols, 0, 4, 'padding_cols')
    terms = _items(searches, 'searches')
    if any(not isinstance(term, str) or not term.strip() or len(term) > 200 for term in terms):
        raise ValueError('searches 每项必须为 1 至 200 字符的字面搜索词')
    requested = _items(ranges, 'ranges')
    groups = _items(sums, 'sums')
    if not terms and not requested and not groups:
        raise ValueError('请提供 searches、ranges 或 sums，不默认返回整份工作簿')
    if sheet is not None and (not isinstance(sheet, str) or not sheet or len(sheet) > 200):
        raise ValueError('工作表名称无效')
    normalized_ranges = [_reference(ref) for ref in requested]
    normalized_groups = []
    for group in groups:
        if not isinstance(group, list) or not 1 <= len(group) <= 16:
            raise ValueError('sums 每组必须包含 1 至 16 个 A1 引用（单格或有限矩形范围）')
        normalized_groups.append([_reference(ref) for ref in group])
    all_references = normalized_ranges + [ref for group in normalized_groups for ref in group]
    prefixes = {prefix for _, prefix in all_references if prefix is not None}
    if len(prefixes) > 1 or (sheet is not None and any(prefix != sheet for prefix in prefixes)):
        raise ValueError('引用包含不同工作表；前缀必须与 sheet 完全一致，不跨工作表合并')
    if all_references and sheet is None:
        if not prefixes or any(prefix is None for _, prefix in all_references):
            raise ValueError('读取范围或计算时必须指定 sheet，或让每个引用都明确包含同一工作表前缀')
        sheet = next(iter(prefixes))
    requested = [ref for ref, _ in normalized_ranges]
    groups = [[ref for ref, _ in group] for group in normalized_groups]
    requested_bounds = [(ref, _bounds(ref)) for ref in requested]
    group_bounds = [[_bounds(ref) for ref in group] for group in groups]
    group_counts = [sum((b[2]-b[0]+1)*(b[3]-b[1]+1) for b in bounds) for bounds in group_bounds]
    if any(count < 2 for count in group_counts):
        raise ValueError('sums 每组展开后至少需要 2 个真实单元格')
    # Count before expansion or any source read: even a syntactically valid
    # million-row range cannot consume memory outside this 400-cell boundary.
    # Overlap between separate groups is counted each time it will be returned.
    required_count = sum((b[2]-b[0]+1)*(b[3]-b[1]+1) for _, b in requested_bounds) + sum(group_counts)
    if required_count > max_cells:
        raise ValueError(f'指定范围及求和输入共 {required_count} 格，超过 max_cells={max_cells}；请缩小范围')
    for bounds in group_bounds:
        seen = set()
        for left, top, right, bottom in bounds:
            for row in range(top, bottom+1):
                for col in range(left, right+1):
                    if (row, col) in seen:
                        raise ValueError('求和组内展开后不能重复坐标，防止范围重叠或重复计数')
                    seen.add((row, col))

    return {'sheet': sheet, 'terms': terms, 'requested': requested,
            'requested_bounds': requested_bounds, 'groups': groups,
            'group_bounds': group_bounds, 'required_count': required_count,
            'max_matches': max_matches, 'max_cells': max_cells,
            'padding_rows': padding_rows, 'padding_cols': padding_cols}


def inspect_source(store, dataset_id, sheet=None, searches=None, ranges=None, sums=None,
                   max_matches=10, max_cells=200, padding_rows=2, padding_cols=2):
    """Never interpret periods, units or competing definitions on the caller's behalf."""
    request = _validate_inspection_request(
        sheet, searches, ranges, sums, max_matches, max_cells, padding_rows, padding_cols)
    scope = {'sheet': request['sheet'], 'ranges': request['requested'],
             'sum_ranges': request['groups'], 'search_count': len(request['terms'])}
    with _source_read_admission(scope):
        return _inspect_source_validated(store, dataset_id, request)


def _inspect_source_validated(store, dataset_id, request):
    sheet = request['sheet']
    terms = request['terms']
    requested = request['requested']
    requested_bounds = request['requested_bounds']
    groups = request['groups']
    group_bounds = request['group_bounds']
    required_count = request['required_count']
    max_matches = request['max_matches']
    max_cells = request['max_cells']
    padding_rows = request['padding_rows']
    padding_cols = request['padding_cols']

    # _path is resolved through the current workspace proxy, never a caller path.
    try:
        folder = store._path(dataset_id)
        manifest = store.manifest(dataset_id)
    except ValueError:
        raise
    except Exception as exc:
        raise SourceInspectionReadError('dataset_lookup', exc) from exc
    original = folder / 'source.xlsx'
    if not original.is_file():
        raise SourceVerificationError('此版本没有保留原 Excel 文件，不能伪造未解析区域的证据')

    def digest():
        try:
            with original.open('rb') as stream:
                return hashlib.file_digest(stream, 'sha256').hexdigest()
        except Exception as exc:
            raise SourceInspectionReadError('source_verification', exc) from exc

    expected = manifest.get('sha256')
    if not expected or digest() != expected:
        raise SourceVerificationError('原 Excel 的 SHA256 与版本记录不一致，拒绝读取')
    started = time.monotonic()
    try:
        formulas = load_workbook(original, read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise SourceInspectionReadError('workbook_open', exc) from exc
    cached = None
    try:
        try:
            cached = load_workbook(original, read_only=True, data_only=True, keep_links=False)
        except Exception as exc:
            raise SourceInspectionReadError('workbook_open', exc) from exc
        if sheet is not None and sheet not in formulas.sheetnames:
            raise ValueError('指定原工作表不存在；不替换为其他工作表')
        selected = [sheet] if sheet else formulas.sheetnames
        output = {'kind': 'structured_source_inspection', 'status': 'ok',
                  'dataset_id': dataset_id, 'source_file': manifest['source_file'],
                  'sha256': expected, 'hash_verified': True, 'sheets': selected,
                  'coverage': {name: {'parsed_ranges': [s['parsed_range'] for s in manifest.get('sheets', [])
                                  if s.get('name') == name and s.get('parsed_range')],
                                    'excluded_regions': [s for s in manifest.get('excluded_regions', [])
                                  if s.get('sheet') == name or s.get('name') == name]}
                               for name in selected},
                  'matches': [], 'ranges': [], 'calculations': [], 'truncated': False,
                  'truncation_reasons': [],
                  'notice': '原件只读证据；公式使用 Excel 保存的缓存值，不重新计算。parsed_ranges 仅代表已解析区域；单位、年度定义和多个表的优先级须依原件或文档核对，不自行推断。'}
        returned = required_count
        reasons = output['truncation_reasons']

        def cell_doc(fc, cc, row, col, name):
            formula = fc.value if fc.data_type == 'f' else None
            value = cc.value if formula else fc.value
            return {'sheet': name, 'cell': f'{get_column_letter(col)}{row}',
                    'value': _json_value(value), 'raw_value': _json_value(fc.value),
                    'formula': formula, 'cached_value': _json_value(cc.value) if formula else None,
                    'value_status': 'missing_formula_cache' if formula and cc.value is None
                                    else 'error' if fc.data_type == 'e' or cc.data_type == 'e' else 'original'}

        def read_range(name, bounds):
            left, top, right, bottom = bounds
            try:
                source_rows = formulas[name].iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right)
                cached_rows = cached[name].iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right)
                cells = []
                for row, (fs, cs) in enumerate(zip(source_rows, cached_rows), top):
                    if time.monotonic() - started > MAX_SCAN_SECONDS:
                        start = f'{get_column_letter(left)}{top}'
                        end = f'{get_column_letter(right)}{bottom}'
                        raise SourceEvidenceReadBudgetExceeded(
                            'original-source cell read exceeded time budget', phase='range_read',
                            scope={'sheet': name, 'range': start if start == end else f'{start}:{end}'})
                    cells.extend(cell_doc(f, c, row, col, name) for col, (f, c) in enumerate(zip(fs, cs), left))
                return cells
            except SourceEvidenceReadBudgetExceeded:
                raise
            except Exception as exc:
                raise SourceInspectionReadError('range_read', exc) from exc

        for ref, bounds in requested_bounds:
            output['ranges'].append({'sheet': sheet, 'range': ref, 'cells': read_range(sheet, bounds)})
        for bounds in group_bounds:
            cells = [cell for ref_bounds in bounds for cell in read_range(sheet, ref_bounds)]
            with localcontext() as context:
                # IEEE doubles may have exponents up to 308; don't round large
                # source integers or cancellation under Decimal's default 28.
                context.prec = 650
                total = Decimal(0)
                for cell in cells:
                    value = cell['value']
                    if isinstance(value, bool) or not isinstance(value, (int, float)) or cell['value_status'] != 'original':
                        raise ValueError(f"{sheet}!{cell['cell']} 不是有效数值；缺失、错误和无公式缓存不补零")
                    try:
                        number = Decimal(str(value))
                    except InvalidOperation:
                        raise ValueError('求和输入不是有限数值') from None
                    if not number.is_finite():
                        raise ValueError('求和输入不是有限数值')
                    total += number
            numeric = float(total)
            output['calculations'].append({'operation': 'sum', 'sheet': sheet, 'cells': cells,
                                           'value': str(total),
                                           **({'numeric_value': numeric} if math.isfinite(numeric) else {}),
                                           'source_file': manifest['source_file'], 'sha256': expected})

        scanned_cells = scanned_rows = 0
        scanned_sheets = []
        incomplete_sheets = []
        if terms:
            stop = False
            normalized = [term.casefold() for term in terms]
            for name in selected:
                ws = formulas[name]
                width = min(ws.max_column or 1, MAX_SCAN_CELLS - scanned_cells)
                if width <= 0:
                    reasons.append('scan_cells_limit'); break
                scanned_sheets.append(name)
                for row_index, row in enumerate(ws.iter_rows(max_col=width), 1):
                    if scanned_rows >= MAX_SCAN_ROWS or time.monotonic() - started > MAX_SCAN_SECONDS:
                        reasons.append('scan_rows_limit' if scanned_rows >= MAX_SCAN_ROWS else 'scan_time_limit')
                        stop = True; break
                    scanned_rows += 1
                    for col, fc in enumerate(row, 1):
                        if scanned_cells >= MAX_SCAN_CELLS:
                            reasons.append('scan_cells_limit'); stop = True; break
                        scanned_cells += 1
                        if fc.value is None:
                            continue
                        text = str(_json_value(fc.value)).casefold()
                        matched = [term for term, normalized_term in zip(terms, normalized) if normalized_term in text]
                        if not matched:
                            continue
                        if len(output['matches']) >= max_matches:
                            reasons.append('max_matches'); stop = True; break
                        left, top = max(1, col-padding_cols), max(1, row_index-padding_rows)
                        right, bottom = min(ws.max_column, col+padding_cols), min(ws.max_row, row_index+padding_rows)
                        context_count = (right-left+1)*(bottom-top+1)
                        if returned + context_count > max_cells:
                            # Keep the matched cell even when its neighbourhood cannot fit.
                            left = right = col; top = bottom = row_index; context_count = 1
                            if 'context_cells_limit' not in reasons:
                                reasons.append('context_cells_limit')
                        if returned + context_count > max_cells:
                            reasons.append('max_cells'); stop = True; break
                        context = read_range(name, (left, top, right, bottom))
                        center = next(c for c in context if c['cell'] == f'{get_column_letter(col)}{row_index}')
                        output['matches'].append({**center, 'matched_terms': matched,
                            'context_range': f'{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}',
                            'context_cells': context})
                        returned += context_count
                    if stop:
                        break
                if width < (ws.max_column or 1):
                    reasons.append('scan_cells_limit'); stop = True
                if stop:
                    incomplete_sheets.append(name)
                    break
        if requested_bounds:
            try:
                chart_refs, chart_refs_omitted, chart_scope = _chart_evidence(
                    original, sheet, requested_bounds, manifest['source_file'], expected)
            except SourceEvidenceReadBudgetExceeded:
                raise
            except Exception as exc:
                raise SourceInspectionReadError('chart_reference_read', exc) from exc
            output['chart_references'] = chart_refs
            output['chart_references_truncated'] = chart_refs_omitted
            if chart_refs_omitted:
                reasons.append('chart_references_limit')
            output['chart_reference_scope'] = chart_scope
            if not chart_scope['complete']:
                reasons.append('chart_scope_incomplete')
        output['truncation_reasons'] = list(dict.fromkeys(reasons))
        output['truncated'] = bool(reasons)
        output['scan'] = {'cells_scanned': scanned_cells, 'rows_scanned': scanned_rows,
                          'sheets_scanned': scanned_sheets, 'complete': not reasons if terms else None,
                          'unscanned_sheets': [name for name in selected if name not in scanned_sheets] if terms else [],
                          'incomplete_sheets': incomplete_sheets,
                          'limits': {'cells': MAX_SCAN_CELLS, 'rows': MAX_SCAN_ROWS, 'seconds': MAX_SCAN_SECONDS},
                          'elapsed_seconds': round(time.monotonic()-started, 3)}
        if output['truncated']:
            chart_omission = ('chart_references_truncated=true 表示超过20条的相交图表曲线已省略。'
                              if output.get('chart_references_truncated') else '')
            output['next_step'] = ('搜索、邻域或图表引用证据可能未完整返回；' + chart_omission +
                                   ('图表候选索引或引用解析范围不完整。' if output.get('chart_reference_scope', {}).get('complete') is False else '') +
                                   'matches=[] 不能证明目标数据或年度标签不存在。'
                                   'scan.unscanned_sheets 是完全未扫描的真实工作表，scan.incomplete_sheets 是尚未扫完的表。'
                                   '已从字段目录定位主题后，请给出该字段的真实 sheet 并沿用搜索词定向检索；'
                                   '不要再次省略 sheet 从工作簿首表重复扫描。匹配或邻域过多时缩小搜索词、邻域或指定 ranges 核对原值。')
        if digest() != expected:
            raise SourceVerificationError('读取期间原 Excel 发生变化，结果已拒绝')
        if len(json.dumps(output, ensure_ascii=False, allow_nan=False).encode()) > MAX_RESULT_BYTES:
            raise ValueError('原件证据超过 512 KiB 输出预算；请缩小范围或邻域，不返回截断的原值')
        return output
    finally:
        formulas.close()
        if cached is not None:
            cached.close()
