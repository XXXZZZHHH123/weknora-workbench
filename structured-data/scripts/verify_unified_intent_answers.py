#!/usr/bin/env python3
"""Grade task intent, independently computed Excel values and cell provenance.

Expected interpretations here are test oracles, never application routing rules.
"""
import argparse
import datetime as dt
import hashlib
import json
import math
from pathlib import Path
import openpyxl


def get_plan(record):
    plans=[]
    for e in record.get('events',[]):
        if e.get('response_type')!='tool_result':continue
        text=e.get('data',{}).get('output','')
        try:obj=json.loads(text[text.index('{'):])
        except (ValueError,TypeError):continue
        if obj.get('kind') in ('structured_field_resolution','structured_tool_notice'):plans.append(obj)
    return plans[-1] if plans else {}


def main():
    p=argparse.ArgumentParser();p.add_argument('results',type=Path);p.add_argument('workbook',type=Path)
    p.add_argument('--cases',type=Path,default=Path(__file__).resolve().parents[1]/'tests/unified_intent_cases_20261010.json')
    args=p.parse_args();bank=json.loads(args.cases.read_text());wb=openpyxl.load_workbook(args.workbook,read_only=True,data_only=True)
    tables={name:list(wb[name].iter_rows(values_only=True)) for name in ('郑棉期货','新疆棉入库和公检量','中国棉花平衡表')};wb.close()
    def column(sheet,col):
        return [dict(date=row[0].strftime('%Y-%m-%d'),value=row[col-1],source_cell=f'{openpyxl.utils.get_column_letter(col)}{n}',date_cell=f'A{n}') for n,row in enumerate(tables[sheet],1) if isinstance(row[0],(dt.date,dt.datetime)) and isinstance(row[col-1],(int,float)) and math.isfinite(row[col-1])]
    price=column('郑棉期货',2);inventory=column('新疆棉入库和公检量',12)
    totals=[column('中国棉花平衡表',i) for i in (3,4)]
    values=[r['value'] for r in price];lo,hi,mean=min(values),max(values),sum(values)/len(values)
    checks=[]
    for case in [c for g in bank['groups'] for c in g['cases']]:
        path=args.results/(case['id']+'.json');issues=[]
        if not path.exists():checks.append(dict(id=case['id'],passed=False,issues=['Missing replay']));continue
        record=json.loads(path.read_text());plan=get_plan(record);a=plan.get('answer',{});kind=case['check'];sources=list(a.get('sources',{}).values())
        def require(ok,msg):
            if not ok:issues.append(msg)
        def close(actual,expected):return isinstance(actual,(int,float)) and math.isclose(actual,expected,rel_tol=1e-10,abs_tol=1e-7)
        require(record.get('status')=='completed_ungraded','Incomplete chat')
        if kind in ('price_choice','time_choice'):
            require(plan.get('status')=='needs_clarification','Expected genuine clarification')
            require(not a,'Executed before clarification')
            if kind=='price_choice':require({c['series_id'] for f in plan.get('fields',[]) for c in f['candidates']}=={'futures.close.CF.CZC','futures.settle.CF.CZC'},'Wrong business alternatives')
            else:require(plan.get('clarification',{}).get('type')=='time_range','Time-only clarification expected')
        else:
            require(plan.get('status')=='resolved',plan.get('message','No resolved result'))
            if kind=='metadata':
                require(a.get('count')==2 and {g['metric'] for g in a.get('groups',[])}=={'收盘价','结算价'},'Wrong price types')
            elif kind.startswith('sum_'):
                start='2022-01-01' if kind=='sum_24' else '2021-01-01';end='2023-12-31'
                wanted=[[r for r in rows if start<=r['date']<=end] for rows in totals]
                require(a.get('start_date')==start and a.get('end_date')==end,'Wrong inherited/updated time range')
                require(close(a.get('value'),sum(r['value'] for rows in wanted for r in rows)),'Wrong two-field total')
                require(a.get('coverage_complete') is False and '已记录月份小计' in a.get('summary',''),'Incomplete calendar total not disclosed')
                require({s.get('series',{}).get('series_id') for s in sources}=={'balance.primary.C','balance.primary.D'},'Wrong balance fields')
                for source in sources:require(source.get('coverage',{}).get('expected_month_count')==(24 if kind=='sum_24' else 36),'Wrong requested month count')
            else:
                if kind not in ('september_trend','september_chart'):require(not a.get('chart') and not a.get('analysis'),'Unrequested analysis/chart')
                expected_sid='generic.a171072121d583505cd7' if kind.startswith('inventory') else 'futures.close.CF.CZC'
                require(bool(sources) and all(s.get('series',{}).get('series_id')==expected_sid for s in sources),'Wrong field or lost context')
                aggregates=[s.get('aggregate') for s in sources]
                if kind=='min':require(len(aggregates)==1 and close(aggregates[0],lo),'Wrong minimum')
                if kind in ('extremes','max_min'):require(len(aggregates)==2 and sorted(aggregates)==[lo,hi],'Missing or wrong extrema')
                if kind=='extremes':require('value' not in a,'Unrequested difference')
                if kind=='max_min':require(close(a.get('value'),hi-lo),'Wrong extreme difference')
                if kind=='max_mean':require(close(a.get('value'),hi-mean),'Wrong maximum-minus-mean')
                if kind in ('latest','inventory_latest'):
                    rows=inventory if kind.startswith('inventory') else price;expected=max(rows,key=lambda r:r['date'])
                    require(len(sources)==1 and len(sources[0].get('rows',[]))==1 and sources[0]['rows'][0]['date']==expected['date'] and close(sources[0]['rows'][0]['value'],expected['value']),'Wrong latest value/date')
                if kind=='last_three':
                    expected=sorted(price,key=lambda r:r['date'])[-3:]
                    require(len(sources)==1 and [(r['date'],r['value']) for r in sources[0].get('rows',[])]==[(r['date'],r['value']) for r in expected],'Wrong last three records')
                if kind=='bottom_three':
                    require(plan.get('query',{}).get('value_filters')==[{'op':'ne','value':0}], 'Excluding zero must not silently exclude negative values')
                    expected=sorted([r for r in price if r['value']!=0],key=lambda r:(r['value'],r['date']))[:3]
                    require(len(sources)==1 and [(r['date'],r['value']) for r in sources[0].get('rows',[])]==[(r['date'],r['value']) for r in expected],'Wrong lowest three nonzero records')
                if kind=='last_five_mean':require(len(aggregates)==1 and close(aggregates[0],sum(r['value'] for r in price[-5:])/5),'Wrong five-record mean')
                if kind.startswith('september_'):
                    month=[r for r in price if '2023-09-01'<=r['date']<='2023-09-30']
                    require(a.get('start_date')=='2023-09-01' and a.get('end_date')=='2023-09-30','Lost September scope')
                    if kind=='september_mean':require(len(aggregates)==1 and close(aggregates[0],sum(r['value'] for r in month)/len(month)),'Wrong monthly mean')
                    else:
                        peak=month[0]['value'];dd=0
                        for r in month:
                            peak=max(peak,r['value']);dd=max(dd,(peak-r['value'])/peak*100)
                        analysis=a.get('analysis',{})
                        require(close(analysis.get('change'),month[-1]['value']-month[0]['value']) and close(analysis.get('max_drawdown_percent'),dd),'Wrong trend/drawdown')
                        if kind=='september_chart':
                            graph=a.get('chart',{}).get('series',[])
                            require(len(graph)==1 and [(r['date'],r['value']) for r in graph[0]['rows']]==[(r['date'],r['value']) for r in month],'Chart differs from saved Excel')
                if kind=='inventory':require(len(aggregates)==1 and close(aggregates[0],max(r['value'] for r in inventory)),'Wrong inventory maximum')
            # Verify every displayed row against its original cell, not another query API.
            for source in sources:
                sheet=source.get('source',{}).get('sheet') or source.get('series',{}).get('sheet')
                if sheet not in tables:issues.append('Unexpected workbook sheet');continue
                for row in source.get('rows',[]):
                    try:
                        col,n=openpyxl.utils.cell.coordinate_from_string(row['source_cell']);original=tables[sheet][n-1][openpyxl.utils.column_index_from_string(col)-1]
                        dcol,dn=openpyxl.utils.cell.coordinate_from_string(row['date_cell']);day=tables[sheet][dn-1][openpyxl.utils.column_index_from_string(dcol)-1]
                        require(close(row['value'],original) and day.strftime('%Y-%m-%d')==row['date'],'Cell value/date mismatch')
                    except (KeyError,IndexError,ValueError,AttributeError):issues.append('Invalid provenance')
        checks.append(dict(id=case['id'],passed=not issues,issues=issues,status=plan.get('status'),seconds=record.get('elapsed_seconds'),resolution_id=plan.get('resolution_id')))
    report=dict(passed=all(c['passed'] for c in checks),source_sha256=hashlib.sha256(args.workbook.read_bytes()).hexdigest(),cases=checks)
    (args.results/'intent-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report,ensure_ascii=False,indent=2))
    if not report['passed']:raise SystemExit(1)

if __name__=='__main__':main()
