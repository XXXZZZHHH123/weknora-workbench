#!/usr/bin/env python3
"""Read-only original workbook evidence, independent of parser/store/resolver.

Run with the bundled workspace Python (openpyxl/numpy). Never recalculates or
modifies Excel; cached values and formula origin are distinguished.
"""
import datetime as dt
import hashlib
import json
from pathlib import Path
import sys
import openpyxl
import numpy as np

SOURCE = Path(sys.argv[1]).resolve()
OUTPUT = Path(sys.argv[2]).resolve()
wb = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
specs = {
    'close': ('郑棉期货', 'A', 'B'), 'close01': ('郑棉期货', 'A', 'C'),
    'settle': ('郑棉期货', 'J', 'K'), 'settle01': ('郑棉期货', 'J', 'L'),
    'settle05': ('郑棉期货', 'J', 'N'),
    'stock': ('新疆棉入库和公检量', 'A', 'L'),
    'entry': ('新疆棉入库和公检量', 'A', 'E'),
    'cum_entry': ('新疆棉入库和公检量', 'A', 'F'),
    'cum_inspect': ('新疆棉入库和公检量', 'A', 'C'),
    'load': ('TTEB下游负荷和库存', 'A', 'B'),
    'raw_stock_days': ('TTEB下游负荷和库存', 'A', 'C'),
    'spot': ('TTEB价格', 'A', 'B'),
    'import_spot': ('进口棉人民币价和进口利润', 'A', 'F'),
    'import_close': ('进口棉人民币价和进口利润', 'A', 'J'),
    'production': ('中国棉花平衡表', 'A', 'C'),
    'import': ('中国棉花平衡表', 'A', 'D'),
}
rows = {k: [] for k in specs}
for sheet in dict.fromkeys(s[0] for s in specs.values()):
    wanted = {k: s for k, s in specs.items() if s[0] == sheet}
    maxcol = max(openpyxl.utils.column_index_from_string(c) for s in wanted.values() for c in s[1:])
    for idx, values in enumerate(wb[sheet].iter_rows(max_col=maxcol, values_only=True), 1):
        for key, (_, dc, vc) in wanted.items():
            date, value = values[openpyxl.utils.column_index_from_string(dc)-1], values[openpyxl.utils.column_index_from_string(vc)-1]
            if isinstance(date, (dt.datetime, dt.date)):
                rows[key].append({'date': date.strftime('%Y-%m-%d'), 'value': value,
                                  'sheet': sheet, 'cell': vc+str(idx), 'date_cell': dc+str(idx)})
wb.close()
numeric = lambda k: [r for r in rows[k] if isinstance(r['value'], (int, float)) and not isinstance(r['value'], bool)]
lookup = lambda k, d: next((r for r in rows[k] if r['date'] == d), None)
extreme = lambda k, fn: [r for r in numeric(k) if r['value'] == fn(x['value'] for x in numeric(k))]
september = [r for r in numeric('close') if r['date'].startswith('2023-09')]
last5 = sorted(numeric('close'), key=lambda r: r['date'])[-5:]
positive = [r for r in numeric('close') if r['value'] > 0]
yearclose = {r['date']: r['value'] for r in numeric('close') if r['date'].startswith('2023-')}
yearload = {r['date']: r['value'] for r in numeric('load') if r['date'].startswith('2023-')}
joint = sorted(set(yearclose) & set(yearload))
peak = september[0]['value']; drawdown = 0; dd_dates = None; peak_date = september[0]['date']
for r in sorted(september, key=lambda r: r['date']):
    if r['value'] > peak: peak, peak_date = r['value'], r['date']
    fall = (peak-r['value']) / peak
    if fall > drawdown: drawdown, dd_dates = fall, [peak_date, r['date']]
close = lookup('close','2023-09-22'); settle = lookup('settle','2023-09-22'); prev = lookup('close','2023-09-21')
truth = {
    'source': str(SOURCE), 'sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    'method': 'openpyxl read-only original saved values; no store/resolver used',
    'latest': {k: max(numeric(k), key=lambda r:r['date']) for k in specs},
    'points': {k: {d: lookup(k,d) for d in ['2022-11-01','2022-11-20','2022-11-21','2023-09-21','2023-09-22','2009-01-06']} for k in specs},
    'stock_max': extreme('stock',max), 'stock_min': extreme('stock',min),
    'close_max': extreme('close',max), 'close_min': extreme('close',min),
    'positive_close_min': [r for r in positive if r['value']==min(x['value'] for x in positive)],
    'top3_close_records': sorted(numeric('close'),key=lambda r:(-r['value'],r['date']))[:3],
    'september_close': september, 'september_mean': sum(r['value'] for r in september)/len(september),
    'last5': last5, 'last5_mean': sum(r['value'] for r in last5)/5,
    '2023_days_gt17000': sum(v>17000 for v in yearclose.values()),
    'settle_minus_close': settle['value']-close['value'],
    'settle_over_close_percent': (settle['value']/close['value']-1)*100,
    'day_change': close['value']-prev['value'], 'day_change_percent': (close['value']/prev['value']-1)*100,
    'inventory_ratio_percent': lookup('cum_inspect','2022-11-20')['value']/lookup('cum_entry','2022-11-20')['value']*100,
    'september_first_last_difference': september[-1]['value']-september[0]['value'],
    'september_max_drawdown_percent': drawdown*100, 'drawdown_dates': dd_dates,
    '2023_price_load_joint_count':len(joint),
    '2023_price_load_pearson':float(np.corrcoef([yearclose[d] for d in joint],[yearload[d] for d in joint])[0,1]),
    'november_cumulative_entry_sum':sum(r['value'] for r in numeric('cum_entry') if r['date'].startswith('2022-11')),
    'balance_2022_23_monthly': {k:[r for r in rows[k] if '2022-09-01'<=r['date']<='2023-08-31'] for k in ('production','import')},
    'balance_2022_23_sum': {k:sum(r['value'] for r in numeric(k) if '2022-09-01'<=r['date']<='2023-08-31') for k in ('production','import')},
}
f = openpyxl.load_workbook(SOURCE, read_only=True, data_only=False)
truth['zero_formula'] = {'sheet':'郑棉期货','cell':'B4484','formula':f['郑棉期货']['B4484'].value}
f.close()
OUTPUT.parent.mkdir(parents=True,exist_ok=True)
OUTPUT.write_text(json.dumps(truth,ensure_ascii=False,indent=2,default=str)+'\n')
print(json.dumps({k:truth[k] for k in ['sha256','last5_mean','september_mean','2023_days_gt17000','day_change','day_change_percent','inventory_ratio_percent','2023_price_load_pearson']},ensure_ascii=False))
