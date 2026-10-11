#!/usr/bin/env python3
"""Original workbook oracle for open-ended agent acceptance. No app/model calls."""
import argparse, datetime as dt, hashlib, json, statistics
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook


def build(source):
    wb=load_workbook(source,read_only=True,data_only=True)
    # Explicit source columns are acceptance evidence, never runtime aliases.
    def extract(sheet,mapping):
        rows=[]
        for n,v in enumerate(wb[sheet].iter_rows(max_col=max(mapping.values()),values_only=True),1):
            if not isinstance(v[0],(dt.datetime,dt.date)):continue
            rows.append(dict(date=v[0].strftime('%Y-%m-%d'),row=n,**{k:v[c-1] for k,c in mapping.items()}))
        return rows
    balance=extract('中国棉花平衡表',{'production':3,'imports':4})
    close=extract('郑棉期货',{'close':2})
    stocks=extract('新疆棉入库和公检量',{'stock':12})
    wb.close()
    b2020=[r for r in balance if r['date']>='2020-01-01']
    b2021=[r for r in balance if '2021-01-01'<=r['date']<='2023-12-31']
    b2022=[r for r in balance if '2022-01-01'<=r['date']<='2023-12-31']
    years=[]
    for year in ['2020','2021','2022']:
        rs=[r for r in balance if r['date'].startswith(year)]
        p=sum(r['production'] for r in rs);i=sum(r['imports'] for r in rs)
        years.append([year,p,i,p+i,100*i/(p+i)])
    imports=[]
    for year in ['2022','2023']:
        rs=[r for r in balance if r['date'].startswith(year) and int(r['date'][5:7])<=9]
        imports.append(sum(r['imports'] for r in rs))
    above=[r for r in balance if '2020-01-01'<=r['date']<='2022-12-31' and r['imports']>r['production']]
    september=[r for r in close if r['date'].startswith('2023-09') and isinstance(r['close'],(int,float))]
    weeks=defaultdict(list)
    for r in september:
        d=dt.date.fromisoformat(r['date']);weeks[(d-dt.timedelta(days=d.weekday())).isoformat()].append(r)
    weekly=[[w,min(r['date'] for r in rs),max(r['date'] for r in rs),len(rs),statistics.mean(r['close'] for r in rs)] for w,rs in sorted(weeks.items())]
    changes=[dict(date=r['date'],previous=prev['date'],value=r['close'],pct=100*(r['close']/prev['close']-1)) for prev,r in zip(september,september[1:])]
    median=statistics.median(r['close'] for r in september);maximum=max(r['close'] for r in september)
    stockmax=max(r['stock'] for r in stocks if isinstance(r['stock'],(int,float)))
    return {'source':str(source),'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'method':'original openpyxl cached values; standard-library calculations; no resolver/store/SQL',
      'coverage2020':dict(first=b2020[0]['date'],last=b2020[-1]['date'],count=len(b2020),production=sum(r['production'] for r in b2020),imports=sum(r['imports'] for r in b2020)),
      'coverage2021':dict(first=b2021[0]['date'],last=b2021[-1]['date'],count=len(b2021),production=sum(r['production'] for r in b2021),imports=sum(r['imports'] for r in b2021)),
      'coverage2022':dict(first=b2022[0]['date'],last=b2022[-1]['date'],count=len(b2022),production=sum(r['production'] for r in b2022),imports=sum(r['imports'] for r in b2022)),
      'annual_rows':years,'imports_jan_sep':imports,'import_change_percent':100*(imports[1]/imports[0]-1),'imports_above_production':above,
      'weekly':weekly,'declines':sorted(changes,key=lambda r:r['pct'])[:3],'last3':september[-3:],'last3_mean':statistics.mean(r['close'] for r in september[-3:]),
      'median':median,'maximum':maximum,'maximum_dates':[r['date'] for r in september if r['close']==maximum],'max_minus_median':maximum-median,
      'stock_max':stockmax,'stock_max_dates':[r['date'] for r in stocks if r['stock']==stockmax]}

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('source',type=Path);p.add_argument('output',type=Path);a=p.parse_args()
    a.output.parent.mkdir(parents=True,exist_ok=True);a.output.write_text(json.dumps(build(a.source),ensure_ascii=False,indent=2)+'\n')
    print('Independent original-source oracle saved')
