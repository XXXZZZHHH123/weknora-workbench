"""Immutable bytes and exact locations; never reconstruct historical evidence."""
import hashlib
import io
import json
import math
import mimetypes
import os
from datetime import date, datetime, timezone
from pathlib import Path
import uuid

from sqlalchemy import select, or_

from tin.config import settings
from tin.models import Observation, ObservationSource, SourceArtifact, SourceBatch
from tin.ownership import scoped_directory, verified_actor

MAX_SOURCE_BYTES = 100 * 1024 * 1024
MAX_SNAPSHOT_CELLS = 2_000_000


def research_visible(artifact):
    """A KB-only center upload never implicitly becomes a research source."""
    return not artifact.details.get("library_restricted") or artifact.details.get("shared_research") is True


def research_visible_clause():
    return or_(SourceArtifact.details["library_restricted"].as_boolean().is_(None),
               SourceArtifact.details["library_restricted"].as_boolean() == False,
               SourceArtifact.details["shared_research"].as_boolean() == True)


def _now():
    return datetime.now(timezone.utc)


def _json(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def stable_hash(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                    separators=(",", ":"), default=_json).encode()).hexdigest()


def source_family(filename):
    import re
    name = Path(filename).name.casefold()
    return re.sub(r"20\d{2}[-_.]?\d{2}[-_.]?\d{2}(?:[-_ ]?\d{6})?", "{date}", name)


def storage_root(session):
    return scoped_directory(Path(settings.data_dir) / "sources", session=session)


def artifact_path(session, artifact):
    root = storage_root(session).resolve()
    path = (root / artifact.storage_key).resolve()
    if not path.is_relative_to(root):
        raise PermissionError("来源路径不合法")
    return path


def verified_artifact_bytes(session, artifact):
    path = artifact_path(session, artifact)
    data = path.read_bytes()
    if len(data) != artifact.size or hashlib.sha256(data).hexdigest() != artifact.id:
        raise ValueError("来源原件与哈希版本不符，拒绝读取或导入")
    return data


def _atomic_bytes(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        if path.read_bytes() != data:
            raise ValueError("不可变来源文件内容已改变，拒绝覆盖")
        return
    temporary = path.with_name(path.name + "." + uuid.uuid4().hex + ".tmp")
    try:
        with temporary.open("xb") as stream:
            stream.write(data)
        temporary.chmod(0o600)
        try:
            os.link(temporary, path)  # no overwrite even across worker processes
        except FileExistsError:
            if path.read_bytes() != data:
                raise ValueError("不可变来源文件内容冲突")
    finally:
        temporary.unlink(missing_ok=True)


def preserve_blob(session, data: bytes, filename: str, *, kind="upload", source_url=None, metadata=None):
    if len(data) > MAX_SOURCE_BYTES:
        raise ValueError("来源文件超过 100 MiB，请分批导入")
    digest = hashlib.sha256(data).hexdigest()
    artifact = session.get(SourceArtifact, digest)
    key = f"objects/{digest[:2]}/{digest}.bin"
    _atomic_bytes(storage_root(session) / key, data)
    if artifact is None:
        values = dict(id=digest, filename=Path(filename).name[:254] or "source.bin",
            kind=kind, media_type=mimetypes.guess_type(filename)[0] or "application/octet-stream",
            size=len(data), storage_key=key, source_url=source_url, details=metadata or {}, created_at=_now())
        dialect = session.get_bind().dialect.name
        if dialect in {"sqlite", "postgresql"}:
            if dialect == "sqlite":
                from sqlalchemy.dialects.sqlite import insert
            else:
                from sqlalchemy.dialects.postgresql import insert
            session.execute(insert(SourceArtifact).values(**values).on_conflict_do_nothing(index_elements=[SourceArtifact.id]))
            artifact = session.get(SourceArtifact, digest)
        else:
            artifact = SourceArtifact(**values)
            session.add(artifact); session.flush()
    alias = {"filename": Path(filename).name[:254], "kind": kind, "source_url": source_url,
             "metadata": metadata or {}}
    # One immutable byte object may be received through several independent
    # files, feeds or knowledge-base links. Keep their capture descriptions.
    aliases = list(artifact.details.get("aliases", []))
    if alias not in aliases:
        aliases.append(alias)
        artifact.details = {**artifact.details, "aliases": aliases}
    return artifact


def start_batch(session, artifact, *, kind, mapping=None, key=None):
    mapping = mapping or {}
    actor = verified_actor(session, "研究导入")
    key = stable_hash({"actor": actor, "batch": key if key is not None else {
        "artifact": artifact.id, "kind": kind, "mapping": mapping}})
    row = session.scalar(select(SourceBatch).where(SourceBatch.idempotency_key == key))
    if row is None:
        values = dict(id=uuid.uuid4().hex, artifact_id=artifact.id, idempotency_key=key,
                         kind=kind, actor=actor, status="preview",
                         mapping=mapping, result={}, created_at=_now())
        dialect = session.get_bind().dialect.name
        if dialect in {"sqlite", "postgresql"}:
            if dialect == "sqlite":
                from sqlalchemy.dialects.sqlite import insert
            else:
                from sqlalchemy.dialects.postgresql import insert
            session.execute(insert(SourceBatch).values(**values).on_conflict_do_nothing(index_elements=[SourceBatch.idempotency_key]))
            row = session.scalar(select(SourceBatch).where(SourceBatch.idempotency_key==key))
        else:
            row = SourceBatch(**values); session.add(row); session.flush()
    return row


def record_with_evidence(session, obs, *, artifact, locator, batch_id=None,
                         mapping_version="v1", transformation=None, cache=None, evidence_ids=None):
    from tin.ingest.record import record
    row = record(session, obs, cache=cache)
    written = row is not None
    if row is None:
        row = cache.get((obs.series_id, obs.as_of)) if cache is not None else session.scalar(
            select(Observation).where(Observation.series_id == obs.series_id, Observation.as_of == obs.as_of)
            .order_by(Observation.revision.desc()).limit(1))
    session.flush()
    transform = {"source": obs.source, "source_url": obs.source_url,
                 "caliber": obs.caliber.dump(), "normalized_value": obs.value,
                 "normalized_as_of": obs.as_of.isoformat(), **(transformation or {})}
    evidence_id = stable_hash({"observation": row.id, "artifact": artifact.id, "locator": locator,
                              "mapping_version": mapping_version, "transformation": transform})
    exists = evidence_id in evidence_ids if evidence_ids is not None else session.get(ObservationSource, evidence_id) is not None
    if not exists:
        values = dict(id=evidence_id, observation_id=row.id, artifact_id=artifact.id,
            batch_id=batch_id, locator=locator, mapping_version=mapping_version,
            transformation=transform, created_at=_now())
        dialect = session.get_bind().dialect.name
        if dialect in {"sqlite", "postgresql"}:
            if dialect == "sqlite":
                from sqlalchemy.dialects.sqlite import insert
            else:
                from sqlalchemy.dialects.postgresql import insert
            session.execute(insert(ObservationSource).values(**values).on_conflict_do_nothing(index_elements=[ObservationSource.id]))
        else:
            session.add(ObservationSource(**values))
        session.flush()
        if evidence_ids is not None:
            evidence_ids.add(evidence_id)
    return row, written


def observation_sources(session, observation_id):
    rows = session.scalars(select(ObservationSource).where(
        ObservationSource.observation_id == observation_id).order_by(ObservationSource.created_at)).all()
    from tin.integration.identity import research_url
    return [{"artifact_id": r.artifact_id, "sha256": r.artifact_id,
             "url": f"/platform/commodity-research?source={r.artifact_id}",
             "filename": r.locator.get("filename") or session.get(SourceArtifact, r.artifact_id).filename,
             "download_url": research_url(f"/api/sources/{r.artifact_id}/download"),
             "locator": r.locator, "mapping_version": r.mapping_version,
             "transformation": r.transformation, "batch_id": r.batch_id} for r in rows
            if research_visible(session.get(SourceArtifact, r.artifact_id))]


def workbook_cells(data, filename):
    """Yield every worksheet and cell, preserving uncached formulas and blanks."""
    suffix = Path(filename).suffix.lower()
    from openpyxl.utils import get_column_letter
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        raw = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for name in raw.sheetnames:
                ws, values = raw[name], cached[name]
                for r, (original, calculated) in enumerate(zip(ws.iter_rows(), values.iter_rows()), 1):
                    for c, (cell, value) in enumerate(zip(original, calculated), 1):
                        formula = cell.value if cell.data_type == "f" else None
                        status = ("formula_uncached" if formula and value.value is None else
                                  "error" if value.data_type == "e" else "empty" if value.value is None else "valid")
                        yield {"sheet": name, "row": r, "column": c,
                               "address": f"{get_column_letter(c)}{r}", "raw_value": _json(cell.value),
                               "value": _json(value.value), "formula": formula,
                               "data_type": cell.data_type, "number_format": cell.number_format, "status": status}
        finally:
            raw.close(); cached.close()
    elif suffix == ".xls":
        import xlrd
        wb = xlrd.open_workbook(file_contents=data)
        from tin.ingest.excel_importer import _xls_cell
        for ws in wb.sheets():
            for r in range(ws.nrows):
                for c in range(ws.ncols):
                    cell = ws.cell(r, c)
                    value = _xls_cell(cell, wb.datemode)
                    yield {"sheet": ws.name, "row": r + 1, "column": c + 1,
                           "address": f"{get_column_letter(c+1)}{r+1}", "raw_value": _json(cell.value),
                           "value": _json(value), "formula": None,
                           "formula_status": "xls_formula_text_unavailable_in_reader; original bytes retained",
                           "data_type": cell.ctype, "status": "empty" if cell.ctype in {0, 6} else "valid"}
    elif suffix in {".csv", ".tsv"}:
        import csv
        from tin.ingest.excel_importer import _decode_csv
        text = _decode_csv(data)
        reader = csv.reader(io.StringIO(text), delimiter="\t" if suffix == ".tsv" else ",")
        for r, row in enumerate(reader, 1):
            for c, value in enumerate(row, 1):
                yield {"sheet": "CSV", "row": r, "column": c, "address": f"{get_column_letter(c)}{r}",
                       "raw_value": value, "value": value, "formula": None,
                       "data_type": "text", "status": "empty" if not value.strip() else "valid"}
    else:
        raise ValueError("仅表格文件可以生成单元格快照")


def snapshot_workbook(session, artifact, *, filename=None):
    if artifact.snapshot_key:
        verified_artifact_bytes(session, artifact)
        return artifact
    key = f"cells/{artifact.id}.jsonl"
    path = storage_root(session) / key
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + "." + uuid.uuid4().hex + ".tmp")
    sheets = {}
    try:
        with temporary.open("x", encoding="utf-8") as stream:
            for count, cell in enumerate(workbook_cells(verified_artifact_bytes(session, artifact), filename or artifact.filename), 1):
                if count > MAX_SNAPSHOT_CELLS:
                    raise ValueError("工作簿超过 200 万单元格，原件已保存，请拆分后解析")
                state = sheets.setdefault(cell["sheet"], {"rows": 0, "columns": 0, "cells": 0})
                state.update(rows=max(state["rows"], cell["row"]), columns=max(state["columns"], cell["column"]), cells=state["cells"]+1)
                stream.write(json.dumps(cell, ensure_ascii=False) + "\n")
        temporary.chmod(0o600)
        os.replace(temporary, path)
        artifact.snapshot_key = key
        with path.open("rb") as stream:
            checksum = hashlib.file_digest(stream, "sha256").hexdigest()
        artifact.details = {**artifact.details, "sheets": sheets, "reader_version": "cells-v1", "snapshot_sha256": checksum}
        session.flush()
        return artifact
    finally:
        temporary.unlink(missing_ok=True)


def source_cells(session, artifact, *, sheet=None, offset=0, limit=200):
    verified_artifact_bytes(session, artifact)
    if not artifact.snapshot_key:
        return {"cells": [], "total": 0, "gap": "此原件没有表格单元格快照"}
    root = storage_root(session).resolve()
    path = (root / artifact.snapshot_key).resolve()
    if not path.is_relative_to(root):
        raise PermissionError("来源路径不合法")
    with path.open("rb") as stream:
        actual_hash = hashlib.file_digest(stream, "sha256").hexdigest()
    if actual_hash != artifact.details.get("snapshot_sha256"):
        raise ValueError("单元格快照缺少校验记录或已被改动；未返回替代内容")
    rows, count = [], 0
    with path.open(encoding="utf-8") as stream:
        for line in stream:
            cell = json.loads(line)
            if sheet is not None and cell["sheet"] != sheet:
                continue
            if offset <= count < offset + limit:
                rows.append(cell)
            count += 1
    return {"cells": rows, "total": count, "offset": offset, "truncated": count > offset+limit}
