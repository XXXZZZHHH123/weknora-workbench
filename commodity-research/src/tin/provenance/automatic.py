"""Locations are resolved from actual response bytes, not historical guesses."""
import csv
import io
import json
from datetime import datetime


def locate_document(data, obs, *, prepared=None):
    """Reuse a run-local parsing index while keeping locators tied to these bytes.

    A caller processing one immutable response supplies the same ``prepared``
    dictionary for its observations. No cache survives the import or crosses
    source documents/workspaces.
    """
    prepared = {} if prepared is None else prepared
    if "text" not in prepared:
        prepared["text"] = data.decode("utf-8", errors="replace")
    text = prepared["text"]
    if obs.source == "SHFE" and (".main." in obs.series_id or obs.caliber.contract):
        if "json" not in prepared:
            prepared["json"] = json.loads(text)
        payload = prepared["json"]
        contract = obs.caliber.contract
        key = {"close": "CLOSEPRICE", "settle": "SETTLEMENTPRICE", "volume": "VOLUME", "oi": "OPENINTEREST"}[obs.series_id.rsplit(".", 1)[1]]
        if "shfe_contracts" not in prepared:
            contracts = {}
            for index, row in enumerate(payload["o_curinstrument"]):
                if row.get("PRODUCTID") == "sn_f":
                    contracts.setdefault(str(row.get("DELIVERYMONTH")), (index, row))
            prepared["shfe_contracts"] = contracts
        selected = prepared["shfe_contracts"].get(contract[2:])
        if selected:
            index, row = selected
            return {"type": "json", "pointer": f"/o_curinstrument/{index}/{key}",
                        "date_pointer": "/report_date", "raw_value": row.get(key),
                        "raw_date": payload["report_date"], "record": row,
                        "selection": "按最大持仓选择主力" if ".main." in obs.series_id else "具体合约"}
    if obs.source == "CFETS":
        if "json" not in prepared:
            prepared["json"] = json.loads(text)
        payload = prepared["json"]
        for index, row in enumerate(payload["records"]):
            if row.get("vrtEName") == "USD/CNY":
                return {"type": "json", "pointer": f"/records/{index}/price", "raw_value": row["price"],
                        "date_pointer": "/data/lastDate", "raw_date": payload["data"]["lastDate"], "record": row}
    if obs.source in {"FRED", "CBOE"}:
        index_key = f"csv:{obs.source}"
        if index_key not in prepared:
            from openpyxl.utils import get_column_letter
            reader = csv.DictReader(io.StringIO(text))
            headers = reader.fieldnames or []
            value_field = "CLOSE" if obs.source == "CBOE" else headers[-1]
            date_field = "DATE" if obs.source == "CBOE" else headers[0]
            value_column = get_column_letter(headers.index(value_field)+1)
            date_column = get_column_letter(headers.index(date_field)+1)
            fmt = "%m/%d/%Y" if obs.source == "CBOE" else "%Y-%m-%d"
            locations = {}
            for number, row in enumerate(reader, 2):
                raw_date = row.get(date_field, "")
                try:
                    day = datetime.strptime(raw_date, fmt).date()
                except ValueError:
                    continue
                # Preserve the original first-match rule for duplicate dates.
                locations.setdefault(day, {"type": "csv", "sheet": "CSV", "row": number,
                        "value_cell": f"{value_column}{number}",
                        "date_cell": f"{date_column}{number}",
                        "header_cells": [f"{value_column}1"],
                        "raw_value": row[value_field], "raw_date": raw_date, "record": row})
            prepared[index_key] = locations
        located = prepared[index_key].get(obs.as_of.date())
        if located is not None:
            return located
    if obs.source == "SHFE":
        from lxml import html
        if "html" not in prepared:
            prepared["html"] = html.fromstring(text)
        root = prepared["html"]
        for tno, table in enumerate(root.xpath("//table")):
            rows = table.xpath(".//tr")
            if not any("锡" == cell.text_content().strip() for cell in table.xpath(".//td|.//th")):
                continue
            for rno, row in enumerate(rows):
                cells = row.xpath("./td|./th")
                if cells and cells[0].text_content().strip() == "总计":
                    matches = []
                    for cno, cell in enumerate(cells):
                        raw = cell.text_content().strip()
                        try:
                            if float(raw.replace(",", "")) == obs.value:
                                matches.append({"xpath": f"(//table)[{tno+1}]//tr[{rno+1}]/*[{cno+1}]", "raw_value": raw})
                        except ValueError:
                            continue
                    return {"type": "html", "table_index": tno, "row_index": rno,
                            "row_text": [c.text_content().strip() for c in cells],
                            "field": "期货" if obs.series_id.endswith("warrant") else "本周库存/小计",
                            "candidates": matches, "location_status": "exact" if len(matches)==1 else "row_and_header; candidate_values_ambiguous"}
    # No synthetic cell locations: retain the original and identify the parser input.
    return {"type": "response", "location_status": "record_location_unavailable",
            "series_id": obs.series_id, "as_of": obs.as_of.isoformat()}
